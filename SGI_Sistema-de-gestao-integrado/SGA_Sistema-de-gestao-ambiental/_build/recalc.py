"""Recalcula ficheiros .xlsx no Microsoft Excel (COM) e lista erros de fórmula."""
import os
import sys
import openpyxl
import win32com.client

ERRS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")


def recalc(paths):
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        for p in paths:
            wb = xl.Workbooks.Open(os.path.abspath(p))
            xl.CalculateFull()
            wb.Save()
            wb.Close(False)
    finally:
        xl.Quit()


def check(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    wbf = openpyxl.load_workbook(path)
    errors, nform = [], 0
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    nform += 1
                    if "NA()" in c.value:  # #N/A intencional (séries de gráficos)
                        continue
                    v = wb[ws.title][c.coordinate].value
                    if isinstance(v, str) and v in ERRS:
                        errors.append(f"{ws.title}!{c.coordinate} {v} :: {c.value[:90]}")
    return nform, errors


if __name__ == "__main__":
    paths = sys.argv[1:]
    recalc(paths)
    for p in paths:
        n, e = check(p)
        print(f"{os.path.basename(p)}: {n} fórmulas, {len(e)} erros")
        for x in e[:25]:
            print("   ", x)
