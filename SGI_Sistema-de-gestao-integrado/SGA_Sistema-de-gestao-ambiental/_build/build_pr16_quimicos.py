"""PR-SGA-16 — Gestão de produtos químicos (REACH, CLP, SVHC e risco químico): procedimento-guia para os engenheiros.
Chamado por build_docs.build(); também pode correr sozinho: python build_pr16_quimicos.py [--force]"""
import os
from docx.shared import Cm, Pt
import build_docs as bd
from build_docs import new_doc, h, para, bullets, table, history

CODE = "PR-SGA-16"
TITLE = "Gestão de produtos químicos — REACH, CLP, SVHC e risco químico"
FNAME = "PR-SGA-16_Gestao_de_produtos_quimicos_REACH_CLP_SVHC_e_risco_quimico.docx"


def steps(doc, items):
    for i, s in enumerate(items, start=1):
        doc.add_paragraph(f"{i}. {s}").paragraph_format.left_indent = Cm(0.4)


def sub(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(10.5)


def build():
    doc = new_doc(CODE, TITLE, "00", "2026-09-24", "Em vigor")

    h(doc, "1. Objetivo")
    para(doc, "Definir como a Plasticom identifica, aprova, regista, usa, armazena e controla os produtos químicos para cumprir o REACH, o CLP e a "
              "legislação portuguesa de segurança e saúde (agentes químicos e cancerígenos/reprotóxicos), e que evidência fica no registo RG-SGA-21. "
              "É o guia de trabalho dos engenheiros e técnicos: cada secção diz o que fazer, em que folha do Excel e com que critério.")

    h(doc, "2. Âmbito")
    para(doc, "Todas as substâncias, misturas, gases, polímeros e aditivos usados na Unidade 1 (produção, manutenção, utilidades, laboratório), os agentes "
              "gerados pelo processo (fumos de moldação) e as embalagens produzidas (artigos) quanto a SVHC. Os produtos usados por prestadores de "
              "serviços (limpeza) são controlados pelo RG-SGA-11 (LIM-01), com lista de FDS entregue pelo prestador.")

    h(doc, "3. Referências")
    bullets(doc, [
        "Reg. (CE) 1907/2006 (REACH) — arts. 3.º, 5.º–8.º, 31.º–39.º, 56.º, 67.º; anexos XIV e XVII (entradas 51, 63, 70, 74, 78, 79)",
        "Reg. (UE) 2020/878 (anexo II do REACH — formato da FDS)",
        "Reg. (CE) 1272/2008 (CLP); Reg. Delegado (UE) 2023/707 (novas classes de perigo); Reg. (UE) 2024/2865 e Reg. (UE) 2025/2439 (adiamento para 2028)",
        "Diretiva 2008/98/CE, art. 9.º (SCIP) e DL 102-D/2020 (RGGR); Reg. (UE) 2019/1021 (POP); Reg. (UE) 528/2012 e DL 140/2017 (biocidas)",
        "Reg. (UE) 2025/2365 (perdas de granulado); Reg. (UE) 2025/40 (PPWR), art. 5.º; Regs. (CE) 1935/2004, (UE) 10/2011 e (UE) 2024/3190 (contacto alimentar, BPA)",
        "DL 293/2009 (execução do REACH; autoridades DGAE, APA, DGS; helpdesk DGAE) e DL 220/2012 (execução do CLP)",
        "DL 24/2012 (agentes químicos) e alterações; DL 301/2000 na redação do DL 102/2024 (cancerígenos, mutagénicos e reprotóxicos); Lei 102/2009; Código do Trabalho, art. 62.º",
        "DL 150/2015 (Seveso III); DL 220/2008 e Portaria 1532/2008 (SCIE); DL 236/2003 (ATEX); DL 41-A/2010 (ADR)",
        "NP EN 689:2018 (estratégia de medição); NP 1796:2014 (valores-limite de exposição)",
        "ECHA: Guia para utilizadores a jusante; Guia prático sobre cenários de exposição; Guia sobre substâncias em artigos; Guia de registo; "
        "'A segurança química na sua empresa — Introdução para PME' (ECHA-15-B-02-PT, pasta REACH); guia de relatório de microplásticos",
        "ONU — GHS, 10.ª revisão ('Purple Book', pasta REACH); INRS ND 2233 e ED 6485 (Seirich) — método de avaliação do risco químico",
        "PR-SGA-03 (obrigações de conformidade), PR-SGA-04 (resíduos), PR-SGA-05 (emergências), PR-SGA-06 (alterações), PR-SGA-09 (controlo operacional e fornecedores)",
    ])

    h(doc, "4. Definições")
    table(doc, ["Termo", "Significado prático"], [
        ("Substância / mistura", "Substância: elemento ou composto (ex.: ciclo-hexanona). Mistura: duas ou mais substâncias (ex.: tinta, diluente)."),
        ("Artigo", "Objeto cuja forma determina a função mais do que a composição (ex.: frasco, tampa, foil de hot stamping). Não tem FDS; tem declaração SVHC."),
        ("Polímero / monómero", "Polímeros estão isentos de registo; os monómeros (≥ 2% e ≥ 1 t/ano) e os aditivos não estão (art. 6.º, n.º 3)."),
        ("Utilizador a jusante (UJ)", "Quem usa químicos na atividade industrial sem os fabricar nem importar — é o papel normal da Plasticom."),
        ("Importador", "Quem introduz a mercadoria no território aduaneiro da UE. Comprar a um fornecedor de fora do EEE torna a Plasticom importadora, salvo RU ou importador UE."),
        ("Representante único (RU)", "Entidade da UE nomeada pelo fabricante extra-UE para registar em seu nome; o importador coberto passa a UJ."),
        ("SVHC", "Substância que suscita elevada preocupação, na lista candidata da ECHA (253 entradas em 04/02/2026)."),
        ("FDS / FDS alargada / CE", "Ficha de dados de segurança (16 secções). A FDS alargada tem cenários de exposição (CE) em anexo: condições em que o uso é seguro."),
        ("PROC / ERC / SU", "Descritores de uso da ECHA (categoria de processo, de libertação para o ambiente, setor de uso)."),
        ("CMR 1A/1B", "Cancerígeno, mutagénico ou tóxico para a reprodução comprovado/presumido (H340, H350, H360) — regime do DL 301/2000."),
        ("VLE", "Valor-limite de exposição profissional (média de 8 h — VLE-MP; curta duração — VLE-CD)."),
        ("RSQ-UJ", "Relatório de segurança química do utilizador a jusante (art. 37.º, n.º 4), para usos não cobertos."),
    ], [4, 13])

    h(doc, "5. Responsabilidades")
    table(doc, ["Função", "Responsabilidade"], [
        ("Gestor do SGA / EHS", "Dono do procedimento e do RG-SGA-21; papel REACH, matriz legal, avaliação de riscos químicos, Seveso, calendário, relatório de microplásticos."),
        ("Técnico de Qualidade/Ambiente", "Verificação das FDS (checklist) e pedidos de correção aos fornecedores."),
        ("Responsável de Compras", "Não encomendar produto novo sem aprovação; verificar origem EEE/RU; obter declarações SVHC/PFAS/FCM."),
        ("Engenheiros de processo e manutenção (Produção, Manutenção, R&D)", "Propor o produto, descrever o uso real (tarefa, PROC, duração, quantidade), verificar os cenários de exposição e implementar as medidas."),
        ("Responsável de R&D", "Triagem SVHC dos artigos (embalagens), comunicação art. 33.º e SCIP; escolha de materiais sem SVHC/PFAS."),
        ("Responsável de Armazém e Logística", "Receção (rótulo = FDS), armazenagem e compatibilidades, stock máximo por local."),
        ("Responsável de Recursos Humanos / Medicina do Trabalho", "Formações obrigatórias (diisocianatos, CMR), vigilância da saúde, proteção da maternidade e de menores, lista de expostos a CMR."),
        ("Diretor Industrial", "Aprova produtos 'Aprovado condicionado' e as substituições; decide sobre CMR/SVHC."),
    ], [5, 12])

    h(doc, "6. Descrição do processo")

    sub(doc, "6.1 Determinar o papel REACH da Plasticom (folha Papel_REACH)")
    para(doc, "O papel define as obrigações. Pergunte, por esta ordem, para cada compra ou atividade:")
    steps(doc, [
        "A Plasticom fabrica a substância? Não (a moldação transforma polímeros em artigos) → não é fabricante.",
        "O fornecedor está fora do EEE (UE + Islândia, Listenstaine, Noruega)? Se sim, quem é o importador na fatura/incoterm/declaração aduaneira? "
        "Se for a Plasticom e não houver RU com a Plasticom na lista, a Plasticom é IMPORTADORA (ver 6.6).",
        "Coloca químicos no mercado (vende, cede, reembala para terceiros)? Não → não é distribuidora nem formuladora. Misturar tinta com diluente para uso próprio não é formular.",
        "Produz objetos (frascos, tampas, potes)? Sim → produtora de artigos (ver 6.7).",
        "Usa granulado de polímeros numa instalação industrial? Sim → relatório anual de microplásticos à ECHA até 31/05 (anexo XVII, entrada 78) e Reg. 2025/2365.",
    ])
    para(doc, "Registar/atualizar a folha Papel_REACH sempre que mude a origem de uma compra (ex.: Brexit tornou o Reino Unido país terceiro) ou o tipo de atividade.")

    sub(doc, "6.2 Aprovação prévia de um produto químico novo ou alterado (porta de entrada)")
    steps(doc, [
        "O engenheiro que precisa do produto pede ao fornecedor a FDS em português (e, se existir, a FDS alargada com os cenários de exposição) e a ficha técnica. "
        "Nenhum produto é encomendado nem entra na fábrica sem este passo (Compras bloqueia a encomenda).",
        "Criar a linha no Inventario (ID QUI-nnn) com Estado_Aprovacao = 'Em avaliação' e preencher as colunas brancas a partir das secções 1, 2, 3, 9 e 15 da FDS. "
        "As colunas cinzentas (perigoso, classe, CMR, sensibilizante, ambiente, proteção da maternidade, alerta) calculam-se sozinhas.",
        "Registar os componentes perigosos (secção 3) na folha Componentes: a coluna SVHC_Vigiada cruza o CAS com a lista vigiada; confirmar sempre na lista completa da ECHA.",
        "Verificar a FDS (6.4) e, se houver CE, o cenário de exposição (6.5).",
        "Triagem de exclusão — o produto NÃO é aprovado se: for CMR 1A/1B (H340, H350, H360) e existir alternativa; contiver SVHC ≥ 0,1% ou substância do anexo XIV; "
        "violar uma restrição do anexo XVII; contiver PFAS em produtos para a linha alimentar; ou for H334 (sensibilizante respiratório) sem medidas de engenharia. "
        "Sem alternativa técnica: 'Aprovado condicionado' pelo Diretor Industrial, com plano de substituição no PAM.",
        "Fazer a avaliação de risco da tarefa (6.9), definir local de armazenagem compatível (6.10), confirmar Seveso (6.11) e formação necessária (6.13).",
        "Se o produto altera processos ou aspetos ambientais, abrir também uma alteração no RG-SGA-18 (PR-SGA-06).",
        "Aprovar: Estado_Aprovacao = 'Aprovado' ou 'Aprovado condicionado', com data. Só então Compras encomenda e o posto recebe a FDS e a ficha-resumo.",
    ])

    sub(doc, "6.3 Inventário (folha Inventario) — como preencher")
    table(doc, ["Coluna", "O que escrever"], [
        ("Tipo_Produto", "Substância, mistura, mistura (aerossol), gás, polímero ou artigo (artigos não têm FDS)."),
        ("Origem_EEE / Papel_REACH", "País do fornecedor; se 'Não', preencher também Registo_Importacao."),
        ("Qtd_Ano_kg / Stock_Max_kg", "Consumo anual e quantidade máxima presente (é esta que conta para o Seveso e a retenção)."),
        ("Volume_Embalagem_L / Local", "Maior embalagem e local (ARM-nn) — calculam a retenção necessária na folha Armazenagem."),
        ("Frases_H", "Códigos da secção 2 separados por ';' exatamente como na FDS (ex.: H225; H319; EUH066). Vazio = não classificado."),
        ("Categoria_Seveso", "Categoria do anexo I do DL 150/2015 (ex.: P5c para líquidos inflamáveis cat. 2/3; E1 para H400/H410)."),
        ("SVHC_Declarado / PFAS", "Sim / Não / Por confirmar, a partir da secção 3/15 ou da declaração do fornecedor."),
        ("Anexo_XVII / Diisocianatos", "Entrada de restrição aplicável; 'Sim' em diisocianatos ≥ 0,1% obriga a formação (6.13)."),
    ], [4.5, 12.5])
    para(doc, "Rever o inventário todo uma vez por ano (OBR-11 do RG-SGA-04) e retirar produtos que já não se usam (Estado 'Descontinuado' — não apagar a linha: a informação conserva-se 10 anos, art. 36.º).")

    sub(doc, "6.4 Verificação da FDS (folha FDS_Verificacao)")
    steps(doc, [
        "Uma linha por VERSÃO de FDS recebida (nunca apagar a versão anterior; a fórmula marca automaticamente a mais recente como Em_Vigor).",
        "Responder aos 12 pontos (Sim / Não / N.A.): língua portuguesa; formato do Reg. 2020/878 (16 secções e subsecções, incl. 9, 11.2 e 12.6 — desregulação endócrina); "
        "identificação igual ao rótulo; telefone de emergência (CIAV 800 250 250); secção 2 igual ao rótulo; secção 3 com CAS/CE, concentrações, n.º de registo e SVHC; "
        "secção 8 com VLE/DNEL e EPI concretos; secções 11 e 12 completas; uso da Plasticom identificado; CE em anexo quando aplicável; secção 15 com restrições/autorização.",
        "Qualquer 'Não' → Estado 'Pedir correção ao fornecedor': enviar pedido escrito (art. 34.º) e registar a data em Pedido_Fornecedor. Não aceitar FDS de outro país sem tradução.",
        "Idade: o REACH não fixa prazo de validade da FDS (o fornecedor atualiza quando há nova informação — art. 31.º, n.º 9). A Plasticom usa 36 meses (parâmetro P-02) "
        "como critério INTERNO para pedir confirmação de que a versão é a mais recente.",
        "Garantir a FDS (ou ficha-resumo do posto) acessível aos trabalhadores (art. 35.º) — coluna Disponivel_Posto.",
    ])

    sub(doc, "6.5 Cenários de exposição (folha Cenarios_Exposicao)")
    para(doc, "Quando a FDS traz cenários de exposição (substâncias registadas ≥ 10 t/ano e perigosas), o utilizador a jusante tem de verificar se o seu uso está coberto "
              "e aplicar as condições no prazo de 12 meses a contar da receção (arts. 37.º, n.º 4, e 39.º). Passos (guia prático da ECHA):")
    steps(doc, [
        "Descrever o uso real: tarefa, PROC (ex.: PROC 10 aplicação com rolo/pano; PROC 7 pulverização industrial; PROC 8b trasfega; PROC 15 laboratório), "
        "duração por turno, quantidade, concentração, temperatura, ventilação/captação e EPI.",
        "Comparar com o título, os descritores e as condições operacionais e medidas de gestão de risco do CE (comparação direta).",
        "Se as condições reais forem diferentes, verificar se estão dentro do CE por escalonamento (scaling) segundo as regras indicadas pelo fornecedor, "
        "ou estimar a exposição com ECETOC TRA.",
        "Resultado: Coberto / Parcialmente coberto / Não coberto. Se não coberto escolher UMA ação: implementar as condições do CE; pedir ao fornecedor a inclusão do uso "
        "(art. 37.º, n.º 2); elaborar RSQ-UJ (art. 37.º, n.º 4) e comunicá-lo à ECHA em 6 meses (art. 38.º); substituir o produto/fornecedor; ou deixar de usar.",
        "O prazo legal e o estado calculam-se sozinhos; qualquer 'Prazo legal ultrapassado' é uma não conformidade legal (abrir ação no PAM).",
    ])

    sub(doc, "6.6 Registo REACH — quando é preciso e como se faz (folha Registo_Importacao)")
    para(doc, "Regra de ouro: 'sem dados, não há mercado' — uma substância fabricada ou importada ≥ 1 t/ano por uma entidade jurídica tem de estar registada por ela. "
              "Para a Plasticom a pergunta é sempre: quem é o importador das compras de fora do EEE?")
    steps(doc, [
        "Para cada fornecedor preencher a linha: país, material, t/ano, importador legal, RU (nome e data da carta) e se a carta cita a Plasticom. A coluna Conclusao indica "
        "'N.A.' (EEE), 'Coberto por RU', 'Utilizador a jusante (importador: …)' ou 'LACUNA'.",
        "LACUNA → primeiro tentar resolver na cadeia: pedir ao fabricante extra-UE que nomeie um RU e inclua a Plasticom na lista de importadores cobertos; ou comprar a um importador UE. "
        "Até estar resolvido, não importar quantidades que somem ≥ 1 t/ano da substância em causa.",
        "Se a Plasticom tiver mesmo de registar (monómeros ≥ 2% e ≥ 1 t/ano, ou aditivos ≥ 1 t/ano não registados a montante):",
        "   a) Criar a conta da entidade jurídica em ECHA Accounts / REACH-IT e o objeto de entidade jurídica no IUCLID (dimensão PME comprovada para taxa reduzida).",
        "   b) Identificar a substância (CAS/CE, composição) e verificar no REACH-IT se já existe registo conjunto; fazer o pedido de informação (inquiry, art. 26.º) se necessário.",
        "   c) Contactar o registante principal (lead registrant) e comprar a carta de acesso (LoA) aos dados, na banda de tonelagem aplicável (1–10, 10–100, 100–1000, > 1000 t/ano).",
        "   d) Preparar o dossier IUCLID de membro: identidade, tonelagem, usos (descritores), orientação de uso seguro; a partir de 10 t/ano também o relatório de segurança química.",
        "   e) Submeter no REACH-IT, pagar a taxa, responder à verificação de integridade técnica (completeness check) e guardar o n.º de registo (01-2119…-xx).",
        "   f) Manter o registo atualizado (art. 22.º: mudança de tonelagem, novos usos, nova informação) e responder a pedidos de avaliação da ECHA.",
        "Recuperação/reciclado: a isenção de registo de substâncias recuperadas (art. 2.º, n.º 7, al. d) só se aplica à recuperação na UE; reciclado importado de fora do EEE segue as regras de importação. "
        "Em caso de dúvida pedir parecer ao helpdesk REACH/CLP da DGAE.",
        "Outros envios à ECHA com o mesmo acesso REACH-IT: relatório anual de microplásticos (6.8) e, se aplicável, notificação SCIP e art. 7.º, n.º 2 (6.7).",
    ])

    sub(doc, "6.7 SVHC em artigos — art. 7.º, n.º 2, art. 33.º e SCIP (Declaracoes_SVHC no RG-SGA-21; conclusão por família em RG-SGA-20 tbl_familias_ppwr)")
    steps(doc, [
        "Em janeiro e junho (atualizações da lista candidata) atualizar o parâmetro P-01 e a lista vigiada (OBR-19 do RG-SGA-04).",
        "Pedir a cada fornecedor declaração SVHC que indique a versão da lista coberta (n.º de entradas), mais metais pesados (PPWR), PFAS, BPA, ftalatos e DoC para contacto alimentar. "
        "Declarações com versão anterior ficam 'Desatualizada — pedir nova'.",
        "Para cada família de embalagens calcular a concentração máxima de cada SVHC no ARTIGO (massa de SVHC ÷ massa do artigo; acórdão C-106/14: o limiar aplica-se a cada artigo). "
        "Ex.: frasco de 25 g com 0,15 g de tinta UV curada a 3% de fotoiniciador → 0,018% no frasco decorado.",
        "Se > 0,1% m/m: comunicar aos clientes (art. 33.º; ao consumidor em 45 dias, a pedido), notificar a base SCIP e, se > 1 t/ano de SVHC no total, notificar a ECHA em 6 meses (art. 7.º, n.º 2).",
        "'Por confirmar' não é aceitável para expedir: obter declaração ou análise laboratorial (ex.: ftalatos por GC-MS, Pb/Cd por XRF).",
    ])

    sub(doc, "6.8 Restrições, autorização e outras obrigações de produto (folha Requisitos_Legais)")
    bullets(doc, [
        "Anexo XIV (autorização): nenhum componente do anexo XIV sem autorização própria ou do fornecedor (notificar a ECHA — art. 66.º).",
        "Anexo XVII — relevantes para a Plasticom: 23 (cádmio em plásticos), 51 (ftalatos ≥ 0,1% em materiais plastificados), 63 (chumbo em PVC), 70 (D4/D5/D6), 74 (diisocianatos — formação), "
        "78 (microplásticos — relatório anual até 31/05 dos dados do ano anterior no REACH-IT), 79 (PFHxA).",
        "POP (Reg. 2019/1021): excluir UV-328, PFOA, PFHxS, decaBDE nas declarações.",
        "Contacto alimentar: só substâncias da lista da União (Reg. 10/2011), sem BPA/bisfenóis (Reg. 2024/3190), declaração de conformidade dos fornecedores.",
        "Biocidas: só produtos autorizados, usados conforme o rótulo (Reg. 528/2012).",
    ])

    sub(doc, "6.9 Avaliação de risco químico — toxicidade (folhas Avaliacao_Risco, Metodologia e Medicoes_VLE)")
    para(doc, "O DL 24/2012 obriga o empregador a avaliar e registar os riscos de todos os agentes químicos perigosos — incluindo os gerados no processo "
              "(ex.: acetaldeído na purga das ISBM) — e o DL 301/2000 (redação do DL 102/2024) acrescenta regras para cancerígenos, mutagénicos e reprotóxicos. "
              "O registo é feito por tarefa:")
    steps(doc, [
        "Uma linha por tarefa × produto: posto, função exposta, n.º de trabalhadores.",
        "A classe de perigo (1–5) vem do inventário (a frase H/EUH mais grave); para agentes gerados usar Classe_Manual.",
        "Escolher volatilidade/estado físico, procedimento (fechado, fechado com aberturas, aberto, dispersivo), proteção coletiva (a mais eficaz existente e com manutenção), "
        "superfície de pele exposta e frequência — SEM contar com os EPI.",
        "O Excel calcula o risco por inalação (10^(classe−1) × volatilidade × procedimento × proteção) e cutâneo (10^(classe−1) × superfície × frequência): "
        "> 1000 → P1 (medidas imediatas); 100–1000 → P2 (plano ≤ 6 meses); < 100 → P3 (manter controlos). O método (adaptado de INRS ND 2233/Seirich) serve para PRIORIZAR; não substitui medições.",
        "Regras que prevalecem sobre a pontuação: CMR 1A/1B → substituir se tecnicamente possível, sistema fechado, lista de trabalhadores expostos (conservar 40 anos; 5 anos para reprotóxicos) e "
        "vigilância da saúde; sensibilizantes (H317, H334, EUH204/208) → vigilância da saúde e luvas do material da secção 8; Protecao_Maternidade = 'Sim' → avaliar a tarefa para trabalhadoras "
        "grávidas, puérperas e lactantes e para menores (Lei 102/2009; CT art. 62.º).",
        "Medidas pela hierarquia: eliminação → substituição → engenharia (captação, confinamento) → organizacionais → EPI (luvas EN ISO 374 do material indicado, filtros A2/P3).",
        "Medir a exposição quando existe VLE e a prioridade é P1/P2 por inalação: estratégia NP EN 689 com laboratório acreditado; teste preliminar com 3 amostras < 10% do VLE (4 < 15%, 5 < 20%) "
        "→ conforme; senão medição periódica; > VLE → ação imediata.",
        "Rever a avaliação em cada alteração (produto, processo, ventilação), após medições desfavoráveis ou resultados da vigilância da saúde, e no máximo a cada 36 meses.",
    ])

    sub(doc, "6.10 Armazenagem e compatibilidades (folhas Armazenagem e Matriz_Compatibilidade)")
    bullets(doc, [
        "Cada produto tem um local ARM-nn; a folha calcula o stock e a retenção necessária = máx(100% do maior recipiente; 50% do total) — critério interno; confirmar condições da licença.",
        "Segregar segundo a matriz e as secções 7 e 10 da FDS: inflamáveis × comburentes; ácidos × bases/hipoclorito (EUH031 liberta cloro); oxigénio × gases inflamáveis; aerossóis longe do calor.",
        "Inflamáveis > 50 L fora de armário EN 14470-1 = local de risco C (RT-SCIE); zonas ATEX identificadas no documento de proteção contra explosões (DL 236/2003).",
        "Recipientes de trasfega sempre rotulados (nome e pictogramas); kit de derrame em cada local; inspeção na ronda ambiental mensal (RG-SGA-10).",
    ])

    sub(doc, "6.11 Seveso (folha Seveso)")
    para(doc, "Verificação anual e em cada alteração: a folha soma o Stock_Max_kg por categoria, divide pelos limiares do DL 150/2015 e aplica a regra da soma por grupo "
              "(saúde, físico, ambiente). Σ ≥ 1 no limiar inferior obriga a notificação à APA — comunicar de imediato ao Diretor Geral.")

    sub(doc, "6.12 Receção, rotulagem e informação no posto")
    bullets(doc, [
        "Na receção confirmar que o rótulo corresponde à FDS em vigor (nome, pictogramas, frases H) e ao ID do inventário; recusar embalagens danificadas ou sem rótulo em português.",
        "Ficha-resumo no posto (perigos, EPI, primeiros socorros, derrame) gerada a partir das secções 2, 4, 6, 7 e 8 da FDS.",
    ])

    sub(doc, "6.13 Formação e informação (registo por colaborador no RG-SGA-08: FOR-01, FOR-03, FOR-12 a FOR-15)")
    bullets(doc, [
        "Todos os utilizadores: FDS, pictogramas CLP, EPI e resposta a derrames antes da primeira utilização.",
        "Diisocianatos (≥ 0,1%, ex.: endurecedor da tinta 2K): formação de nível adequado ANTES do uso, renovada a cada 5 anos, com certificado (anexo XVII, entrada 74).",
        "CMR: informação específica e registo; expedição de resíduos perigosos: ADR 1.3.",
    ])

    sub(doc, "6.14 Emergência e resíduos")
    para(doc, "Derrames e incêndios segundo PR-SGA-05, IT-SGA-01 e IT-SGA-02 (secções 4, 5 e 6 da FDS; CIAV 800 250 250). Resíduos de químicos classificados pelas propriedades HP "
              "com base na FDS e encaminhados segundo PR-SGA-04 (ex.: 14 06 03*, 15 01 10*, 13 01 10*).")

    sub(doc, "6.15 Vigilância regulamentar e calendário (RG-SGA-04 tbl_obrigacoes OBR-11 e OBR-19 a OBR-30; parâmetros no RG-SGA-21)")
    bullets(doc, [
        "Semestral: lista candidata (janeiro e junho), novas restrições e ATP do CLP; novas classes de perigo do Reg. 2023/707 (substâncias já no mercado até 01/11/2026; misturas até 01/05/2028); "
        "regras de rótulo do CLP revisto adiadas para 01/01/2028 (Reg. 2025/2439). Em 27/04/2026 a Comissão anunciou que não abrirá o texto do REACH, preferindo alterar os anexos — acompanhar.",
        "Anual: inventário, declarações, Seveso, relatório de microplásticos (31/05), medições; granulado: notificação da instalação e declaração de conformidade até 17/12/2027 (Reg. 2025/2365).",
        "As alterações relevantes entram na matriz Requisitos_Legais e no RG-SGA-04 (PR-SGA-03).",
    ])

    h(doc, "7. Indicadores (folha Painel)")
    para(doc, "Produtos CMR 1A/1B (meta 0); produtos com SVHC (meta 0); % de FDS em vigor conformes (meta ≥ 95%); cenários de exposição por implementar (meta 0); lacunas de registo (meta 0); "
              "declarações desatualizadas; tarefas P1; produtos perigosos sem avaliação; medições não conformes/inconclusivas; locais de armazenagem não conformes; formações em falta; "
              "obrigações em atraso; conclusão Seveso. Os indicadores são apresentados na revisão pela gestão (PR-SGA-11).")

    h(doc, "8. Informação documentada (evidência de controlo REACH)")
    table(doc, ["Evidência exigível", "Onde fica (RG-SGA-21)", "Conservação"], [
        ("Papel REACH e verificação de fornecedores extra-UE (RU, importador)", "Papel_REACH; Registo_Importacao; cartas de RU (anexas)", "10 anos (art. 36.º)"),
        ("Inventário de produtos químicos e componentes perigosos", "Inventario; Componentes", "10 anos após a última utilização"),
        ("FDS recebidas (todas as versões) e verificação", "FDS_Verificacao + arquivo PDF das FDS", "10 anos"),
        ("Verificação dos cenários de exposição e medidas implementadas", "Cenarios_Exposicao", "10 anos"),
        ("Declarações de fornecedores e triagem SVHC dos artigos; SCIP", "Declaracoes_SVHC; RG-SGA-20 tbl_familias_ppwr", "10 anos"),
        ("Avaliação de riscos químicos e medições", "Avaliacao_Risco; Medicoes_VLE; relatórios do laboratório", "Enquanto válida + 10 anos"),
        ("Lista de trabalhadores expostos a CMR / reprotóxicos", "Anexo à Avaliacao_Risco (RH / medicina do trabalho)", "40 anos (reprotóxicos: 5 anos)"),
        ("Armazenagem, compatibilidades e Seveso", "Armazenagem; Matriz_Compatibilidade; Seveso", "5 anos após substituição"),
        ("Formação (diisocianatos, CMR, ADR)", "RG-SGA-08 tbl_registo_formacao + certificados", "Duração do vínculo + 5 anos"),
        ("Relatório de microplásticos e notificações à ECHA", "RG-SGA-04 OBR-20 + comprovativos REACH-IT", "10 anos"),
        ("Matriz legal, ações e análise crítica", "Requisitos_Legais; Analise_Critica; ações no RG-SGA-06 (PAM)", "5 anos após substituição"),
    ], [6, 7, 4])

    h(doc, "Anexo A — Verificação rápida antes de uma inspeção (ASAE / IGAMAOT / ACT)")
    bullets(doc, [
        "Painel sem 'Atenção' nos indicadores de CMR, SVHC, FDS, cenários e lacunas de registo — ou com ação no PAM dentro do prazo.",
        "Para 5 produtos ao acaso: rótulo no armazém = FDS em vigor = linha do inventário = local ARM correto = ficha no posto.",
        "Para cada fornecedor extra-UE: carta de RU com a Plasticom na lista, ou prova de que o importador é outra entidade.",
        "Cenários de exposição recebidos há mais de 12 meses: todos 'Conforme' ou 'Implementado'.",
        "Comprovativo do relatório de microplásticos do último 31/05 e certificados de formação em diisocianatos válidos.",
    ])

    history(doc, [("00", "2026-09-24", "Emissão. Criado com o registo RG-SGA-21, após análise crítica do ficheiro de referência REACH 2024.xlsx e das fontes ECHA, EUR-Lex, DGAE, ACT e INRS.")])
    bd._save(doc, os.path.join(bd.OUT, FNAME))
    return FNAME


if __name__ == "__main__":
    import sys
    bd.FORCE = "--force" in sys.argv
    print(build())
