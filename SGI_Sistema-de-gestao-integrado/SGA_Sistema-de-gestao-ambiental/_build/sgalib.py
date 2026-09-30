"""Biblioteca comum para gerar os registos do SGA da Plasticom (ISO 14001:2026).

Convenções de todos os ficheiros:
- Cada registo é uma Tabela Excel (ListObject) com cabeçalho na linha 1, uma linha por registo,
  sem células mescladas, tipos consistentes por coluna -> leitura direta por pandas / Power BI.
- Colunas calculadas têm fundo cinzento e fórmula Excel (não valores fixos).
- Listas controladas ficam na folha "Listas" (nomes lst_*) e alimentam a validação de dados.
- Cada ficheiro tem "00_LEIA-ME" (controlo documental e guia) e "Dicionario_Dados".
"""
import re
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment

EMPRESA = "Plasticom — Unidade 1 (Marinha Grande, Portugal)"
NORMA = "NP EN ISO 14001:2026"
DATA_REF = dt.date(2026, 12, 31)         # data de referência do SGA: fecho do ano de 2026 (a revisão pela gestão foi em 23/09/2026)
AUTOR = "Gestor(a) do SGA / EHS"
APROVADOR = "Diretor Geral"

FONT = "Arial"
C_HEAD = "1F4E5F"      # cabeçalho de tabela (verde-petróleo escuro)
C_HEAD_CALC = "4F6F7A"  # cabeçalho de coluna calculada
C_CALC = "EDEDED"      # fundo de célula calculada
C_BAND = "DCEBEF"      # faixas de título
C_TITLE = "0F2F3A"

thin = Side(style="thin", color="B7C4C9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True, color=C_TITLE)
F_SUB = Font(name=FONT, size=10, italic=True, color="4A5A60")
F_LINK = Font(name=FONT, size=10, color="0563C1", underline="single")

FILL_HEAD = PatternFill("solid", fgColor=C_HEAD)
FILL_HEAD_CALC = PatternFill("solid", fgColor=C_HEAD_CALC)
FILL_CALC = PatternFill("solid", fgColor=C_CALC)
FILL_BAND = PatternFill("solid", fgColor=C_BAND)

WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

NUMFMT = {
    "date": "yyyy-mm-dd",
    "int": "0",
    "num": "#,##0.00",
    "num1": "#,##0.0",
    "num0": "#,##0",
    "num3": "#,##0.000",
    "pct": "0%",
    "pct1": "0.0%",
    "eur": "#,##0 €",
    "text": "@",
}

# Paletas para formatação condicional (fundo, fonte)
CF_COLORS = {
    "red": ("F4C7C3", "7F1D1D"),
    "orange": ("FCE4C4", "7C3E00"),
    "yellow": ("FFF2B3", "5C4B00"),
    "green": ("CDEBD3", "14532D"),
    "blue": ("CFE2F3", "0B3A66"),
    "gray": ("E7E7E7", "404040"),
    "purple": ("E4D7F5", "3B1F66"),
}


def col(name, w=14, t="text", f=None, dv=None, desc="", dom="", req=True, key="", note=None):
    """Especificação de coluna. f = modelo de fórmula com tokens @Col@ (célula da linha) e
    #Col# (intervalo absoluto da coluna na tabela)."""
    return dict(name=name, w=w, t=t, f=f, dv=dv, desc=desc, dom=dom, req=req, key=key, note=note)


class Book:
    def __init__(self, code, title, activities, clauses, purpose, links=None, version="01",
                 elaborated=AUTOR, approved=APROVADOR, date=DATA_REF, retention="5 anos após substituição"):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.code, self.title = code, title
        self.meta = dict(activities=activities, clauses=clauses, purpose=purpose, links=links or [],
                         version=version, elaborated=elaborated, approved=approved, date=date,
                         retention=retention)
        self.sheets_info = []      # (sheet, descrição)
        self.dictionary = []       # linhas do dicionário
        self.lists = {}            # nome -> valores
        self.tables = {}           # nome da tabela -> (ws, header_row, first_row, last_row, colmap)
        self.readme = self.wb.create_sheet("00_LEIA-ME")
        self.lists["DataReferencia"] = [DATA_REF]

    # ------------------------------------------------------------------ listas
    def add_list(self, name, values):
        self.lists[name] = list(values)

    def _write_lists(self):
        if not self.lists:
            return
        ws = self.wb.create_sheet("Listas")
        ws["A1"] = "LISTAS CONTROLADAS (domínios de validação de dados) — editar aqui para alterar opções"
        ws["A1"].font = F_BOLD
        for i, (name, vals) in enumerate(self.lists.items()):
            c = 1 + i
            letter = get_column_letter(c)
            h = ws.cell(row=2, column=c, value=name)
            h.font, h.fill, h.alignment = F_HEAD, FILL_HEAD, CENTER
            for j, v in enumerate(vals):
                cell = ws.cell(row=3 + j, column=c, value=v)
                cell.font = F_BASE
                cell.alignment = WRAP_TOP
            ws.column_dimensions[letter].width = max(16, min(48, max(len(str(v)) for v in vals + [name]) + 2))
            ref = f"Listas!${letter}$3:${letter}${2 + len(vals)}"
            self.wb.defined_names[f"lst_{name}"] = DefinedName(f"lst_{name}", attr_text=ref)
            if name == "DataReferencia":
                ws.cell(row=3, column=c).number_format = "yyyy-mm-dd"
                self.wb.defined_names["DataRef"] = DefinedName("DataRef", attr_text=f"Listas!${letter}$3")
        ws.freeze_panes = "A3"
        self.sheets_info.append(("Listas", "Domínios (listas de valores) usados na validação de dados das tabelas."))

    # ------------------------------------------------------------------ tabelas
    def table(self, sheet_name, table_name, cols, rows, description, title=None, subtitle=None,
              cf=None, row_height=None, freeze_col=1, extra_rows=0, tab_color=None, print_landscape=True):
        """Escreve uma tabela. Se title for dado, a tabela começa na linha 4 (título nas linhas 1-2);
        caso contrário começa na linha 1 (formato base de dados puro)."""
        ws = self.wb.create_sheet(sheet_name)
        if tab_color:
            ws.sheet_properties.tabColor = tab_color
        hr = 1
        if title:
            ws["A1"] = title
            ws["A1"].font = F_TITLE
            if subtitle:
                ws["A2"] = subtitle
                ws["A2"].font = F_SUB
            hr = 4
        colmap = {c["name"]: get_column_letter(i + 1) for i, c in enumerate(cols)}
        n = len(rows) + extra_rows
        first, last = hr + 1, hr + max(n, 1)

        def render(tpl, r):
            s = re.sub(r"@([^@]+)@", lambda m: f"{colmap[m.group(1)]}{r}", tpl)
            s = re.sub(r"#([^#]+)#", lambda m: f"${colmap[m.group(1)]}${first}:${colmap[m.group(1)]}${last}", s)
            return s

        for i, c in enumerate(cols):
            cell = ws.cell(row=hr, column=i + 1, value=c["name"])
            cell.font = F_HEAD
            cell.fill = FILL_HEAD_CALC if c["f"] else FILL_HEAD
            cell.alignment = CENTER
            cell.border = BORDER
            if c.get("note") or c["desc"]:
                cell.comment = Comment(((c.get("note") or c["desc"]) + ("\n[calculado]" if c["f"] else ""))[:600], "SGA")
            ws.column_dimensions[get_column_letter(i + 1)].width = c["w"]
        for k in range(max(n, 1)):
            r = first + k
            data = rows[k] if k < len(rows) else {}
            for i, c in enumerate(cols):
                cell = ws.cell(row=r, column=i + 1)
                if c["f"]:
                    cell.value = render(c["f"], r)
                    cell.fill = FILL_CALC
                else:
                    v = data.get(c["name"])
                    if isinstance(v, str) and v.startswith("="):
                        v = render(v, r)
                    cell.value = v
                cell.font = F_BASE
                cell.border = BORDER
                cell.alignment = WRAP_TOP
                if c["t"] in NUMFMT:
                    cell.number_format = NUMFMT[c["t"]]
            if row_height:
                ws.row_dimensions[r].height = row_height
        ref = f"A{hr}:{get_column_letter(len(cols))}{last}"
        tab = Table(displayName=table_name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(tab)
        ws.freeze_panes = f"{get_column_letter(freeze_col + 1)}{hr + 1}"
        # validação de dados
        for i, c in enumerate(cols):
            if c["dv"]:
                L = get_column_letter(i + 1)
                if c["dv"].startswith("="):
                    dv = DataValidation(type="list", formula1=c["dv"], allow_blank=True)
                else:
                    dv = DataValidation(type="list", formula1=f"=lst_{c['dv']}", allow_blank=True)
                    if c["dv"] not in self.lists:
                        raise KeyError(f"lista {c['dv']} não definida")
                dv.error = "Valor fora da lista controlada (ver folha Listas)."
                dv.errorTitle = "Valor inválido"
                dv.showErrorMessage = True
                ws.add_data_validation(dv)
                dv.add(f"{L}{first}:{L}{last + 200}")
        # formatação condicional: cf = [(colname, {texto: cor})] ou (colname, "expr", cor)
        for rule in (cf or []):
            cname, spec = rule[0], rule[1]
            L = colmap[cname]
            rng = f"{L}{first}:{L}{last + 200}"
            if isinstance(spec, dict):
                for txt, color in spec.items():
                    bg, fg = CF_COLORS[color]
                    ws.conditional_formatting.add(rng, FormulaRule(
                        formula=[f'ISNUMBER(SEARCH("{txt}",{L}{first}))'],
                        fill=PatternFill("solid", fgColor=bg), font=Font(name=FONT, color=fg, bold=True)))
            else:
                bg, fg = CF_COLORS[rule[2]]
                expr = spec.replace("@", f"{L}{first}")
                ws.conditional_formatting.add(rng, FormulaRule(
                    formula=[expr], fill=PatternFill("solid", fgColor=bg), font=Font(name=FONT, color=fg, bold=True)))
        ws.sheet_view.zoomScale = 90
        ws.page_setup.orientation = "landscape" if print_landscape else "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"{hr}:{hr}"
        self.tables[table_name] = dict(ws=ws, hr=hr, first=first, last=last, colmap=colmap, sheet=sheet_name)
        self.sheets_info.append((sheet_name, description))
        for c in cols:
            self.dictionary.append(dict(Folha=sheet_name, Tabela=table_name, Coluna=c["name"],
                                        Tipo={"text": "Texto", "int": "Inteiro", "num": "Decimal", "num0": "Decimal",
                                              "num1": "Decimal", "num3": "Decimal", "date": "Data (aaaa-mm-dd)",
                                              "pct": "Percentagem (fração 0-1)", "pct1": "Percentagem (fração 0-1)",
                                              "eur": "Moeda (EUR)"}.get(c["t"], c["t"]),
                                        Origem="Calculado (fórmula)" if c["f"] else "Entrada",
                                        Descricao=c["desc"], Dominio_Unidade=c["dom"] or (f"Lista lst_{c['dv']}" if c["dv"] and not c["dv"].startswith("=") else ""),
                                        Obrigatorio="Sim" if c["req"] else "Não", Chave=c["key"]))
        return ws

    def ref(self, table_name, colname, absolute=True, sheet=True):
        t = self.tables[table_name]
        L = t["colmap"][colname]
        s = f"'{t['sheet']}'!" if sheet else ""
        return f"{s}${L}${t['first']}:${L}${t['last']}"

    def sheet(self, name, description, tab_color=None):
        ws = self.wb.create_sheet(name)
        if tab_color:
            ws.sheet_properties.tabColor = tab_color
        self.sheets_info.append((name, description))
        ws.sheet_view.showGridLines = False
        return ws

    # ------------------------------------------------------------------ LEIA-ME e dicionário
    def _write_readme(self):
        ws = self.readme
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = C_HEAD
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 110
        ws["A1"] = f"{self.code} — {self.title}"
        ws["A1"].font = F_TITLE
        ws["A2"] = f"Sistema de Gestão Ambiental {NORMA} · {EMPRESA} · Registo do SGA (fábrica simulada)"
        ws["A2"].font = F_SUB
        r = 4

        def band(text):
            nonlocal r
            ws.cell(row=r, column=1, value=text).font = F_BOLD
            for c in (1, 2):
                ws.cell(row=r, column=c).fill = FILL_BAND
            r += 1

        def kv(k, v, link=False):
            nonlocal r
            a = ws.cell(row=r, column=1, value=k)
            a.font, a.alignment = F_BOLD, WRAP_TOP
            b = ws.cell(row=r, column=2, value=v)
            b.font, b.alignment = (F_LINK if link else F_BASE), WRAP_TOP
            if isinstance(v, (dt.date, dt.datetime)):
                b.number_format = "yyyy-mm-dd"
                b.alignment = Alignment(horizontal="left", vertical="top")
            r += 1

        m = self.meta
        band("CONTROLO DO DOCUMENTO (7.5 — informação documentada)")
        kv("Código do registo", self.code)
        kv("Título", self.title)
        kv("Versão / Revisão", m["version"])
        kv("Data de emissão", m["date"])
        kv("Elaborado por", m["elaborated"])
        kv("Aprovado por", m["approved"])
        kv("Estado", "Em vigor")
        kv("Retenção", m["retention"])
        kv("Localização", "Repositório controlado do SGA > Registos_SGA_Plasticom (versão eletrónica é a versão controlada; cópias impressas são não controladas)")
        r += 1
        band("FINALIDADE E ÂMBITO")
        kv("Finalidade", m["purpose"])
        kv("Atividades do curso", m["activities"])
        kv("Requisitos ISO 14001:2026", m["clauses"])
        r += 1
        band("ESTRUTURA DO FICHEIRO")
        for s, d in self.sheets_info:
            kv(s, d)
        kv("Dicionario_Dados", "Descrição de cada coluna de cada tabela: tipo, origem (entrada/calculado), domínio, obrigatoriedade e chaves.")
        r += 1
        band("CONVENÇÕES (prontas para análise de dados e machine learning)")
        kv("Tabelas", "Cada registo é uma Tabela Excel com nome próprio (tbl_*): uma linha = um registo, cabeçalho único, sem células mescladas. Ler com pandas (pd.read_excel(ficheiro, sheet_name=folha, header=<linha do cabeçalho - 1>)) ou Power BI (Obter Dados > Tabela).")
        kv("Cabeçalho cinzento", "Coluna CALCULADA por fórmula — não escrever por cima. As restantes colunas são de entrada.")
        kv("Validação", "Colunas com lista controlada só aceitam valores da folha Listas (evita categorias duplicadas por grafia diferente).")
        kv("Chaves", "IDs estáveis e únicos (PK) e referências a IDs de outros registos (FK) permitem juntar tabelas entre ficheiros — ver SGA-00_Indice_Modelo_Dados.xlsx.")
        kv("Datas e unidades", "Datas em formato data (aaaa-mm-dd); números como números; a unidade está no nome da coluna ou no dicionário. Data de referência do SGA: 2026-12-31 (fecho do ano).")
        kv("Dados simulados", "Fábrica simulada Plasticom. Dados operacionais derivados do dataset do projeto (datasets/silver); dados ambientais e de conformidade são simulação didática — confirmar aplicabilidade legal e números antes de uso real.")
        if m["links"]:
            r += 1
            band("LIGAÇÕES A OUTROS REGISTOS")
            for k, v in m["links"]:
                kv(k, v)

    def _write_dictionary(self):
        ws = self.wb.create_sheet("Dicionario_Dados")
        cols = ["Folha", "Tabela", "Coluna", "Tipo", "Origem", "Descricao", "Dominio_Unidade", "Obrigatorio", "Chave"]
        widths = [22, 24, 30, 16, 18, 70, 36, 11, 12]
        for i, (c, w) in enumerate(zip(cols, widths)):
            cell = ws.cell(row=1, column=i + 1, value=c)
            cell.font, cell.fill, cell.alignment, cell.border = F_HEAD, FILL_HEAD, CENTER, BORDER
            ws.column_dimensions[get_column_letter(i + 1)].width = w
        for j, d in enumerate(self.dictionary):
            for i, c in enumerate(cols):
                cell = ws.cell(row=2 + j, column=i + 1, value=d[c])
                cell.font, cell.alignment, cell.border = F_BASE, WRAP_TOP, BORDER
        tab = Table(displayName="tbl_dicionario", ref=f"A1:{get_column_letter(len(cols))}{1 + max(1, len(self.dictionary))}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(tab)
        ws.freeze_panes = "D2"

    def save(self, path):
        self._write_lists()
        self._write_readme()
        self._write_dictionary()
        # ordem: LEIA-ME primeiro, Listas e Dicionário no fim
        order = [self.readme] + [s for s in self.wb.worksheets if s.title not in ("00_LEIA-ME", "Listas", "Dicionario_Dados")]
        for n in ("Listas", "Dicionario_Dados"):
            if n in self.wb.sheetnames:
                order.append(self.wb[n])
        self.wb._sheets = order
        self.wb.active = 0
        for ws in self.wb.worksheets:
            ws.sheet_view.tabSelected = ws is self.readme
        self.wb.properties.creator = AUTOR
        self.wb.properties.title = f"{self.code} {self.title}"
        self.wb.properties.subject = f"SGA {NORMA} — Plasticom"
        self.wb.save(path)
        return path


# ---------------------------------------------------------------------- utilidades de layout
def form_block(ws, r, label, value, lw=1, vw=6, height=None, label_fill=FILL_BAND, bold_value=False):
    """Linha de formulário (rótulo | valor mesclado) para folhas de apresentação/impressão."""
    a = ws.cell(row=r, column=1, value=label)
    a.font, a.fill, a.alignment, a.border = F_BOLD, label_fill, WRAP_TOP, BORDER
    if lw > 1:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=lw)
    b = ws.cell(row=r, column=lw + 1, value=value)
    b.font = F_BOLD if bold_value else F_BASE
    b.alignment, b.border = WRAP_TOP, BORDER
    ws.merge_cells(start_row=r, start_column=lw + 1, end_row=r, end_column=lw + vw)
    for c in range(1, lw + vw + 1):
        ws.cell(row=r, column=c).border = BORDER
    if height:
        ws.row_dimensions[r].height = height
    return b


def header_row(ws, r, labels, widths=None, fill=FILL_HEAD):
    for i, t in enumerate(labels):
        c = ws.cell(row=r, column=i + 1, value=t)
        c.font, c.fill, c.alignment, c.border = F_HEAD, fill, CENTER, BORDER
        if widths:
            ws.column_dimensions[get_column_letter(i + 1)].width = widths[i]


def doc_header(ws, code, title, model, ncols, version="01", date=DATA_REF):
    ws["A1"] = "PLASTICOM"
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=C_HEAD)
    mid = max(2, ncols // 2 - 1)
    ws.cell(row=1, column=mid, value=title).font = Font(name=FONT, size=13, bold=True, color=C_TITLE)
    ws.cell(row=1, column=ncols, value=model).font = F_BOLD
    ws.cell(row=2, column=ncols, value=f"Rev. {version} · {date:%d/%m/%Y}").font = F_BASE
    ws.cell(row=2, column=mid, value=f"{EMPRESA} · {NORMA}").font = F_SUB
    for c in range(1, ncols + 1):
        ws.cell(row=3, column=c).border = Border(bottom=Side(style="medium", color=C_HEAD))
