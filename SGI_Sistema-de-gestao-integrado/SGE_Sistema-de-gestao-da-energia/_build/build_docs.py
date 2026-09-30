"""Gera os documentos controlados do SGE (Word): Manual MAN-SGE-01, Política POL-SGE-01, procedimentos PR-SGE-01 a 12 e instruções IT-SGE-01 a 03.
Mesmo formato dos documentos do SGA (build_docs.py do SGA): cabeçalho, tabela de controlo, secções numeradas, registos e histórico.
Uso: python build_docs.py            → cria só os documentos que ainda não existem (preserva edições manuais)
     python build_docs.py --force    → regenera todos
Cada documento cita as normas e os requisitos legais usados e remete para os registos (RG-SGE-nn) e tabelas (tbl_*) onde fica a evidência."""
import os
import re
import unicodedata
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from sgelib import EMPRESA, NORMA, AUTOR, APROVADOR

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(os.path.join(HERE, "..", "Documentos_SGE_Plasticom"))
TEAL = RGBColor(0x1F, 0x4E, 0x5F)
FORCE = False


def slug(text, n=60):
    t = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9]+", "_", t).strip("_")[:n]


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def new_doc(code, title, version, date, state, approver=APROVADOR, author=AUTOR):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
        hp = s.header.paragraphs[0]
        hp.text = f"PLASTICOM — Sistema de Gestão da Energia ({NORMA})    |    {code} rev. {version}"
        hp.runs[0].font.size, hp.runs[0].font.color.rgb = Pt(8), RGBColor(0x60, 0x60, 0x60)
        fp = s.footer.paragraphs[0]
        fp.text = "Documento controlado: a versão eletrónica em Documentos_SGE_Plasticom é a versão em vigor; cópias impressas não são controladas (PR-SGA-08)."
        fp.runs[0].font.size, fp.runs[0].font.color.rgb = Pt(7), RGBColor(0x80, 0x80, 0x80)
    p = doc.add_paragraph()
    r = p.add_run(title.upper())
    r.bold, r.font.size, r.font.color.rgb = True, Pt(15), TEAL
    t = doc.add_table(rows=2, cols=6)
    t.style = "Table Grid"
    for i, (hh, v) in enumerate(zip(["Código", "Versão", "Data de aprovação", "Elaborado por", "Aprovado por", "Estado"], [code, version, date, author, approver, state])):
        c = t.cell(0, i)
        c.text = hh
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].runs[0].font.size = Pt(8)
        shade(c, "DCEBEF")
        t.cell(1, i).text = str(v)
        t.cell(1, i).paragraphs[0].runs[0].font.size = Pt(8)
    doc.add_paragraph()
    return doc


def h(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(11.5), TEAL
    p.paragraph_format.space_before = Pt(8)


def sub(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(10.5)
    p.paragraph_format.space_before = Pt(4)


def para(doc, text):
    doc.add_paragraph(text).paragraph_format.space_after = Pt(4)


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def steps(doc, items):
    for i, s in enumerate(items, start=1):
        doc.add_paragraph(f"{i}. {s}").paragraph_format.left_indent = Cm(0.4)


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, hh in enumerate(header):
        c = t.cell(0, i)
        c.text = hh
        run = c.paragraphs[0].runs[0]
        run.bold, run.font.size, run.font.color.rgb = True, Pt(8.5), RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "1F4E5F")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            for pr in cells[i].paragraphs:
                for rr in pr.runs:
                    rr.font.size = Pt(8.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()


def _save(doc, path):
    if os.path.exists(path) and not FORCE:
        return
    try:
        doc.save(path)
    except PermissionError:
        print("aviso: documento aberto no Word, não atualizado:", os.path.basename(path))


# ------------------------------------------------------------------------------------------------ procedimentos
# code: dict(title, ver, date, obj, amb, refs[], legal[], defs[(termo, definição)], resp[(função, resp)], steps[] ou secções{título: [passos]}, crit (header, rows) opcional,
#            regs[(registo, tabelas)], hist[(v, data, alteração)])
BASE_REFS = ["NP EN ISO 50001:2019 (ISO 50001:2018 + Amd 1:2024 — Climate action changes)", "ISO 50004:2020 — orientação para implementar, manter e melhorar um SGE"]
PROC = {
    "PR-SGE-01": dict(
        title="Contexto, partes interessadas, âmbito e fronteiras do SGE", ver="00", date="2026-09-29",
        obj="Determinar as questões externas e internas, as partes interessadas e os seus requisitos (incluindo os legais e os relacionados com as alterações climáticas), e definir e manter o âmbito e as fronteiras do SGE.",
        amb="Unidade 1 da Plasticom (Marinha Grande), todas as atividades, tipos de energia e frota de serviço.",
        refs=BASE_REFS + ["ISO 50001 4.1, 4.2, 4.3 e 4.4", "ISO 31000:2018", "RG-SGA-01 (contexto do SGI — fonte única)"],
        legal=["DL 71/2008 (SGCIE)", "Diretiva (UE) 2023/1791, art. 11.º (transposição em curso)"],
        defs=[("Fronteira", "Limites físicos ou organizacionais definidos pela organização (ISO 50001 3.1.3)."),
              ("Âmbito", "Atividades, instalações e decisões abrangidas pelo SGE, que podem incluir várias fronteiras.")],
        resp=[("Diretor Geral", "Aprova o âmbito e as fronteiras."), ("Gestor(a) de Energia", "Mantém o RG-SGE-01 e liga cada questão ao contexto do SGI."),
              ("Gestor do SGA / EHS", "Mantém o contexto do SGI (RG-SGA-01) e o registo legal (RG-SGA-04).")],
        steps=["Ler o PESTEL, SWOT e TOWS do RG-SGA-01 e selecionar os fatores com efeito na energia; registar em tbl_contexto_energia com o ID_Contexto_SGI.",
               "Acrescentar as questões só de energia (prefixo CTX-E): preço e tarifas, tecnologia das máquinas, dados, legislação de eficiência.",
               "Determinar se as alterações climáticas são pertinentes (Amd 1:2024) e justificar em tbl_clima.",
               "Identificar as partes interessadas, os requisitos e se o SGE os trata (tbl_partes_interessadas); encaminhar os legais para o RG-SGE-10.",
               "Definir o âmbito e as fronteiras (tbl_ambito, tbl_fronteiras), confirmando a autoridade para controlar e que nenhum tipo de energia é excluído.",
               "Rever anualmente antes da revisão pela gestão e sempre que haja alteração relevante (nova linha, novo tipo de energia, legislação)."],
        regs=[("RG-SGE-01", "tbl_contexto_energia, tbl_partes_interessadas, tbl_clima, tbl_ambito, tbl_fronteiras")],
        hist=[("00", "2026-09-29", "Emissão inicial do SGE.")]),
    "PR-SGE-02": dict(
        title="Revisão energética e identificação dos usos significativos de energia", ver="00", date="2026-09-29",
        obj="Estabelecer os métodos e critérios da revisão energética (6.3), que são mantidos como informação documentada, e a forma de reter os seus resultados.",
        amb="Todos os tipos e usos de energia dentro das fronteiras do SGE.",
        refs=BASE_REFS + ["ISO 50001 6.3 a)–e)", "ISO 50002-1:2025 (auditorias energéticas; substitui a ISO 50002:2014) e EN 16247-1:2022", "ISO 50006:2023 §5.1 e §5.4",
                          "ISO 50046:2019 (previsão de poupanças)", "ISO 11011:2013 (ar comprimido)", "Kent, R. — Energy Management in Plastics Processing (BPF)"],
        legal=["DL 71/2008 (SGCIE) — a auditoria energética obrigatória (8 anos) usa esta revisão", "Despacho 17313/2008 — fatores de conversão para tep"],
        defs=[("Uso significativo de energia (USE)", "Uso com consumo substancial e/ou potencial considerável de melhoria (ISO 50001 3.5.6); critérios definidos pela organização."),
              ("Uso de energia", "Aplicação da energia (ventilação, iluminação, processos...) — 'para que é usada'."), ("Consumo de energia", "Quantidade de energia aplicada — 'quanto'.")],
        resp=[("Gestor(a) de Energia", "Coordena a revisão, aplica os critérios e propõe os USE."), ("Técnico de Utilidades", "Campanhas de medição e dados do ar comprimido e do frio."),
              ("Diretor Industrial", "Aprova os USE e a priorização das oportunidades.")],
        steps=["Identificar os tipos de energia atuais e futuros (tbl_tipos_energia) — 6.3 a 1.",
               "Recolher o consumo mensal por uso (tbl_consumo_uso): fatura e analisadores (medido) ou rateio horas × kW (estimado); gasóleo pelo cartão de frota — 6.3 a 2.",
               "Calcular o consumo de 12 meses, o peso de cada uso e o Pareto (tbl_usos).",
               "Aplicar os critérios de USE: uso ≥ 10% do consumo total OU potencial de melhoria identificado ≥ 100 MWh/ano (tbl_parametros) — 6.3 b.",
               "Para cada USE preencher a ficha: variáveis relevantes, desempenho atual (IDE), pessoas que o influenciam e controlos (tbl_use) — 6.3 c.",
               "Medir por máquina com o analisador portátil PA-01 (48 h: kW em produção e em espera) e ratear o consumo por máquina (tbl_campanha, tbl_energia_maquina).",
               "Registar as oportunidades com base, % de poupança, investimento e referência técnica; o registo calcula MWh, €, retorno, CO2 e prioridade (tbl_oportunidades) — 6.3 d.",
               "Estimar o consumo futuro com o crescimento da produção e as ações aprovadas (tbl_previsao) — 6.3 e.",
               "Atualizar anualmente (antes da revisão pela gestão) e após alterações maiores; a revisão de 2027 usa 12 meses de submedição."],
        crit=(["Critério", "Regra"], [("Significância (USE)", "≥ 10% do consumo total OU potencial ≥ 100 MWh/ano"),
                                      ("Prioridade das oportunidades", "Pontos = min(3; MWh/50) + retorno (≤1 ano 3; ≤3 anos 2; >3 anos 1) + facilidade (1–3); Alta ≥ 7; Média ≥ 5"),
                                      ("Qualidade dos dados", "Medido (fatura, analisador, caudalímetro) vs estimado (rateio) — EN 17267")]),
        regs=[("RG-SGE-04", "tbl_tipos_energia, tbl_consumo_uso, tbl_usos, tbl_use, tbl_campanha, tbl_energia_maquina, tbl_oportunidades, tbl_previsao, Metodo_Criterios")],
        hist=[("00", "2026-09-29", "Emissão inicial.")]),
    "PR-SGE-04": dict(
        title="Riscos e oportunidades, objetivos, metas energéticas e planos de ação", ver="00", date="2026-09-29",
        obj="Determinar os riscos e oportunidades do SGE e estabelecer objetivos, metas energéticas e planos de ação com o método de verificação da melhoria.",
        amb="SGE da Unidade 1.",
        refs=BASE_REFS + ["ISO 50001 6.1, 6.2", "ISO 31000:2018", "ISO 50015:2014 e IPMVP (método de verificação)"],
        legal=["DL 71/2008 (SGCIE) — metas do PREn consideradas nos objetivos (6.2.2 c)"],
        defs=[("Meta energética", "Objetivo quantificado de melhoria do desempenho energético (ISO 50001 3.4.15)."),
              ("Plano de ação", "O quê, recursos, responsável, prazo e como os resultados são avaliados, incluindo o método de verificação (6.2.3).")],
        resp=[("Gestor(a) de Energia", "Propõe riscos, objetivos, metas e planos."), ("Donos dos riscos / planos", "Executam e avaliam a eficácia."), ("Diretor Geral", "Aprova objetivos, metas e recursos.")],
        steps=["Identificar riscos e oportunidades a partir do contexto, partes interessadas, revisão energética e desvios; usar os IDs do RG-SGA-02 quando já existem (fonte única).",
               "Avaliar por probabilidade × impacto (1–5) e decidir o tratamento; definir a integração no processo e o método de avaliação da eficácia (tbl_riscos_e, tbl_oport_e).",
               "Estabelecer objetivos coerentes com a política, mensuráveis, que considerem os USE, as oportunidades e os requisitos aplicáveis (6.2.2 a–h) — tbl_objetivos.",
               "Estabelecer metas energéticas por IDE com valor e prazo (tbl_metas_energeticas).",
               "Para cada objetivo, criar planos de ação com o quê, recursos, responsável, prazo, como avaliar e o método de verificação IPMVP/ISO 50015 (tbl_planos_acao).",
               "Acompanhar mensalmente na reunião da equipa de gestão de energia e rever na revisão pela gestão."],
        regs=[("RG-SGE-03", "tbl_riscos_e, tbl_oport_e"), ("RG-SGE-05", "tbl_objetivos, tbl_metas_energeticas, tbl_planos_acao")],
        hist=[("00", "2026-09-29", "Emissão inicial.")]),
    "PR-SGE-05": dict(
        title="Plano de recolha de dados energéticos, medição e confirmação dos equipamentos", ver="01", date="2026-12-15",
        obj="Assegurar que as características-chave que afetam o desempenho energético são identificadas, medidas, monitorizadas e analisadas com dados exatos e repetíveis.",
        amb="Todos os contadores, analisadores, campanhas e fontes de dados de energia e das variáveis relevantes.",
        refs=BASE_REFS + ["ISO 50001 6.6 e 9.1.1", "EN 17267:2019 (plano de medição e monitorização)", "ISO 50006:2023 §5.6", "ISO 50015:2014", "IEC 62053-22 / IEC 61557-12"],
        legal=["Diretiva 2014/32/UE (MID) e controlo metrológico legal do contador de faturação"],
        defs=[("Característica-chave", "Parâmetro de operação que afeta o desempenho energético (9.1.1 a: eficácia dos planos, IDE, operação dos USE, real vs esperado).")],
        resp=[("Gestor(a) de Energia", "Mantém o plano de recolha e o plano de monitorização."), ("Gerente de Dados / TI", "Integração dos analisadores no BI e retenção dos dados."),
              ("Técnico de Utilidades", "Leituras, reconciliação e campanhas.")],
        steps=["Manter a árvore de contadores (nível 0 fronteira, 1 USE, 2 equipamento) e a fração do consumo que cada ponto mede (tbl_contadores; KPI-E-01).",
               "Definir os dados a recolher nas 5 categorias de 6.6 a)–e), com fonte, método, frequência de recolha e de análise, retenção e qualidade (tbl_plano_recolha).",
               "Definir o plano de monitorização 9.1.1 com o valor esperado e o limite de desvio significativo de cada item (tbl_plano_monitorizacao).",
               "Confirmar os equipamentos (calibração/verificação, classe, incerteza, repetibilidade) — o estado é calculado em tbl_equipamentos.",
               "Reconciliar todos os meses a leitura interna com a fatura (±1%) e, desde 15/12/2026, Σ analisadores com o contador geral (±2%) — tbl_reconciliacao.",
               "Rever o plano anualmente e quando mudarem os USE ou os contadores."],
        regs=[("RG-SGE-06", "tbl_contadores, tbl_plano_recolha, tbl_plano_monitorizacao, tbl_equipamentos, tbl_reconciliacao")],
        hist=[("00", "2026-09-29", "Emissão inicial."), ("01", "2026-12-15", "Analisadores M01–M06 em serviço; reconciliação Σ analisadores × M00.")]),
    "PR-SGE-06": dict(
        title="Controlo operacional dos usos significativos de energia", ver="01", date="2026-12-01",
        obj="Definir, comunicar e controlar os critérios de operação e manutenção dos USE para evitar desvios significativos do desempenho energético.",
        amb="USE-01 sopro, USE-02 injeção, USE-03 ar comprimido, USE-04 arrefecimento e serviços gerais.",
        refs=BASE_REFS + ["ISO 50001 8.1", "ISO 11011:2013", "Kent, R. / BPF (standby, pressão, setpoint da água gelada)"],
        legal=["Reg. (UE) 2024/573 (gases fluorados) — manutenção do chiller (tratada pelo SGA)"],
        defs=[("Desvio significativo", "Afastamento do critério definido pela organização que pode alterar o desempenho energético (Nota de 8.1).")],
        resp=[("Chefes de turno", "Standby, ramais de ar e iluminação."), ("Técnico de Utilidades", "Pressão, fugas, setpoints e rondas."), ("Gerente de Manutenção", "Critérios de manutenção.")],
        steps=["Definir por USE o parâmetro, o critério, o desvio significativo, o método, a frequência e o responsável (tbl_criterios_operacionais).",
               "Comunicar os critérios nas IT-SGE-01 a 03, no quadro SQDC e na formação FOR-E-01.",
               "Fazer a ronda de energia semanal (7 pontos) e registar o resultado e a ação imediata (tbl_rondas_energia).",
               "Tratar as não conformidades repetidas ou com efeito relevante como NC (PR-SGE-12); taxa de conformidade < 80% obriga a rever o critério.",
               "Controlar as alterações planeadas pela checklist do SGI (PR-SGA-06, secção de energia) e rever as consequências das não intencionais (folha Alteracoes_Nao_Intencionais).",
               "Garantir que os serviços subcontratados com impacto nos USE têm critérios energéticos (PR-SGE-07)."],
        regs=[("RG-SGE-08", "tbl_criterios_operacionais, tbl_rondas_energia, Alteracoes_Nao_Intencionais")],
        hist=[("00", "2026-09-29", "Emissão inicial."), ("01", "2026-12-01", "Pressão de 6,8 bar e setpoint do chiller ≥ 10 °C.")]),
    "PR-SGE-07": dict(
        title="Projeto, aquisições e compra de energia com critérios de desempenho energético", ver="01", date="2026-12-18",
        obj="Considerar o desempenho energético no projeto de instalações, equipamentos e processos, e nas aquisições de produtos, equipamentos, serviços e energia.",
        amb="Projetos novos, modificados ou renovados com impacto significativo; aquisições com impacto nos USE; contrato de eletricidade.",
        refs=BASE_REFS + ["ISO 50001 8.2 e 8.3", "EN 17463:2021 (VALERI — avaliação de investimentos)", "EUROMAP 60.1/60.2", "ISO 1217 (anexo E)", "ISO 55001:2024 (gestão de ativos)"],
        legal=["Reg. (UE) 2019/1781 (motores e variadores)", "Reg. (UE) 2016/2281 (chillers — SEPR)", "DL 15/2022 alterado pelo DL 130/2026 (autoconsumo)", "Regulamentos da ERSE (tarifas)"],
        defs=[("Custo do ciclo de vida (LCC)", "Investimento + valor atual da energia (com escalada) + valor atual da manutenção durante a vida útil.")],
        resp=[("Diretor Industrial", "Aprova os projetos com a avaliação energética."), ("Responsável de Compras", "Aplica os critérios e informa os fornecedores."),
              ("Gestor(a) de Energia", "Avalia o desempenho energético e o LCC."), ("Diretor Financeiro", "Compra de energia e tarifas.")],
        steps=["Na checklist de alterações do SGI, responder à secção de energia: impacto significativo? oportunidades? controlo operacional? (tbl_projetos).",
               "Para impacto significativo, avaliar o desempenho energético ao longo da vida e incorporar o resultado na especificação (8.2).",
               "Aplicar os critérios por categoria (tbl_criterios_aquisicao) e informar os fornecedores de que o desempenho energético é critério de avaliação (8.3).",
               "Calcular o LCC (tbl_lcc) em investimentos > € 50 000 e escolher pelo menor custo anual equivalente.",
               "Especificar a compra de energia (origem, preço, dados de 15 min, potência contratada, autoconsumo) — tbl_compra_energia.",
               "Verificar mensalmente as faturas: preço médio, quota de ponta e utilização da potência (alerta > 95%) — tbl_faturas."],
        regs=[("RG-SGE-09", "tbl_projetos, tbl_criterios_aquisicao, tbl_aquisicoes, tbl_lcc, tbl_compra_energia, tbl_precos, tbl_faturas")],
        hist=[("00", "2026-09-29", "Emissão inicial."), ("01", "2026-12-18", "Secção de energia obrigatória na checklist e LCC > € 50 000 (NCE-26-01; RD-E-26-D05); categoria CAQ-06 para serviços (NCE-26-02).")]),
    "PR-SGE-08": dict(
        title="Desvios significativos e medição e verificação de poupanças", ver="00", date="2026-10-01",
        obj="Investigar e responder aos desvios significativos do desempenho energético e verificar as poupanças das ações de forma reprodutível.",
        amb="IDE da instalação e dos USE; todas as ações de melhoria com poupança declarada.",
        refs=BASE_REFS + ["ISO 50001 9.1.1, 6.2.3, 10.2", "ISO 50015:2014 (M&V)", "ISO 50047:2016 (poupanças)", "EVO — IPMVP Core Concepts", "ASHRAE Guideline 14", "ISO 50006:2023 §10"],
        legal=["Regulamentos de avisos de financiamento (M&V das poupanças financiadas)"],
        defs=[("Opção A / B / C (IPMVP)", "A: medir o parâmetro-chave e estimar os restantes; B: medir todos os parâmetros na fronteira da medida; C: instalação completa com modelo."),
              ("Ajuste não rotineiro", "Correção da LBE por alteração de fatores estáticos (ex.: novas máquinas).")],
        resp=[("Gestor(a) de Energia", "Deteta e coordena a investigação; aprova os planos de M&V."), ("Donos das ações", "Recolhem os dados de M&V.")],
        steps=["Todos os meses comparar o real com o esperado (RG-SGE-05); é desvio significativo |real − esperado| > t(95%) × erro-padrão, um mês fora do domínio, ou um critério operacional falhado.",
               "Registar o desvio (tbl_desvios), investigar a causa e responder em ≤ 30 dias (KPI-E-02); desvios favoráveis também são investigados para confirmar as causas.",
               "Para cada ação com poupança, elaborar o plano de M&V (fronteira, opção IPMVP, períodos, ajustes, instrumentos, incerteza) — tbl_mv_planos.",
               "Calcular as poupanças verificadas (tbl_poupancas) e a sua incerteza; comparar a soma com a melhoria normalizada da instalação (coerência opção C vs A/B).",
               "Reportar na reunião mensal e na revisão pela gestão; poupança verificada < 50% da prevista abre ação (PR-SGE-12)."],
        regs=[("RG-SGE-11", "tbl_desvios, tbl_mv_planos, tbl_ar_mv, tbl_poupancas"), ("RG-SGE-05", "tbl_ide_mensal, LBE_Modelo")],
        hist=[("00", "2026-10-01", "Emissão inicial.")]),
    "PR-SGE-09": dict(
        title="Competência, consciencialização e comunicação do SGE", ver="00", date="2026-09-29",
        obj="Assegurar a competência das pessoas que afetam o desempenho energético, a consciencialização de todos, a comunicação e o processo de sugestões.",
        amb="Todas as pessoas que trabalham para a Plasticom ou em seu nome.",
        refs=BASE_REFS + ["ISO 50001 7.2, 7.3, 7.4", "ISO 10015 (gestão da competência)", "ISO 50003:2021 (competência de auditores)"],
        legal=["Código do Trabalho, arts. 130.º–134.º (formação contínua)"],
        defs=[],
        resp=[("Responsável de RH", "Plano e registos de formação."), ("Gestor(a) de Energia", "Requisitos de competência, consciencialização, comunicação e sugestões.")],
        steps=["Definir por função os requisitos de competência ligados aos USE (tbl_competencias).",
               "Planear e realizar a formação (tbl_formacoes); registar e avaliar a eficácia até ao nível 3 de Kirkpatrick (tbl_registo_formacao).",
               "Consciencializar para a política, a contribuição de cada um, o impacto do seu comportamento e as consequências de não cumprir (7.3 a–d) — tbl_consciencializacao.",
               "Comunicar segundo a matriz o quê/quando/a quem/como/quem, com informação coerente com os registos do SGE (tbl_comunicacao).",
               "Receber sugestões por caixa, app ou e-mail; responder ao proponente em ≤ 15 dias e ligar a oportunidades ou controlos (tbl_sugestoes)."],
        regs=[("RG-SGE-07", "tbl_competencias, tbl_formacoes, tbl_registo_formacao, tbl_consciencializacao, tbl_comunicacao, tbl_sugestoes")],
        hist=[("00", "2026-09-29", "Emissão inicial.")]),
    "PR-SGE-10": dict(
        title="Requisitos legais e outros requisitos de energia e avaliação da conformidade (SGCIE)", ver="00", date="2026-09-29",
        obj="Identificar os requisitos legais e outros requisitos de energia, determinar como se aplicam e avaliar a conformidade em intervalos planeados.",
        amb="Legislação nacional e europeia de energia e compromissos subscritos (clientes, certificação).",
        refs=BASE_REFS + ["ISO 50001 4.2, 9.1.2", "ISO 50003:2021"],
        legal=["DL 71/2008 (SGCIE), alterado pela Lei 7/2013 e pelo DL 68-A/2015", "Despacho 17313/2008", "DL 68-A/2015", "Diretiva (UE) 2023/1791 (arts. 8.º e 11.º)",
               "DL 15/2022 alterado pelo DL 130/2026", "Reg. (UE) 2019/1781", "Reg. (UE) 2016/2281", "Reg. (UE) 2024/573", "Diretiva (UE) 2022/2464 (CSRD) — via clientes"],
        defs=[("CIE", "Consumidor intensivo de energia: instalação com mais de 500 tep/ano (SGCIE)."),
              ("PREn / ARCE", "Plano de Racionalização do Consumo de Energia; Acordo de Racionalização dos Consumos de Energia (metas a 8 anos).")],
        resp=[("Gestor(a) de Energia", "Vigilância legal, SGCIE, REP."), ("Gestor do SGA / EHS", "Inclusão dos requisitos no RG-SGA-04 (fonte única)."), ("Diretor Financeiro", "Critério de grande empresa; UPAC.")],
        steps=["Vigiar semestralmente o Diário da República, a DGEG, a ADENE e o EUR-Lex (OBRE-04).",
               "Registar cada requisito com a mesma estrutura do RG-SGA-04 (tbl_legal) e propor a sua inclusão no registo do SGI.",
               "Registar as obrigações com prazo (auditoria SGCIE, PREn, REP, controlo prévio da UPAC) — tbl_obrigacoes.",
               "Calcular todos os anos os indicadores do PREn (tep, consumo específico, intensidade energética e carbónica) e comparar com a trajetória (−6% em 8 anos) — tbl_sgcie.",
               "Avaliar a conformidade de cada requisito com evidência datada; o estado e a próxima avaliação são calculados.",
               "Tratar qualquer incumprimento como NC (PR-SGE-12) e reportar na revisão pela gestão."],
        crit=(["Indicador do PREn", "Cálculo e meta (≥ 1 000 tep)"], [("Consumo específico (CEE)", "kgep / t de produto — −6% em 8 anos"), ("Intensidade energética (IE)", "kgep / mil € de VAB — −6% em 8 anos"),
                                                                     ("Intensidade carbónica (IC)", "kgCO2e / kgep — manter"), ("Fatores", "Eletricidade 0,215 kgep/kWh e 0,47 kgCO2e/kWh; gasóleo 0,864 kgep/L (Despacho 17313/2008)")]),
        regs=[("RG-SGE-10", "tbl_legal, tbl_obrigacoes, tbl_sgcie"), ("RG-SGA-04", "tbl_legal (fonte única do SGI)")],
        hist=[("00", "2026-09-29", "Emissão inicial.")]),
    "PR-SGE-11": dict(
        title="Auditoria interna e revisão pela gestão do SGE", ver="00", date="2026-09-29",
        obj="Planear e realizar auditorias internas que verifiquem a conformidade, a eficácia e a melhoria do desempenho energético, e conduzir a revisão pela gestão.",
        amb="Todo o SGE; auditorias e revisões integradas com o SGA e o SGQ sempre que possível.",
        refs=BASE_REFS + ["ISO 50001 9.2 e 9.3", "ISO 19011 (edição em vigor)", "ISO 50003:2021"],
        legal=["—"],
        defs=[],
        resp=[("Gestor(a) de Energia", "Programa de auditorias e preparação da revisão."), ("Auditores do SGE", "Auditam com imparcialidade."), ("Diretor Geral", "Conduz a revisão e decide.")],
        steps=["Elaborar o programa anual com base na importância dos processos (USE) e em auditorias anteriores (tbl_programa_auditorias).",
               "Qualificar os auditores (ISO 19011 + competência em desempenho energético — ISO 50003) e garantir que não auditam o próprio trabalho (tbl_auditores).",
               "Auditar com checklist por requisito, verificando sempre se o SGE melhora o desempenho energético (9.2.1 a) — tbl_checklist_aud.",
               "Registar as constatações e abrir as NC no RG-SGE-14 (tbl_constatacoes).",
               "Preparar a revisão pela gestão com as entradas 9.3.2 a)–e) e 9.3.3 (objetivos, desempenho e IDE, planos) — tbl_agenda.",
               "Registar as decisões por 9.3.4 a)–g), com responsável, prazo e recursos, e a ata (tbl_decisoes)."],
        regs=[("RG-SGE-12", "tbl_auditores, tbl_programa_auditorias, tbl_checklist_aud, tbl_constatacoes"), ("RG-SGE-13", "tbl_agenda, tbl_decisoes, tbl_acoes_anteriores, Ata")],
        hist=[("00", "2026-09-29", "Emissão inicial.")]),
    "PR-SGE-12": dict(
        title="Não conformidade, ação corretiva e melhoria contínua do SGE", ver="00", date="2026-09-29",
        obj="Tratar as não conformidades e demonstrar a melhoria contínua do SGE e do desempenho energético.",
        amb="NC de auditorias, rondas, desvios significativos, requisitos legais e reclamações de partes interessadas.",
        refs=BASE_REFS + ["ISO 50001 10.1, 10.2", "ISO 50006:2023 §10.3"],
        legal=["—"],
        defs=[],
        resp=[("Gestor(a) de Energia", "Registo e acompanhamento."), ("Responsáveis das ações", "Causas, ações e eficácia.")],
        steps=["Reagir: controlar, corrigir e tratar as consequências (10.1 a).",
               "Analisar a causa (5 Porquês ou Ishikawa) e verificar se existem ou podem ocorrer NC semelhantes (10.1 b) — tbl_5porques.",
               "Implementar a ação corretiva proporcional aos efeitos (10.1 c) e, se necessário, alterar o SGE (10.1 e).",
               "Rever a eficácia com critério objetivo e data (10.1 d); não eficaz → nova análise.",
               "Demonstrar a melhoria contínua: IDE-01 por trimestre e por ano, IDE dos USE, poupanças verificadas e maturidade do sistema (tbl_melhoria; ISO 50005)."],
        regs=[("RG-SGE-14", "tbl_nc, tbl_5porques, tbl_melhoria")],
        hist=[("00", "2026-09-29", "Emissão inicial.")]),
}

IT = {
    "IT-SGE-01": ("Standby das máquinas nas pausas e paragens", "01", "2026-10-05", "Gerente de Produção",
                  ["Paragem prevista > 30 min (pausa do 3.º turno, falta de molde, avaria longa): ativar o modo STANDBY no HMI.",
                   "Injetoras hidráulicas: resistências do canhão a 60% e bomba hidráulica desligada. Elétricas: modo eco.",
                   "ISBM: fornos de pré-forma a 60% e bomba desligada; confirmar a recuperação de ar ativa nas elétricas.",
                   "Desligar a extração e a iluminação da ilha se não houver outra máquina em marcha.",
                   "No regresso: sair do standby 15 min antes do arranque (tempo de aquecimento validado).",
                   "O chefe de turno assina o checklist de pausa; a ronda semanal verifica (CO-01, CO-05, CO-16)."],
                  "RG-SGE-08 tbl_rondas_energia; checklist de pausa; RG-SGE-11 MV-02"),
    "IT-SGE-02": ("Ar comprimido: pressão, fugas e ramais", "01", "2026-11-16", "Gerente de Manutenção",
                  ["Setpoint dos compressores: 6,8 bar (não alterar sem autorização do Gestor de Energia).",
                   "Ronda de fugas mensal com o detetor ultrassónico EQP-05; etiquetar cada fuga com data.",
                   "Reparar as fugas etiquetadas em ≤ 5 dias úteis; registar no SAP PM.",
                   "Teste de vazio ao domingo (1.º do mês): ler o caudal no EQP-04 com a produção parada; fugas = caudal ÷ caudal médio de produção.",
                   "Sexta às 22 h: fechar as válvulas dos ramais das áreas sem produção ao fim de semana.",
                   "Não usar ar comprimido para limpar roupa ou bancadas (usar aspirador).",
                   "Verificar a perda de carga dos filtros (≤ 0,4 bar) semanalmente."],
                  "RG-SGE-08 (CO-09 a CO-12); RG-SGE-11 tbl_ar_mv"),
    "IT-SGE-03": ("Operação eficiente do chiller e da torre de arrefecimento", "01", "2026-12-01", "Gerente de Manutenção",
                  ["Setpoint da água gelada: 10 °C (validado por DOE com a qualidade); não baixar sem autorização da engenharia de processo.",
                   "Qualquer alteração de setpoint gera alarme no BI e deve ser justificada no registo do chiller.",
                   "Verificar o ΔT ida/retorno (4–6 °C) semanalmente; ΔT < 3 °C → verificar caudal e bombas.",
                   "Limpar os permutadores da torre antes do verão e sempre que o ΔT cair > 1 °C.",
                   "Com temperatura exterior < 9 °C, colocar o free-cooling em serviço (após instalação — PA-E-06).",
                   "Fugas de fluido frigorigéneo: seguir o RG-SGA-10 (técnico certificado)."],
                  "RG-SGE-08 (CO-13 a CO-15); IDE-06"),
}


def build_proc(code, s):
    doc = new_doc(code, s["title"], s["ver"], s["date"], "Em vigor")
    h(doc, "1. Objetivo"); para(doc, s["obj"])
    h(doc, "2. Âmbito"); para(doc, s["amb"])
    h(doc, "3. Referências normativas e requisitos"); bullets(doc, s["refs"])
    sub(doc, "Requisitos legais e outros requisitos aplicáveis"); bullets(doc, s["legal"])
    n = 4
    if s.get("defs"):
        h(doc, f"{n}. Definições"); table(doc, ["Termo", "Definição"], s["defs"], [4.5, 12.5]); n += 1
    h(doc, f"{n}. Responsabilidades"); table(doc, ["Função", "Responsabilidade"], s["resp"], [5, 12]); n += 1
    h(doc, f"{n}. Descrição do processo"); steps(doc, s["steps"]); n += 1
    if s.get("crit"):
        h(doc, f"{n}. Critérios"); table(doc, s["crit"][0], s["crit"][1], [5, 12]); n += 1
    h(doc, f"{n}. Informação documentada (evidência)"); table(doc, ["Registo", "Tabelas (Registos_SGE_Plasticom)"], s["regs"], [4, 13]); n += 1
    h(doc, "Histórico de revisões"); table(doc, ["Versão", "Data", "Alteração"], s["hist"], [2, 3, 12])
    _save(doc, os.path.join(OUT, f"{code}_{slug(s['title'])}.docx"))


def build_pr03():
    """PR-SGE-03 — IDE, LBE e normalização (ISO 50006:2023): procedimento-guia detalhado."""
    code, title = "PR-SGE-03", "Indicadores de desempenho energético, linhas de base e normalização (ISO 50006)"
    doc = new_doc(code, title, "01", "2026-12-18", "Em vigor")
    h(doc, "1. Objetivo")
    para(doc, "Estabelecer a metodologia para determinar, atualizar e usar os indicadores de desempenho energético (IDE) e as linhas de base energéticas (LBE), normalizar pelas variáveis "
              "relevantes e demonstrar a melhoria do desempenho energético. Esta metodologia é a informação documentada mantida exigida pela ISO 50001 6.4 e aplica a orientação da ISO 50006:2023.")
    h(doc, "2. Âmbito")
    para(doc, "Todos os IDE do SGE (instalação e USE) e respetivas LBE (RG-SGE-05).")
    h(doc, "3. Referências normativas e requisitos")
    bullets(doc, ["ISO 50001:2018 + Amd 1:2024 — 6.2, 6.4, 6.5, 9.1.1, 10.2", "ISO 50006:2023 — secções 4 a 10 e anexos A a G (2.ª edição; substitui a ISO 50006:2014)",
                  "ISO 50015:2014 e ISO 50047:2016 — M&V e poupanças", "EVO — IPMVP Core Concepts", "ASHRAE Guideline 14 — critérios estatísticos de modelos mensais",
                  "LBNL — EnPI Lite 'Valid model requirements' e 50001 Ready Navigator (tarefa 11)", "ISO 50003:2021 — melhoria demonstrada para a certificação",
                  "Kent, R. / BPF — 'impressão digital' energética de fábricas de plásticos (carga de base + consumo proporcional)"])
    sub(doc, "Requisitos legais e outros requisitos aplicáveis")
    bullets(doc, ["DL 71/2008 (SGCIE) — consumo específico e intensidade energética do PREn (RG-SGE-10)"])
    h(doc, "4. Definições (ISO 50006:2023, secção 3)")
    table(doc, ["Termo", "Definição"], [
        ("Período de referência", "Período usado para comparação com o período de reporte (3.1.1)."), ("Período de reporte", "Período em que se avalia o desempenho e a sua melhoria (3.1.16)."),
        ("LBE", "Valor que serve de base de comparação do desempenho energético (3.1.4); os dados e o método são retidos."),
        ("IDE", "Medida usada para quantificar o desempenho energético (3.1.10); pode ser calculado com um modelo energético."),
        ("Modelo energético", "Representação matemática da relação entre variáveis relevantes e consumo ou eficiência num período (3.1.8)."),
        ("Variável relevante", "Fator quantificável que afeta significativamente o desempenho e muda com frequência (3.1.15) — ex.: produção, temperatura."),
        ("Fator estático", "Fator que afeta significativamente o desempenho mas não muda com frequência (3.1.18) — ex.: n.º de máquinas, turnos, mix."),
        ("Normalização", "Processo que permite a análise em condições equivalentes ou normalizadas (3.1.13)."),
        ("Melhoria do desempenho energético", "Melhoria mensurável da eficiência ou do consumo, face à LBE (3.1.11).")], [4.5, 12.5])
    h(doc, "5. Responsabilidades")
    table(doc, ["Função", "Responsabilidade"], [("Gestor(a) de Energia", "Dono da metodologia e do IDE-01; calcula e valida os modelos."),
                                                ("Donos dos IDE (tbl_kpi)", "Acompanham, analisam e reportam o seu IDE (ISO 50006 tab. 1 — 'EnPI owner')."),
                                                ("Equipa de gestão de energia", "Aprova IDE, LBE e alterações."), ("Diretor Geral", "Assegura que os IDE representam o desempenho (5.1 l).")], [5, 12])
    h(doc, "6. Descrição do processo")
    sub(doc, "6.1 Utilizadores e fronteiras (ISO 50006 §5.2–5.3)")
    steps(doc, ["Identificar os utilizadores de cada IDE (gestão de topo, equipa, produção, engenheiros, externos) e as decisões que o IDE suporta.",
                "Definir a fronteira de cada IDE: instalação (IDE-01/02/05/09), USE (IDE-03/04/06/07) ou fronteira de medida (IDE-08). A fronteira deve poder ser medida."])
    sub(doc, "6.2 Variáveis relevantes e fatores estáticos (§5.5)")
    steps(doc, ["Listar as variáveis candidatas (tbl_variaveis) e calcular a correlação com o consumo nos 12 meses de referência.",
                "Aceitar uma variável se for significativa no modelo (p < 0,2) e tiver explicação física; preferir variáveis de saída útil (unidades) a variáveis de esforço (horas), para não esconder perdas.",
                "Listar os fatores estáticos e a sua condição na LBE (tbl_fatores_estaticos); revê-los trimestralmente."])
    sub(doc, "6.3 Escolha do tipo de IDE (§6.2)")
    table(doc, ["Tipo de IDE", "Quando usar", "Exemplo na Plasticom"], [
        ("Valor de energia medido", "Relato e consumo absoluto", "IDE-09 energia final (MWh)"),
        ("Rácio de valores medidos", "Uma entrada e uma saída; poucas variáveis", "IDE-02 kWh/1.000 un; IDE-08 kWh/Nm³"),
        ("Modelo estatístico", "Várias variáveis relevantes", "IDE-01 regressão com unidades e graus-dia"),
        ("Modelo de engenharia", "Projetos, simulação, sistemas interdependentes", "Previsão da UPAC (futuro)")], [4, 6, 7])
    sub(doc, "6.4 Período de referência e LBE (§7)")
    steps(doc, ["Usar 12 meses completos para cobrir o ciclo anual de temperatura; a LBE-01 usa mar/2025–fev/2026 (termina antes dos compressores VSD para que o seu efeito seja visível).",
                "Calcular o modelo com LINEST no próprio Excel (RG-SGE-05 LBE_Modelo): kWh = a + b1 × mil unidades (INJ+SOP) + b2 × graus-dia base 15 °C.",
                "Registar a LBE em tbl_lbe com o método, o período, o R² e o resultado dos testes."])
    sub(doc, "6.5 Validade estatística do modelo")
    table(doc, ["Teste", "Critério", "Fonte"], [("R²", "≥ 0,50 (mínimo); > 0,75 recomendado", "EnPI Lite (LBNL); ASHRAE 14"), ("p-valor das variáveis", "Todas < 0,20; pelo menos uma < 0,10", "EnPI Lite"),
                                                ("Teste F", "p < 0,10", "EnPI Lite"), ("CV(RMSE)", "≤ 15% (dados mensais)", "ASHRAE 14"), ("NMBE", "entre −5% e +5%", "ASHRAE 14"),
                                                ("Resultado", "Válido; válido com reserva (1 teste falha — justificar); não válido", "Decisão da equipa")], [4, 7, 6])
    sub(doc, "6.6 Normalização e domínio de validade (§8)")
    steps(doc, ["Energia esperada do mês = modelo da LBE aplicado às variáveis do mês; melhoria = (esperado − real) ÷ esperado.",
                "Só normalizar dentro do domínio: variáveis dentro do intervalo da referência ± 10%. Fora do domínio o mês não entra na demonstração da melhoria.",
                "Aplicar ajustes não rotineiros quando mudam os fatores estáticos (ex.: ALE-01 — 4 máquinas novas: retirar do real a energia atribuível e das variáveis as suas unidades)."])
    sub(doc, "6.7 Desvio significativo, CUSUM e demonstração da melhoria (§10)")
    steps(doc, ["Desvio significativo mensal: |real − esperado| > t(95%) × erro-padrão do modelo → investigar (PR-SGE-08).",
                "Acompanhar o CUSUM (soma acumulada de real − esperado): declive negativo = poupança sustentada.",
                "Demonstração: somar o esperado e o real dos meses válidos do período; a melhoria é demonstrada se a poupança exceder a incerteza a 95% ≈ t × erro-padrão × √m.",
                "Confirmar com as poupanças verificadas por ação (opções A/B) — devem ser coerentes com a instalação (opção C)."])
    sub(doc, "6.8 Revisão das LBE (6.5)")
    bullets(doc, ["a) Os IDE deixam de representar o desempenho (ex.: submedição desde 15/12/2026 → LBE por USE medidas).",
                  "b) Alterações significativas de fatores estáticos (tbl_fatores_estaticos).", "c) Método predeterminado: revisão anual com a revisão energética.",
                  "Cada alteração é retida em tbl_alteracoes_lbe com a justificação."])
    h(doc, "7. Resultado de referência (31/12/2026)")
    para(doc, "LBE-01: R² 0,75; CV(RMSE) 2,8%; carga de base ≈ 17%; válida com reserva (p-valor dos graus-dia 0,27). Mar–dez/2026: melhoria de 1,9% (100 MWh), dentro da incerteza de 106 MWh (não demonstrada). "
              "Out–dez/2026, após o standby e o programa de ar comprimido: melhoria de 5,4% (92 MWh) > incerteza (61 MWh) — demonstrada e coerente, dentro da incerteza, com 71 MWh verificados por ação.")
    h(doc, "8. Informação documentada (evidência)")
    table(doc, ["Registo", "Tabelas"], [("RG-SGE-05", "tbl_base_energia, tbl_ide_mensal, LBE_Modelo, tbl_variaveis, tbl_kpi, tbl_lbe, tbl_alteracoes_lbe, tbl_fatores_estaticos, Metodologia_IDE_LBE"),
                                        ("RG-SGE-11", "tbl_desvios, tbl_mv_planos, tbl_poupancas")], [4, 13])
    h(doc, "Histórico de revisões")
    table(doc, ["Versão", "Data", "Alteração"], [("00", "2026-09-29", "Emissão inicial (ISO 50006:2023)."),
                                                 ("01", "2026-12-18", "Demonstração no período após as ações; decisão RD-E-26-D03 sobre a reserva da LBE-01.")], [2, 3, 12])
    _save(doc, os.path.join(OUT, f"{code}_{slug(title)}.docx"))


def build_it(code, spec):
    title, ver, date, appr, st, regs = spec
    doc = new_doc(code, title, ver, date, "Em vigor", appr)
    h(doc, "Passos")
    for i, s in enumerate(st, start=1):
        p = doc.add_paragraph(f"{i}. {s}")
        p.runs[0].font.size = Pt(11)
    h(doc, "Registos"); para(doc, regs)
    _save(doc, os.path.join(OUT, f"{code}_{slug(title)}.docx"))


def build_politica():
    doc = new_doc("POL-SGE-01", "Política energética da Plasticom", "00", "2026-09-29", "Em vigor", author="Diretor Geral")
    para(doc, "A Plasticom fabrica embalagens plásticas num processo 100% elétrico. A energia é o nosso primeiro custo ambiental e uma alavanca de competitividade e de descarbonização. "
              "A gestão de topo compromete-se a:")
    bullets(doc, ["Melhorar continuamente o desempenho energético e o sistema de gestão da energia, com objetivos e metas energéticas revistos todos os anos (5.2 a, b, e).",
                  "Assegurar a informação e os recursos necessários para atingir os objetivos e as metas, incluindo a medição dos usos significativos (5.2 c).",
                  "Cumprir os requisitos legais aplicáveis e outros requisitos subscritos relacionados com a eficiência, o uso e o consumo de energia, incluindo o SGCIE (5.2 d).",
                  "Apoiar a aquisição de produtos e serviços energeticamente eficientes, avaliados pelo custo do ciclo de vida (5.2 f).",
                  "Apoiar atividades de projeto que considerem a melhoria do desempenho energético (5.2 g).",
                  "Contribuir para a descarbonização aumentando a eletricidade renovável e envolvendo trabalhadores, fornecedores e clientes."])
    para(doc, "Esta política é comunicada a todas as pessoas que trabalham para a Plasticom, está disponível às partes interessadas e é revista anualmente na revisão pela gestão.")
    para(doc, "Marinha Grande, 29 de setembro de 2026 — O Diretor Geral")
    _save(doc, os.path.join(OUT, "POL-SGE-01_Politica_energetica.docx"))


def build_manual():
    doc = new_doc("MAN-SGE-01", "Manual do Sistema de Gestão da Energia — Âmbito, Estrutura e Correspondência", "01", "2026-12-18", "Em vigor")
    h(doc, "1. A organização e o perfil energético")
    para(doc, f"{EMPRESA}: embalagens plásticas (frascos por injeção-sopro ISBM, tampas e potes por injeção) decoradas por serigrafia e hot foil. Processo 100% elétrico: ≈ 7,8 GWh "
              "em 2026 (≈ 1 680 tep; ≈ 28 TJ), 99,6% eletricidade e 0,4% gasóleo (frota e gerador). Consumidor intensivo de energia (SGCIE ≥ 1 000 tep).")
    h(doc, "2. Âmbito e fronteiras (4.3)")
    para(doc, "“Uso e consumo de energia na conceção, produção e decoração de embalagens plásticas (frascos, tampas e potes) por injeção, injeção-sopro, serigrafia e hot foil, incluindo "
              "utilidades, armazenagem, expedição, serviços de apoio e frota de serviço da Unidade 1 da Plasticom, Marinha Grande.”")
    para(doc, "Fronteiras: terreno de 45 000 m² da Unidade 1 e frota própria; nenhum tipo de energia excluído. Detalhe e mapa das fronteiras no RG-SGE-01.")
    h(doc, "3. Usos significativos de energia (6.3)")
    table(doc, ["USE", "Uso", "Peso em 2026", "IDE"], [("USE-01", "Sopro (ISBM)", "≈ 40%", "IDE-03"), ("USE-02", "Injeção", "≈ 27%", "IDE-04"),
                                                       ("USE-03", "Ar comprimido", "≈ 13%", "IDE-08"), ("USE-04", "Arrefecimento (pelo potencial)", "≈ 7%", "IDE-06")], [2.5, 6, 3, 5.5])
    h(doc, "4. Política, liderança e equipa (5)")
    para(doc, "Política energética POL-SGE-01. Equipa de gestão de energia nomeada em 01/07/2026 (OS-2026-07), liderada pelo(a) Gestor(a) de Energia com reporte ao Diretor Industrial; "
              "RACI dos 14 processos do SGE no RG-SGE-02.")
    h(doc, "5. Estrutura do SGE e integração no SGI")
    para(doc, "O SGE é o terceiro subsistema do SGI da Plasticom, com o SGA (ISO 14001:2026) e o SGQ (ISO 9001:2026). Partilha com o SGI: contexto (RG-SGA-01), riscos (RG-SGA-02), "
              "requisitos legais (RG-SGA-04), controlo da informação documentada (PR-SGA-08), planeamento de alterações (PR-SGA-06 / RG-SGA-18), auditorias e revisões integradas. "
              "As tabelas com o mesmo assunto têm o mesmo nome e colunas (RG-SGE-00 Matriz_Harmonizacao_SGI).")
    h(doc, "6. Correspondência requisito → documento → registo")
    table(doc, ["Cláusula", "Documento", "Registo"], [
        ("4.1–4.4", "MAN-SGE-01; PR-SGE-01", "RG-SGE-01"), ("5.1–5.3", "POL-SGE-01; MAN-SGE-01", "RG-SGE-02"), ("6.1", "PR-SGE-04", "RG-SGE-03"),
        ("6.2", "PR-SGE-04", "RG-SGE-05"), ("6.3", "PR-SGE-02", "RG-SGE-04"), ("6.4 / 6.5", "PR-SGE-03", "RG-SGE-05"), ("6.6", "PR-SGE-05", "RG-SGE-06"),
        ("7.1–7.4", "PR-SGE-09", "RG-SGE-07"), ("7.5", "PR-SGA-08 (SGI)", "RG-SGE-00; RG-SGA-09"), ("8.1", "PR-SGE-06; IT-SGE-01 a 03", "RG-SGE-08"),
        ("8.2 / 8.3", "PR-SGE-07", "RG-SGE-09"), ("9.1.1", "PR-SGE-05; PR-SGE-08", "RG-SGE-05; RG-SGE-06; RG-SGE-11"), ("9.1.2", "PR-SGE-10", "RG-SGE-10"),
        ("9.2 / 9.3", "PR-SGE-11", "RG-SGE-12; RG-SGE-13"), ("10.1 / 10.2", "PR-SGE-12", "RG-SGE-14")], [3, 6.5, 7.5])
    h(doc, "7. Normas e requisitos que suportam o SGE")
    bullets(doc, ["Núcleo certificável: NP EN ISO 50001:2019 (ISO 50001:2018 + Amd 1:2024).",
                  "Desempenho energético: ISO 50006:2023 (IDE/LBE), ISO 50015:2014 (M&V), ISO 50047:2016 (poupanças), ISO 50046:2019 (previsão), IPMVP, ASHRAE Guideline 14.",
                  "Implementação e auditoria: ISO 50004:2020, ISO 50005:2021, ISO/TS 50011:2023, ISO 50002-1:2025, EN 16247-1, ISO 50003:2021, ISO 19011.",
                  "Técnicas: EN 17267:2019 (medição), EN 17463:2021 (investimentos), ISO 11011:2013 (ar comprimido), ISO 1217, EUROMAP 60, IEC 60034-30-1, IEC 62053-22.",
                  "Integração: ISO 14001:2026, ISO 9001:2026, ISO 14064-1, ISO 14067, ISO 31000, ISO 55001 (lista completa em RG-SGE-00 Matriz_Normativa).",
                  "Legais: DL 71/2008 (SGCIE), Despacho 17313/2008, DL 68-A/2015, Diretiva (UE) 2023/1791, DL 15/2022 + DL 130/2026, Reg. (UE) 2019/1781, 2016/2281 e 2024/573 (RG-SGE-10)."])
    h(doc, "8. Estado à data de referência (31/12/2026)")
    bullets(doc, ["Prontidão para a ISO 50001: 84% (20 requisitos conformes, 7 parciais, 1 lacuna — 8.2) — RG-SGE-00.",
                  "Melhoria do desempenho energético demonstrada em out–dez/2026 (5,4%); certificação prevista para 09/2027 (RD-E-26-D08)."])
    h(doc, "Histórico de revisões")
    table(doc, ["Versão", "Data", "Alteração"], [("00", "2026-09-29", "Emissão inicial."), ("01", "2026-12-18", "Estado após a 1.ª revisão pela gestão.")], [2, 3, 12])
    _save(doc, os.path.join(OUT, "MAN-SGE-01_Manual_do_SGE.docx"))


def build(out_dir=OUT, force=False):
    global FORCE
    FORCE = force
    os.makedirs(out_dir, exist_ok=True)
    build_manual()
    build_politica()
    build_pr03()
    for code, spec in PROC.items():
        build_proc(code, spec)
    for code, spec in IT.items():
        build_it(code, spec)
    return sorted(os.listdir(out_dir))


if __name__ == "__main__":
    import sys
    for f in build(force="--force" in sys.argv):
        print(f)
