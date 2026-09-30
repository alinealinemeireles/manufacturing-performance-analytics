"""RG-SGE-04 — Revisão energética: tipos de energia, consumo por uso, usos significativos de energia (USE), variáveis relevantes,
pessoas que os influenciam, campanha de medição por máquina, oportunidades de melhoria priorizadas e previsão do consumo.
ISO 50001:2018 6.3 a)–e) · ISO 50002-1:2025 (método de auditoria) · ISO 50006:2023 §5.1 · ISO 11011 (ar comprimido) · EN 17267."""
from openpyxl.chart import BarChart, Reference
from sgelib import *
from dimse import *
import edata

P12 = (dt.date(2026, 1, 1), dt.date(2026, 12, 1))   # ano civil de 2026 (12 meses fechados)
CRIT_PCT, CRIT_POT = 0.10, 100                      # critérios de USE: ≥ 10% do consumo OU potencial ≥ 100 MWh/ano

OPORT = [
    # (ID, descrição, usos afetados, tipo, % poupança, investimento €, facilidade 1-3, referência técnica, plano, estado, nota)
    ("OPE-01", "Programa de deteção e reparação de fugas de ar comprimido (de ≈ 25% para ≤ 10%)", "UTL-AR", "Operacional", 0.15, 6000, 3,
     "BPF/Kent: fugas usam 20–40% do ar; ISO 11011:2013 (avaliação de sistemas de ar comprimido)", "PA-E-02", "Em curso", "Teste de vazio: 25% (jun/2026) → 14% (dez/2026) (RG-SGE-11)"),
    ("OPE-02", "Reduzir a pressão da rede de 7,4 para 6,8 bar (≈ 7% por bar)", "UTL-AR", "Operacional", 0.04, 500, 3,
     "BPF/Kent (pressão mínima necessária); ISO 11011", "PA-E-02", "Concluída", "6,8 bar desde 16/11/2026, sem queixas das máquinas (ISBM 40 bar têm booster próprio)"),
    ("OPE-03", "Modo standby nas pausas do 3.º turno e em paragens > 30 min (KZ-01)", "SOP;INJ", "Comportamental", 0.03, 0, 3,
     "BPF/Kent: máquina parada consome até 80–90% da potência de marcha; campanha PA-01 (kW em espera)", "PA-E-03", "Concluída", "Sugestão KZ-01 do RG-SGA-16; em vigor desde 05/10/2026"),
    ("OPE-04", "Fechar os ramais de ar das áreas paradas ao fim de semana (KZ-04)", "UTL-AR", "Operacional", 0.03, 1200, 3,
     "BPF/Kent (fugas alimentadas sem produção)", "PA-E-02", "Concluída", "Sugestão KZ-04 do RG-SGA-16; em vigor desde 09/10/2026"),
    ("OPE-05", "Isolamento térmico (mantas) dos canhões das 8 injetoras", "INJ", "Investimento", 0.05, 7200, 2,
     "BPF/Kent: isolamento do canhão reduz 5–10% da energia de aquecimento", "PA-E-04", "Aprovada", ""),
    ("OPE-06", "Servo-bomba nas injetoras hidráulicas de bomba fixa IM-002 e IM-004", "INJ", "Investimento", 0.08, 56000, 2,
     "Campanha PA-01: IM-002/004 consomem ≈ 60 kW vs ≈ 41 kW das servo; BPF/Kent (servo 30–50%)", "PA-E-04", "Aprovada", ""),
    ("OPE-07", "Recuperação do ar de alta pressão do sopro nas ISBM hidráulicas (5 máquinas)", "UTL-AR", "Investimento", 0.10, 45000, 2,
     "BPF/Kent: recuperação do ar de 40 bar no sopro de PET", "—", "Em estudo", "As ISBM elétricas já têm recuperação"),
    ("OPE-08", "Subir o setpoint da água gelada de 7 °C para 11 °C (≈ 3% por °C)", "UTL-FRIO", "Operacional", 0.12, 2000, 2,
     "BPF/Kent: +1 °C na água gelada ≈ −3% no chiller", "PA-E-06", "Em curso", "DOE concluído em 11/2026: 10 °C sem efeito na qualidade; 11 °C em ensaio"),
    ("OPE-09", "Free-cooling na torre quando T exterior < 9 °C", "UTL-FRIO", "Investimento", 0.20, 40000, 2,
     "BPF/Kent: pré-arrefecimento por ar com água > 12 °C, retorno < 2 anos", "PA-E-06", "Aprovada", "Depende de OPE-08"),
    ("OPE-10", "Variadores de velocidade nas bombas do circuito de arrefecimento", "UTL-FRIO", "Investimento", 0.10, 12000, 2,
     "BPF/Kent: VSD em bombas (−20% velocidade ≈ −49% potência); Reg. (UE) 2019/1781", "—", "Em estudo", ""),
    ("OPE-11", "Iluminação LED com sensores de presença nos armazéns", "GER", "Investimento", 0.06, 25000, 3,
     "BPF/Kent: iluminação ≈ 5% da energia de uma fábrica de plásticos", "—", "Em estudo", ""),
    ("OPE-12", "Central fotovoltaica de autoconsumo (UPAC ≈ 1 MWp, ≈ 1 400 MWh/ano)", "—", "Energia renovável", 0, 850000, 1,
     "SGA ALT-2026-07; EVI-010 do RG-SGA-17; DL 15/2022 alterado pelo DL 130/2026", "PA-E-07", "Em estudo", "Não reduz o consumo: substitui a origem (IDE-10)"),
    ("OPE-13", "Recuperação de calor dos compressores para águas sanitárias e AVAC", "GER", "Investimento", 0.10, 14000, 2,
     "ISO 11011; BPF/Kent (até 80% da energia do compressor é recuperável como calor)", "PA-E-05", "Aprovada", ""),
    ("OPE-14", "Substituir as ISBM hidráulicas mais antigas (ISBM-003, ISBM-005) por ISBM elétricas", "SOP", "Investimento", 0.07, 900000, 1,
     "Campanha PA-01: ISBM hidráulicas ≈ 60–66 kW vs ≈ 40 kW das elétricas", "—", "Em estudo", "Decisão com o plano de capacidade 2028 (LCC no RG-SGE-09)"),
    ("OPE-15", "Formação e sensibilização energética dos operadores e chefes de turno", "SOP;INJ;SER;HFS", "Comportamental", 0.02, 3000, 3,
     "BPF/Kent: formação pode reduzir até 20% (valor conservador de 2%)", "—", "Aprovada", "Plano de formação RG-SGE-07"),
    ("OPE-16", "Gestão da potência em horas de ponta (deslastre de cargas não críticas)", "—", "Custo (não energia)", 0, 8000, 2,
     "Faturas RG-SGE-09: potência tomada de 1 429 kW em ago/2026 (102% da contratada — ultrapassagem)", "—", "Em estudo", "Reduz custo e risco de ultrapassagem, não kWh"),
]

USE_INFO = {
    "SOP": ("USE-01", "Frascos produzidos; horas de marcha; massa do frasco; tecnologia (hidráulica/elétrica)", "IDE-03 (kWh/1.000 frascos); kW por máquina (campanha)",
            f"{GPROD}; {F['FE-17']}; {EPROC}; {GMAN}; {CTURNO}", "Temperatura dos fornos de pré-forma, pressão de sopro, standby em paragens, recuperação de ar (CO-01 a CO-04)"),
    "INJ": ("USE-02", "Peças produzidas; horas de marcha; tempo de ciclo; tecnologia (hidráulica/servo/elétrica)", "IDE-04 (kWh/1.000 un); kW por máquina (campanha)",
            f"{GPROD}; {F['FE-16']}; {EPROC}; {GMAN}; {CTURNO}", "Temperaturas do canhão por molde, standby em paragens, isolamento, manutenção hidráulica (CO-05 a CO-08)"),
    "UTL-AR": ("USE-03", "Procura de ar das máquinas; fugas; pressão de serviço", "IDE-08 (kWh/Nm³); % fugas no teste de vazio",
               f"{TUTL}; {F['FE-20']}; {GMAN}; operadores (uso de ar para limpeza)", "Pressão 6,8–7,2 bar, ronda de fugas, ramais fechados ao fim de semana, secadores (CO-09 a CO-12)"),
    "UTL-FRIO": ("USE-04", "Carga térmica dos moldes; temperatura exterior (graus-dia)", "IDE-06 (kWh frio ÷ kWh INJ+SOP)",
                 f"{TUTL}; {EPROC}; {F['FE-20']}", "Setpoint da água gelada, ΔT, limpeza dos permutadores, free-cooling (CO-13 a CO-15)"),
}


def build(out):
    b = Book("RG-SGE-04", "Revisão Energética e Usos Significativos de Energia",
             activities="Analisar o uso e o consumo de energia com base em medições e outros dados; identificar os USE; determinar variáveis relevantes, desempenho atual e pessoas que "
                        "influenciam cada USE; priorizar oportunidades de melhoria; estimar o consumo futuro. Atualização anual e após alterações maiores.",
             clauses="6.3 a) 1) tipos de energia atuais; a) 2) uso e consumo passado e atual; b) identificar USE; c) 1) variáveis relevantes, 2) desempenho atual, 3) pessoas que influenciam; "
                     "d) determinar e priorizar oportunidades; e) estimar uso e consumo futuros; métodos e critérios mantidos, resultados retidos; atualização em intervalos definidos e após alterações maiores.",
             purpose=f"Consumo mensal por uso e tipo de energia {PER_INI[:7]} a {PER_FIM[:7]} (fonte única RG-SGA-13 e RG-SGA-19); critérios de USE aplicados por fórmula sobre os 12 meses "
                     "jan–dez/2026 (Pareto); ficha de cada USE; campanha de medição por máquina e rateio mensal por máquina com ranking de consumo específico; registo de 16 oportunidades "
                     "com poupança, custo, retorno e prioridade calculados; previsão do consumo de 2027.",
             links=[("RG-SGA-13", "Eletricidade por uso (rateio) — fonte única."), ("RG-SGA-19", "Gasóleo da frota e do gerador (âmbito 1)."),
                    ("RG-SGA-03", "Aspetos ambientais de energia AA-007, AA-010, AA-021, AA-030 a AA-032 (mesmos usos)."),
                    ("RG-SGE-05", "IDE e LBE por USE; objetivos e planos de ação que tratam as oportunidades."), ("RG-SGE-08", "Critérios operacionais de cada USE (CO-xx)."),
                    ("RG-SGE-09", "LCC das oportunidades de investimento e compra de energia.")],
             guidance=[("ISO 50002-1:2025 (substitui a ISO 50002:2014) e EN 16247-1", "Estrutura da análise: dados, visita/medições, análise, oportunidades com poupança, custo e retorno."),
                       ("ISO 50004:2020 (anexo sobre revisão energética)", "Critérios de significância documentados; atualização anual e em alterações maiores."),
                       ("ISO 50006:2023 §5.1 e §5.4", "Tipos de energia, fluxos e SEU como base dos IDE."),
                       ("ISO 11011:2013", "Avaliação da central de ar comprimido (procura, fugas, pressão, potência específica)."),
                       ("EN 17267:2019", "Qualidade dos dados: medido vs estimado (coluna Qualidade_Dados)."),
                       ("Kent, R. / BPF — Energy Management in Plastics Processing", "Distribuição típica (ar ≈ 10%, frio 10–15%, iluminação ≈ 5%), perdas em espera, benchmarks de medidas."),
                       ("UNIDO — EnMS toolkit / DOE 50001 Ready Navigator (tarefas 8–10)", "Modelo do registo de oportunidades e da ficha de USE."),
                       ("Despacho 17313/2008", "Fatores de conversão para tep e intensidade carbónica (tbl_tipos_energia).")],
             legal=[("DL 71/2008 (SGCIE) — LEG-08", "A auditoria energética SGCIE (OBR-06, próxima 2030) usa esta revisão como base; o PREn integra as oportunidades."),
                    ("Diretiva (UE) 2023/1791, art. 11.º", "Empresas > 10 TJ/ano sem SGE: auditoria energética de 4 em 4 anos; com SGE ISO 50001 certificado ficam isentas (transposição em curso).")])
    b.add_list("Funcao", FUNC_NAMES)
    b.add_list("Uso", USO_CODES)
    b.add_list("TipoOport", ["Comportamental", "Operacional", "Investimento", "Energia renovável", "Custo (não energia)"])
    b.add_list("EstadoOport", ["Em estudo", "Aprovada", "Em curso", "Concluída", "Rejeitada"])
    b.add_list("Facilidade", [1, 2, 3])

    bm = edata.base_mensal()
    # ------------------------------------------------------------------ parâmetros
    PAR = [("PRECO_MEDIO", edata.preco_medio(), "€/kWh", "Preço médio da eletricidade (energia + redes) jan–dez/2026", "RG-SGE-09 (faturas simuladas)"),
           ("FE_ELET_LB", 0.110, "kgCO2e/kWh", "Fator de emissão location-based (cópia do FE-01 do RG-SGA-19)", "RG-SGA-19 tbl_fatores_emissao"),
           ("CRIT_USE_PCT", CRIT_PCT, "fração", "Critério 1 de USE: uso ≥ 10% do consumo total de energia", "Decisão da equipa de gestão de energia (6.3 b)"),
           ("CRIT_USE_POT", CRIT_POT, "MWh/ano", "Critério 2 de USE: potencial de melhoria identificado ≥ 100 MWh/ano", "Decisão da equipa de gestão de energia (6.3 b)"),
           ("CRESC_2027", 0.12, "fração", "Crescimento da produção em 2027 (ano completo das 4 máquinas novas, que em 2026 só produziram 6 meses)", "Plano de capacidade 2027 (RG-SGA-18 ALT-2026-03)")]
    parcols = [col("Codigo", 16, key="PK", desc="Parâmetro."), col("Valor", 10, "num3", desc="Valor."), col("Unidade", 11, desc="Unidade."),
               col("Descricao", 60, desc="Descrição."), col("Fonte", 40, desc="[Só SGE] Origem.")]
    b.table("Parametros", "tbl_parametros", parcols, [dict(zip(["Codigo", "Valor", "Unidade", "Descricao", "Fonte"], p_)) for p_ in PAR],
            "Parâmetros e critérios da revisão energética (mesma estrutura do RG-SGA-13 tbl_parametros).", title="PARÂMETROS E CRITÉRIOS DA REVISÃO ENERGÉTICA", row_height=18)
    tp = b.tables["tbl_parametros"]
    for i, p_ in enumerate(PAR):
        nm = p_[0].lower()
        b.wb.defined_names[nm] = DefinedName(nm, attr_text=f"Parametros!${tp['colmap']['Valor']}${tp['first'] + i}")

    # ------------------------------------------------------------------ tipos de energia (6.3 a 1)
    tcols = [col("ID_Tipo", 7, key="PK", desc="Tipo de energia."), col("Tipo_Energia", 36, desc="Tipo de energia (6.3 a 1)."), col("Origem", 36, desc="Origem / fornecedor."),
             col("Unidade_Compra", 9, desc="Unidade faturada."), col("kWh_por_Unidade", 9, "num1", desc="Conversão para kWh.", req=False),
             col("kgep_por_Unidade", 9, "num3", desc="Conversão para kgep (Despacho 17313/2008).", req=False),
             col("kgCO2e_por_Unidade_SGCIE", 10, "num3", desc="Intensidade carbónica SGCIE (Despacho 17313/2008).", req=False),
             col("kgCO2e_por_Unidade_GEE", 10, "num3", desc="Fator do inventário de GEE (RG-SGA-19).", req=False),
             col("Dentro_Ambito", 10, desc="Dentro do âmbito e fronteiras (4.3 — nenhum tipo de energia pode ser excluído)."),
             col("kWh_12m", 12, "kwh", f='=IF(@ID_Tipo@="","",SUMIFS(tbl_consumo_uso[kWh],tbl_consumo_uso[ID_Tipo],@ID_Tipo@,tbl_consumo_uso[Mes],">="&DATE(2026,1,1)))',
                 desc="Consumo em 2026 (kWh)."),
             col("Pct_12m", 8, "pct1", f='=IF(@ID_Tipo@="","",IFERROR(@kWh_12m@/SUM(#kWh_12m#),""))', desc="Peso no consumo total."),
             col("tep_12m", 9, "num1", f='=IF(OR(@ID_Tipo@="",@kWh_por_Unidade@=""),"",@kWh_12m@/@kWh_por_Unidade@*@kgep_por_Unidade@/1000)', desc="Energia primária (tep) — SGCIE.")]
    b.table("Tipos_Energia", "tbl_tipos_energia", tcols, [dict(zip(input_names(tcols), t_)) for t_ in TIPOS_ENERGIA],
            "Tipos de energia atuais e futuros com fatores de conversão (6.3 a 1).", title="TIPOS DE ENERGIA (6.3 a 1)",
            subtitle="Eletricidade: 0,215 kgep/kWh e 0,47 kgCO2e/kWh (Despacho 17313/2008, usados no SGCIE) · Gasóleo: 0,864 kgep/L · O inventário GEE usa fatores atuais (RG-SGA-19)", row_height=30)

    # ------------------------------------------------------------------ consumo mensal por uso (6.3 a 2)
    MAP = {"SOP": "kWh_SOP", "INJ": "kWh_INJ", "UTL-AR": "kWh_UTL_AR", "UTL-FRIO": "kWh_UTL_FRIO", "GER": "kWh_GER", "SER": "kWh_SER", "HFS": "kWh_HFS"}
    rows = []
    for _, r in bm.iterrows():
        for u in USO_CODES:
            if u in MAP:
                v, tipo, fonte, q = r[MAP[u]], "TE-01", "RG-SGA-13 tbl_dados_ambientais (rateio por horas de marcha)", "Estimado (rateio)"
            elif u == "FRO":
                v, tipo, fonte, q = r["L_Gasoleo_Frota"] * 10, "TE-02", "RG-SGA-19 S1-01 (cartão de frota) × 10 kWh/L", "Medido"
            else:
                v, tipo, fonte, q = r["L_Gasoleo_Gerador"] * 10, "TE-03", "RG-SGA-19 S1-02 (registo de testes) × 10 kWh/L", "Estimado"
            rows.append(dict(Mes=r["Mes"], ID_Uso=u, ID_Tipo=tipo, kWh=round(float(v)), Fonte=fonte, Qualidade_Dados=q))
    ccols = [col("Mes", 11, "date", desc="Mês.", key="PK (com ID_Uso)"), col("ID_Uso", 9, dv="Uso", desc="Uso de energia.", key="FK → tbl_usos"),
             col("ID_Tipo", 7, desc="Tipo de energia.", key="FK → tbl_tipos_energia"), col("kWh", 11, "kwh", desc="Consumo (kWh)."),
             col("Fonte", 50, desc="Origem."), col("Qualidade_Dados", 16, desc="Medido / estimado (EN 17267).")]
    b.table("Consumo_Mensal_Uso", "tbl_consumo_uso", ccols, rows, "Consumo mensal por uso e tipo de energia (formato longo, pronto para Power BI).",
            title="CONSUMO MENSAL POR USO DE ENERGIA (6.3 a 2)", subtitle="Eletricidade: RG-SGA-13 (soma dos usos = fatura) · Gasóleo: RG-SGA-19 · Uma linha por mês × uso",
            cf=[("Qualidade_Dados", {"Estimado": "orange", "Medido": "green"})], row_height=15, freeze_col=2)

    # ------------------------------------------------------------------ usos e critérios de USE (6.3 b)
    ucols = [col("ID_Uso", 9, key="PK", desc="Uso de energia (mesmo código do RG-SGA-13)."), col("Uso", 50, desc="Descrição."), col("Tipo_Energia", 16, desc="Tipo de energia."),
             col("Equipamentos", 24, desc="Equipamentos."),
             col("kWh_12m", 11, "kwh", f='=IF(@ID_Uso@="","",SUMIFS(tbl_consumo_uso[kWh],tbl_consumo_uso[ID_Uso],@ID_Uso@,tbl_consumo_uso[Mes],">="&DATE(2026,1,1)))', desc="Consumo de 2026 (12 meses)."),
             col("Pct_Total", 8, "pct1", f='=IF(@ID_Uso@="","",@kWh_12m@/SUM(#kWh_12m#))', desc="Peso no consumo total de energia."),
             col("Ranking", 7, "int", f='=IF(@ID_Uso@="","",RANK(@kWh_12m@,#kWh_12m#,0))', desc="Posição (1 = maior)."),
             col("Pct_Acumulada", 9, "pct1", f='=IF(@ID_Uso@="","",SUMIF(#kWh_12m#,">="&@kWh_12m@)/SUM(#kWh_12m#))', desc="Pareto: peso acumulado até este uso."),
             col("Potencial_MWh", 10, "num0", f='=IF(@ID_Uso@="","",SUMPRODUCT(ISNUMBER(SEARCH(@ID_Uso@&";",tbl_oportunidades[Usos_Afetados]&";"))*tbl_oportunidades[Poupanca_MWh_Parcela]))',
                 desc="Potencial de melhoria identificado (Σ das oportunidades, repartido pelos usos afetados)."),
             col("Criterio_Consumo", 9, f='=IF(@ID_Uso@="","",IF(@Pct_Total@>=crit_use_pct,"Sim","Não"))', desc="≥ 10% do consumo?"),
             col("Criterio_Potencial", 9, f='=IF(@ID_Uso@="","",IF(@Potencial_MWh@>=crit_use_pot,"Sim","Não"))', desc="Potencial ≥ 100 MWh/ano?"),
             col("E_USE", 7, f='=IF(@ID_Uso@="","",IF(OR(@Criterio_Consumo@="Sim",@Criterio_Potencial@="Sim"),"USE","—"))', desc="Uso significativo (6.3 b)."),
             col("Qualidade_Dados", 16, desc="Medido / estimado."), col("Aspeto_SGA", 16, desc="Aspeto ambiental equivalente (RG-SGA-03).", req=False)]
    ASP = {"SOP": "AA-010", "INJ": "AA-007", "UTL-AR": "AA-021", "GER": "AA-030", "HFS": "AA-019"}
    urows = [dict(ID_Uso=u[0], Uso=u[1], Tipo_Energia=u[2], Equipamentos=u[4], Qualidade_Dados="Medido" if u[0] == "FRO" else "Estimado (rateio)" if u[2] == "Eletricidade" else "Estimado",
                  Aspeto_SGA=ASP.get(u[0])) for u in USOS]
    b.table("Usos_Energia", "tbl_usos", ucols, urows, "Usos de energia com critérios de significância calculados (6.3 b).",
            title="USOS DE ENERGIA E DETERMINAÇÃO DOS USE (6.3 b)", subtitle="USE = uso ≥ 10% do consumo total OU potencial de melhoria ≥ 100 MWh/ano · 12 meses jan–dez/2026 · Critérios em tbl_parametros",
            cf=[("E_USE", {"USE": "red"}), ("Criterio_Consumo", {"Sim": "orange"}), ("Criterio_Potencial", {"Sim": "orange"})], row_height=30, freeze_col=2)
    ws = b.wb["Usos_Energia"]
    tu = b.tables["tbl_usos"]
    ch = BarChart()
    ch.type, ch.title, ch.height, ch.width = "col", "Consumo de energia por uso — 12 meses (kWh)", 8, 18
    ci = list(tu["colmap"]).index("kWh_12m") + 1
    ch.add_data(Reference(ws, min_col=ci, min_row=tu["hr"], max_row=tu["last"]), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=tu["first"], max_row=tu["last"]))
    ch.legend = None
    ch.y_axis.numFmt = "#,##0"
    fix_chart(ch)
    ws.add_chart(ch, f"B{tu['last'] + 3}")

    # ------------------------------------------------------------------ ficha dos USE (6.3 c)
    scols = [col("ID_USE", 8, key="PK", desc="Uso significativo."), col("ID_Uso", 9, desc="Uso.", key="FK → tbl_usos"),
             col("Pct_Consumo", 8, "pct1", f='=IF(@ID_USE@="","",IFERROR(INDEX(tbl_usos[Pct_Total],MATCH(@ID_Uso@,tbl_usos[ID_Uso],0)),""))', desc="Peso no consumo."),
             col("Confirmado_USE", 9, f='=IF(@ID_USE@="","",IFERROR(INDEX(tbl_usos[E_USE],MATCH(@ID_Uso@,tbl_usos[ID_Uso],0)),""))', desc="Confirmação pelos critérios."),
             col("Variaveis_Relevantes", 40, desc="Variáveis relevantes (6.3 c 1)."), col("Desempenho_Atual", 34, desc="IDE e valor atual (6.3 c 2) — ver RG-SGE-05."),
             col("Pessoas_Influenciam", 44, desc="Pessoas que trabalham sob o controlo da organização e influenciam o USE (6.3 c 3) — competência no RG-SGE-07."),
             col("Controlos_Operacionais", 44, desc="Critérios de operação e manutenção (8.1) — RG-SGE-08."),
             col("IDE", 10, desc="IDE do USE.", key="FK → RG-SGE-05 tbl_kpi")]
    srows = [dict(ID_USE=v[0], ID_Uso=k, Variaveis_Relevantes=v[1], Desempenho_Atual=v[2], Pessoas_Influenciam=v[3], Controlos_Operacionais=v[4],
                  IDE={"SOP": "IDE-03", "INJ": "IDE-04", "UTL-AR": "IDE-08", "UTL-FRIO": "IDE-06"}[k]) for k, v in USE_INFO.items()]
    b.table("Ficha_USE", "tbl_use", scols, srows, "Ficha de cada USE: variáveis relevantes, desempenho atual, pessoas que influenciam (6.3 c).",
            title="FICHA DOS USOS SIGNIFICATIVOS DE ENERGIA (6.3 c 1–3)", row_height=60, freeze_col=2)

    # ------------------------------------------------------------------ campanha de medição e máquinas
    camp = edata.campanha()
    mcols = [col("ID_Maquina", 10, key="PK", desc="Máquina (ID do dataset)."), col("Processo", 7, desc="Uso."), col("Ano_Instalacao", 8, "int", desc="Ano de instalação (dim_machine_profile)."),
             col("Tecnologia", 30, desc="Acionamento (simulado coerente com o ano)."), col("Data_Medicao", 11, "date", desc="Início da medição."), col("Duracao_h", 7, "int", desc="Duração (h)."),
             col("Instrumento", 9, desc="Analisador portátil (RG-SGE-06 tbl_equipamentos).", key="FK → RG-SGE-06"),
             col("kW_Producao", 9, "num1", desc="Potência média em produção (medida)."), col("kW_Espera_Aquecida", 9, "num1", desc="Potência em espera com resistências ligadas."),
             col("kW_Desligada", 8, "num1", desc="Consumo residual desligada (comandos)."),
             col("Fracao_Espera", 8, "pct", f='=IF(@ID_Maquina@="","",@kW_Espera_Aquecida@/@kW_Producao@)', desc="Espera ÷ produção (Kent: até 80–90% nas hidráulicas)."),
             col("kW_vs_Media_Processo", 9, "pct", f='=IF(@ID_Maquina@="","",@kW_Producao@/AVERAGEIFS(#kW_Producao#,#Processo#,@Processo@)-1)', desc="Desvio face à média do processo."),
             col("Horas_Marcha_12m", 9, "num0", f='=IF(@ID_Maquina@="","",SUMIFS(tbl_energia_maquina[Horas_Marcha],tbl_energia_maquina[ID_Maquina],@ID_Maquina@,tbl_energia_maquina[Mes],">="&DATE(2026,1,1)))', desc="Horas de marcha em 2026."),
             col("Horas_Paragem_12m", 9, "num0", f='=IF(@ID_Maquina@="","",SUMIFS(tbl_energia_maquina[Horas_Paragem],tbl_energia_maquina[ID_Maquina],@ID_Maquina@,tbl_energia_maquina[Mes],">="&DATE(2026,1,1)))', desc="Horas de paragem + setup em 2026."),
             col("kWh_Espera_12m", 10, "kwh", f='=IF(@ID_Maquina@="","",@kW_Espera_Aquecida@*@Horas_Paragem_12m@)', desc="Energia gasta em espera (oportunidade OPE-03)."),
             col("kWh_12m", 10, "kwh", f='=IF(@ID_Maquina@="","",SUMIFS(tbl_energia_maquina[kWh_Rateado],tbl_energia_maquina[ID_Maquina],@ID_Maquina@,tbl_energia_maquina[Mes],">="&DATE(2026,1,1)))', desc="Eletricidade rateada em 2026."),
             col("Unid_12m", 11, "num0", f='=IF(@ID_Maquina@="","",SUMIFS(tbl_energia_maquina[Unid],tbl_energia_maquina[ID_Maquina],@ID_Maquina@,tbl_energia_maquina[Mes],">="&DATE(2026,1,1)))', desc="Unidades em 2026."),
             col("SEC_kWh_1000", 9, "num1", f='=IF(OR(@ID_Maquina@="",@Unid_12m@=0),"",@kWh_12m@/@Unid_12m@*1000)', desc="kWh por 1.000 unidades (depende também da peça e do molde)."),
             col("Prioridade_Maquina", 12, f='=IF(@ID_Maquina@="","",IF(AND(@kW_vs_Media_Processo@>0.15,@Fracao_Espera@>0.5),"Alta",IF(OR(@kW_vs_Media_Processo@>0.15,@Fracao_Espera@>0.5),"Média","Baixa")))',
                 desc="Alta = potência > média do processo + 15% E espera > 50%.")]
    b.table("Campanha_Maquinas", "tbl_campanha", mcols, [dict(**{k: r[k] for k in input_names(mcols) if k in r}) for r in camp.to_dict("records")],
            "Campanha de medição por máquina com analisador portátil (simulada, jun–jul/2026) e indicadores por máquina (6.3 a 2, c 2).",
            title="CAMPANHA DE MEDIÇÃO POR MÁQUINA (analisador PA-01, 48 h por máquina) — SIMULADA",
            subtitle="Calibrada para reproduzir os kW médios por processo do RG-SGA-13 · Energia por máquina = rateio do valor do processo por kW × horas (tbl_energia_maquina)",
            cf=[("Prioridade_Maquina", {"Alta": "red", "Média": "orange", "Baixa": "green"})], row_height=16, freeze_col=2)
    em = edata.energia_maquinas()
    erows = [dict(Mes=r["Mes"], ID_Maquina=r["MachineId"], Processo=r["Processo"], Unid=int(r["Unid"]), Horas_Marcha=round(float(r["Horas_Marcha"]), 1),
                  Horas_Paragem=round(float(r["Horas_Paragem"] + r["Horas_Setup"]), 1), kWh_Rateado=float(r["kWh_Rateado"])) for _, r in em.iterrows()]
    ecols = [col("Mes", 11, "date", desc="Mês.", key="PK (com ID_Maquina)"), col("ID_Maquina", 10, desc="Máquina.", key="FK → tbl_campanha"), col("Processo", 7, desc="Uso."),
             col("Unid", 11, "num0", desc="Unidades produzidas (dataset)."), col("Horas_Marcha", 9, "num1", desc="Horas de marcha (dataset)."),
             col("Horas_Paragem", 9, "num1", desc="Paragens não planeadas + setup (dataset)."), col("kWh_Rateado", 10, "kwh", desc="Eletricidade do processo (RG-SGA-13) × peso da máquina."),
             col("SEC_kWh_1000", 9, "num1", f='=IF(OR(@Mes@="",@Unid@=0),"",@kWh_Rateado@/@Unid@*1000)', desc="kWh por 1.000 unidades.")]
    b.table("Energia_Maquina_Mensal", "tbl_energia_maquina", ecols, erows, "Eletricidade mensal por máquina (rateio pelo kW medido × horas) jul/2025–dez/2026.",
            title="ELETRICIDADE MENSAL POR MÁQUINA (rateio ponderado pela campanha)", subtitle="Σ por processo e mês = RG-SGA-13 · Substituído pela medição dos analisadores a partir de 12/2026", row_height=15)

    # ------------------------------------------------------------------ oportunidades (6.3 d)
    ocols = [col("ID_Oportunidade", 9, key="PK", desc="Oportunidade de melhoria do desempenho energético."), col("Descricao", 50, desc="Descrição."),
             col("Usos_Afetados", 14, desc="Usos afetados (códigos separados por ';')."), col("Tipo", 14, dv="TipoOport", desc="Tipo."),
             col("Base_kWh", 11, "kwh", f='=IF(@ID_Oportunidade@="","",SUMPRODUCT(ISNUMBER(SEARCH(tbl_usos[ID_Uso]&";",@Usos_Afetados@&";"))*tbl_usos[kWh_12m]))', desc="Consumo 12 meses dos usos afetados."),
             col("Pct_Poupanca", 8, "pct", desc="Poupança estimada (% da base)."),
             col("Poupanca_MWh", 9, "num0", f='=IF(@ID_Oportunidade@="","",@Base_kWh@*@Pct_Poupanca@/1000)', desc="Poupança anual (MWh)."),
             col("Poupanca_EUR", 10, "eur", f='=IF(@ID_Oportunidade@="","",@Poupanca_MWh@*1000*preco_medio)', desc="Poupança anual (€) ao preço médio."),
             col("Investimento_EUR", 10, "eur", desc="Investimento estimado."),
             col("Payback_Anos", 8, "num1", f='=IF(OR(@ID_Oportunidade@="",@Poupanca_EUR@=0),"",@Investimento_EUR@/@Poupanca_EUR@)', desc="Retorno simples (anos)."),
             col("tCO2e_Evitadas", 8, "num1", f='=IF(@ID_Oportunidade@="","",@Poupanca_MWh@*fe_elet_lb)', desc="Emissões evitadas (location-based)."),
             col("Facilidade", 7, "int", dv="Facilidade", desc="1 difícil · 2 média · 3 fácil."),
             col("Pontuacao", 8, "num1", f='=IF(@ID_Oportunidade@="","",MIN(3,@Poupanca_MWh@/50)+IF(@Payback_Anos@="",0,IF(@Payback_Anos@<=1,3,IF(@Payback_Anos@<=3,2,1)))+@Facilidade@)',
                 desc="Priorização (6.3 d): energia (0–3) + retorno (1–3) + facilidade (1–3)."),
             col("Prioridade", 9, f='=IF(@ID_Oportunidade@="","",IF(@Pontuacao@>=7,"Alta",IF(@Pontuacao@>=5,"Média","Baixa")))', desc="Alta ≥ 7 · Média ≥ 5 · Baixa."),
             col("Referencia_Tecnica", 44, desc="Fonte da estimativa."), col("ID_Plano", 9, desc="Plano de ação (RG-SGE-05).", key="FK → RG-SGE-05 tbl_planos_acao"),
             col("Estado", 10, dv="EstadoOport", desc="Estado."), col("Nota", 36, desc="Nota.", req=False),
             col("N_Usos", 6, "int", f='=IF(@ID_Oportunidade@="","",IF(@Usos_Afetados@="—",0,LEN(@Usos_Afetados@)-LEN(SUBSTITUTE(@Usos_Afetados@,";",""))+1))', desc="N.º de usos afetados."),
             col("Poupanca_MWh_Parcela", 9, "num0", f='=IF(OR(@ID_Oportunidade@="",@N_Usos@=0),0,@Poupanca_MWh@/@N_Usos@)', desc="Parcela por uso (para o potencial de cada uso).")]
    ON = ["ID_Oportunidade", "Descricao", "Usos_Afetados", "Tipo", "Pct_Poupanca", "Investimento_EUR", "Facilidade", "Referencia_Tecnica", "ID_Plano", "Estado", "Nota"]
    b.table("Oportunidades", "tbl_oportunidades", ocols, rows_from(ON, OPORT), "Oportunidades de melhoria do desempenho energético determinadas e priorizadas (6.3 d).",
            title="REGISTO DE OPORTUNIDADES DE MELHORIA DO DESEMPENHO ENERGÉTICO (6.3 d)",
            subtitle="Base = consumo de 12 meses dos usos afetados · Poupança, €, retorno, CO2 e prioridade calculados · Estimativas a confirmar por M&V (RG-SGE-11)",
            cf=[("Prioridade", {"Alta": "green", "Média": "yellow", "Baixa": "gray"})], row_height=45, freeze_col=2)

    # ------------------------------------------------------------------ previsão 2027 (6.3 e)
    pcols = [col("ID_Uso", 9, key="PK", desc="Uso."), col("kWh_12m", 11, "kwh", f='=IF(@ID_Uso@="","",INDEX(tbl_usos[kWh_12m],MATCH(@ID_Uso@,tbl_usos[ID_Uso],0)))', desc="Consumo atual."),
             col("Elasticidade_Producao", 9, "num", desc="Quanto o uso acompanha a produção (1 = proporcional; 0 = fixo)."),
             col("kWh_Tendencia_2027", 11, "kwh", f='=IF(@ID_Uso@="","",@kWh_12m@*(1+cresc_2027*@Elasticidade_Producao@))', desc="Sem novas ações."),
             col("Poupanca_Aprovada_MWh", 10, "num0", f='=IF(@ID_Uso@="","",SUMPRODUCT(ISNUMBER(SEARCH(@ID_Uso@&";",tbl_oportunidades[Usos_Afetados]&";"))*((tbl_oportunidades[Estado]="Aprovada")+(tbl_oportunidades[Estado]="Em curso"))*tbl_oportunidades[Poupanca_MWh_Parcela]))',
                 desc="Oportunidades aprovadas ou em curso (efeito anual; as concluídas em 2026 já estão parcialmente no consumo de 2026)."),
             col("kWh_Previsto_2027", 11, "kwh", f='=IF(@ID_Uso@="","",@kWh_Tendencia_2027@-@Poupanca_Aprovada_MWh@*1000)', desc="Previsão com as ações aprovadas."),
             col("Pressuposto", 50, desc="Pressupostos.")]
    PREV = [("SOP", 1.0, "Frascos +12% (ISBM-009/010 em ano completo)"), ("INJ", 1.0, "Tampas e potes +12% (IM-007/008)"), ("UTL-AR", 1.0, "Acompanha a procura das máquinas"),
            ("UTL-FRIO", 1.0, "Acompanha a carga térmica; verão igual a 2026"), ("GER", 0.1, "Quase fixo"), ("SER", 0.5, "Decoração cresce metade da produção"),
            ("HFS", 0.5, "Idem"), ("FRO", 0.0, "Frota igual"), ("GER-EMG", 0.0, "Testes iguais")]
    b.table("Previsao_2027", "tbl_previsao", pcols, [dict(ID_Uso=u, Elasticidade_Producao=e, Pressuposto=p_) for u, e, p_ in PREV],
            "Estimativa do uso e consumo futuros (6.3 e).", title="PREVISÃO DO CONSUMO DE ENERGIA 2027 (6.3 e)",
            subtitle="Crescimento da produção em tbl_parametros (CRESC_2027) · A UPAC não reduz o consumo (reduz a compra à rede)", row_height=18)
    ws = b.wb["Previsao_2027"]
    tq = b.tables["tbl_previsao"]
    r = tq["last"] + 2
    cell(ws, r, 1, "Total", bold=True)
    for cname in ("kWh_12m", "kWh_Tendencia_2027", "Poupanca_Aprovada_MWh", "kWh_Previsto_2027"):
        L_ = tq["colmap"][cname]
        cell(ws, r, list(tq["colmap"]).index(cname) + 1, f"=SUM({L_}{tq['first']}:{L_}{tq['last']})", fmt="#,##0", bold=True)

    ws = b.sheet("Metodo_Criterios", "Métodos e critérios da revisão energética (6.3 — manter como informação documentada; detalhe no PR-SGE-02).")
    title(ws, "MÉTODOS E CRITÉRIOS DA REVISÃO ENERGÉTICA (6.3 — informação documentada mantida)")
    notes(ws, 3, [
        ("Frequência", "Anual (antes da revisão pela gestão) e sempre que houver alteração maior em instalações, equipamentos, sistemas ou processos (ligação ao PR-SGA-06 / RG-SGA-18)."),
        ("Dados usados", "Fatura (medido), rateio por horas × kW do RG-SGA-13 (estimado), campanha PA-01 por máquina (medido, 48 h), caudalímetro EQP-04 e teste de vazio (ar), dataset de produção, IPMA."),
        ("Critério de USE", "Uso com ≥ 10% do consumo total de energia OU com potencial de melhoria identificado ≥ 100 MWh/ano (tbl_parametros). O Pareto (Pct_Acumulada) é informativo."),
        ("Variáveis relevantes", "Candidatas testadas por correlação e regressão (RG-SGE-05 tbl_variaveis); aceites se significativas e com explicação física."),
        ("Pessoas que influenciam", "Funções que operam, mantêm, compram ou projetam equipamentos do USE — ligação às competências (RG-SGE-07)."),
        ("Priorização das oportunidades", "Pontuação = energia (MWh/50, máx. 3) + retorno (≤ 1 ano 3; ≤ 3 anos 2; > 3 anos 1) + facilidade (1–3). Alta ≥ 7; Média ≥ 5."),
        ("Previsão", "Consumo de 12 meses × (1 + crescimento × elasticidade) − poupanças aprovadas."),
        ("Resultados principais (jan–dez/2026)", "Sopro ≈ 40%, injeção ≈ 26%, ar comprimido ≈ 13% e frio ≈ 7% da energia: 4 USE (o frio pelo potencial). Máquinas hidráulicas antigas (IM-002/004, ISBM-002/003/005/006/008) "
                                                    "consomem 40–60% mais que as elétricas e perdem 60–78% da potência em espera. Carga de base ≈ 17% (RG-SGE-05)."),
        ("Limitação", "Até 15/12/2026 os consumos por uso são rateio (analisadores M01–M06 em serviço desde então): a revisão de 2027 será refeita com 12 meses de submedição (ALE-03)."),
    ])
    return b.save(out)


if __name__ == "__main__":
    import sys
    print(build(sys.argv[1]))
