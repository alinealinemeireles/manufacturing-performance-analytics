"""RG-SGE-00 — Índice do SGE e modelo de dados: catálogo de ficheiros e tabelas, integridade referencial (inclui ligações ao SGA),
matriz de requisitos ISO 50001:2018 + Amd 1:2024 com prontidão, matriz mestra de informação documentada, matriz normativa (família ISO 50000
e normas de integração), fontes citadas, fonte única SGE ↔ SGA ↔ SGQ, harmonização de colunas e autoavaliação de maturidade (ISO 50005)."""
import glob
import os
import re
import openpyxl
import pandas as pd
from sgelib import *
from dimse import *

FILES = [
    ("SGE-01_Contexto_Ambito_Fronteiras.xlsx", "RG-SGE-01", "Contexto de energia ligado ao RG-SGA-01, 10 partes interessadas, alterações climáticas (Amd 1:2024), âmbito e mapa de fronteiras", "4.1; 4.2; 4.3"),
    ("SGE-02_Lideranca_Politica_Equipa.xlsx", "RG-SGE-02", "Liderança 5.1 a–m, política energética POL-SGE-01 (5.2 a–g), equipa de gestão de energia e RACI", "5.1; 5.2; 5.3"),
    ("SGE-03_Riscos_Oportunidades.xlsx", "RG-SGE-03", "Riscos e oportunidades do SGE (vista do RG-SGA-02 + RE/OE novos) com mapa de calor", "6.1"),
    ("SGE-04_Revisao_Energetica_USE.xlsx", "RG-SGE-04", "Revisão energética: tipos de energia, consumo por uso, 4 USE, campanha por máquina, 16 oportunidades priorizadas, previsão 2027", "6.3"),
    ("SGE-05_IDE_LBE_Objetivos_Metas.xlsx", "RG-SGE-05", "Base mensal de energia, LBE-01 por LINEST com testes, IDE mensais (esperado, CUSUM, domínio), catálogo de IDE, LBE, objetivos, metas e planos", "6.2; 6.4; 6.5; 9.1.1"),
    ("SGE-06_Recolha_Dados_Medicao.xlsx", "RG-SGE-06", "Árvore de contadores (EN 17267), plano de recolha 6.6 a–e, plano de monitorização 9.1.1, equipamentos e reconciliação", "6.6; 9.1.1"),
    ("SGE-07_Competencias_Comunicacao.xlsx", "RG-SGE-07", "Competências por função e USE, formação com eficácia, consciencialização, comunicação e sugestões", "7.1–7.4"),
    ("SGE-08_Controlo_Operacional_USE.xlsx", "RG-SGE-08", "16 critérios operacionais com desvio significativo, rondas de energia jul–dez/2026, alterações não intencionais", "8.1"),
    ("SGE-09_Projeto_Aquisicoes_Compra_Energia.xlsx", "RG-SGE-09", "Projetos (8.2), critérios e aquisições (8.3), LCC (EN 17463), compra de energia, faturas por período tarifário", "8.2; 8.3"),
    ("SGE-10_Conformidade_Legal_SGCIE.xlsx", "RG-SGE-10", "Requisitos legais de energia (estrutura do RG-SGA-04), obrigações e trajetória do PREn (SGCIE)", "4.2; 9.1.2"),
    ("SGE-11_Desvios_MV_Poupancas.xlsx", "RG-SGE-11", "Desvios significativos investigados, planos de M&V (IPMVP A/B/C), ar comprimido em 3 fases, poupanças verificadas", "9.1.1; 6.2.3; 10.2"),
    ("SGE-12_Auditoria_Interna.xlsx", "RG-SGE-12", "Auditores, programa 2026–2027, checklist e constatações da AUD-E-2026-01", "9.2"),
    ("SGE-13_Revisao_pela_Gestao.xlsx", "RG-SGE-13", "Revisão pela gestão RD-E-2026-01: agenda 9.3.2/9.3.3, decisões 9.3.4 a–g, ata", "9.3"),
    ("SGE-14_NC_Acao_Corretiva_Melhoria.xlsx", "RG-SGE-14", "NC com o ciclo 10.1 a–e, 5 Porquês, evidência trimestral da melhoria contínua", "10.1; 10.2"),
]

REL = [
    ("tbl_contexto_energia", "ID_Contexto_SGI", "SGA:tbl_pestel|SGA:tbl_swot|SGA:tbl_tows", "ID_PESTEL|ID_SWOT|ID_TOWS"),
    ("tbl_partes_interessadas", "ID_RO", "tbl_riscos_e|tbl_oport_e", "ID_Risco|ID_Oport"), ("tbl_partes_interessadas", "ID_PI_SGI", "SGA:tbl_partes_interessadas", "ID_PI"),
    ("tbl_partes_interessadas", "IDs_Legal", "tbl_legal|SGA:tbl_legal", "ID_Legal|ID_Legal"),
    ("tbl_riscos_e", "ID_Risco", "SGA:tbRiscos", "ID", r"^R\d"), ("tbl_oport_e", "ID_Oport", "SGA:tbOportunidades", "ID", r"^O\d"),
    ("tbl_riscos_e", "ID_Plano", "tbl_planos_acao", "ID_Plano"), ("tbl_riscos_e", "KPI", "tbl_kpi", "ID_KPI"),
    ("tbl_usos", "Aspeto_SGA", "SGA:tbl_aspetos", "ID_Aspeto"), ("tbl_consumo_uso", "ID_Uso", "tbl_usos", "ID_Uso"),
    ("tbl_use", "ID_Uso", "tbl_usos", "ID_Uso"), ("tbl_use", "IDE", "tbl_kpi", "ID_KPI"), ("tbl_oportunidades", "ID_Plano", "tbl_planos_acao", "ID_Plano"),
    ("tbl_energia_maquina", "ID_Maquina", "tbl_campanha", "ID_Maquina"),
    ("tbl_kpi", "ID_OBJ", "tbl_objetivos", "ID_OBJ"), ("tbl_kpi", "ID_MED", "tbl_plano_recolha", "ID_Dado"), ("tbl_kpi", "ID_LBE", "tbl_lbe", "ID_LBE"),
    ("tbl_objetivos", "ID_KPI", "tbl_kpi", "ID_KPI"), ("tbl_objetivos", "ID_Politica", "tbl_politica", "ID_Politica"), ("tbl_objetivos", "Acoes_PAM", "tbl_planos_acao", "ID_Plano"),
    ("tbl_objetivos", "Origem_IDs", "tbl_use|tbl_oportunidades|tbl_riscos_e|SGA:tbl_pestel|SGA:tbOportunidades|SGA:tbRiscos|SGA:tbl_alteracoes|SGA:tbl_kaizen|tbl_alteracoes_lbe",
     "ID_USE|ID_Oportunidade|ID_Risco|ID_PESTEL|ID|ID|ID_Alteracao|ID_KZ|ID_Alteracao_LBE"),
    ("tbl_metas_energeticas", "ID_OBJ", "tbl_objetivos", "ID_OBJ"), ("tbl_metas_energeticas", "ID_KPI", "tbl_kpi", "ID_KPI"),
    ("tbl_planos_acao", "ID_OBJ", "tbl_objetivos", "ID_OBJ"), ("tbl_planos_acao", "ID_Oportunidade", "tbl_oportunidades", "ID_Oportunidade"), ("tbl_planos_acao", "ID_PAM_SGA", "SGA:tbl_pam", "ID_Acao"),
    ("tbl_lbe", "ID_KPI", "tbl_kpi", "ID_KPI"), ("tbl_fatores_estaticos", "ID_Alteracao_LBE", "tbl_alteracoes_lbe", "ID_Alteracao_LBE"),
    ("tbl_contadores", "ID_Equipamento", "tbl_equipamentos", "ID_Equipamento"), ("tbl_plano_monitorizacao", "ID_KPI", "tbl_kpi", "ID_KPI"),
    ("tbl_plano_monitorizacao", "Equipamentos", "tbl_equipamentos", "ID_Equipamento"), ("tbl_equipamentos", "ID_Equipamento", "SGA:tbl_equipamentos", "ID_Equipamento", r"^EQP-"),
    ("tbl_competencias", "Formacao_Obrigatoria", "tbl_formacoes", "ID_Formacao"), ("tbl_competencias", "ID_USE", "tbl_use", "ID_USE"),
    ("tbl_registo_formacao", "ID_Competencia", "tbl_competencias", "ID_Competencia"), ("tbl_registo_formacao", "ID_Formacao", "tbl_formacoes", "ID_Formacao"),
    ("tbl_criterios_operacionais", "ID_USE", "tbl_use", "ID_USE"), ("tbl_rondas_energia", "ID_Criterio", "tbl_criterios_operacionais", "ID_Criterio"),
    ("tbl_projetos", "ID_Alteracao_SGI", "SGA:tbl_alteracoes", "ID_Alteracao"), ("tbl_aquisicoes", "ID_Criterio", "tbl_criterios_aquisicao", "ID_Criterio"),
    ("tbl_legal", "ID_Legal", "SGA:tbl_legal", "ID_Legal", r"^LEG-\d"), ("tbl_obrigacoes", "ID_Legal", "tbl_legal", "ID_Legal"), ("tbl_obrigacoes", "ID_Obrigacao", "SGA:tbl_obrigacoes", "ID_Obrigacao", r"^OBR-\d"),
    ("tbl_desvios", "ID_NC", "tbl_nc", "ID_NC"), ("tbl_poupancas", "ID_MV", "tbl_mv_planos", "ID_MV"),
    ("tbl_constatacoes", "ID_NC", "tbl_nc", "ID_NC"), ("tbl_constatacoes", "ID_Auditoria", "tbl_programa_auditorias", "ID_Auditoria"),
    ("tbl_checklist_aud", "ID_Auditoria", "tbl_programa_auditorias", "ID_Auditoria"), ("tbl_nc", "ID_Origem", "tbl_constatacoes|tbl_desvios", "ID_Constatacao|ID_Desvio"),
    ("tbl_5porques", "ID_NC", "tbl_nc", "ID_NC"),
]
PAT = (r"^(R\d+$|O\d+$|RE-|OE-|PES-|SWT-|TOWS-|CTX-E-|PIE-|PI-|LEG-|OBR|USE-|OPE-|IDE-|KPI-E-|OBJ-E-|PA-E-|PAM-|LBE-|ALE-|DAD-|MEDE-|EQP-|EQE-|COMPE-|FOR-E-|CO-\d|CAQ-|ALT-|AA-|"
       r"AUD-E-|NCE-|DSV-|MV-|CONE-A-|POL-E-|KZ-|ISBM-|IM-|SS-|HF-|INJ$|SOP$|UTL-|GER|SER$|HFS$|FRO$)")

# matriz de requisitos: (ID, cláusula, requisito, informação documentada, classe, documento, evidência, estado, observação, ação)
REQ = [
    ("RQE-01", "4.1", "Questões externas e internas; alterações climáticas pertinentes (Amd 1:2024)", "—", "C", "PR-SGE-01", "RG-SGE-01 tbl_contexto_energia, tbl_clima", "Conforme", "11 questões ligadas ao contexto do SGI", ""),
    ("RQE-02", "4.2", "Partes interessadas, requisitos, requisitos legais e outros (acesso; como se aplicam)", "—", "C", "PR-SGE-01; PR-SGE-10", "RG-SGE-01 tbl_partes_interessadas; RG-SGE-10 tbl_legal", "Conforme", "", ""),
    ("RQE-03", "4.3", "Âmbito e fronteiras; nenhum tipo de energia excluído; autoridade para controlar", "Manter", "A", "MAN-SGE-01", "RG-SGE-01 tbl_ambito, tbl_fronteiras", "Conforme", "Frota e gerador incluídos", ""),
    ("RQE-04", "4.4", "SGE estabelecido, implementado, mantido e melhorado", "—", "C", "MAN-SGE-01", "RG-SGE-00 a 14", "Conforme", "", ""),
    ("RQE-05", "5.1 a)–h)", "Liderança e compromisso (âmbito, política, integração, planos, recursos, comunicação, resultados, melhoria)", "—", "C", "MAN-SGE-01", "RG-SGE-02 tbl_lideranca", "Parcial", "Integração nos processos de negócio parcial (LIDE-03)", "RD-E-26-D05"),
    ("RQE-06", "5.1 i)–m)", "Equipa de gestão de energia; apoio às pessoas; IDE representativos; alterações", "—", "C", "MAN-SGE-01", "RG-SGE-02 tbl_equipa_energia; RG-SGE-05 LBE_Modelo", "Parcial", "LBE-01 válida com reserva; alterações sem energia até 12/2026", "NCE-26-01"),
    ("RQE-07", "5.2", "Política energética a)–g), disponível, comunicada, revista", "Manter (disponível)", "A", "POL-SGE-01", "RG-SGE-02 Politica_Energetica, tbl_verif_politica", "Conforme", "Comunicação reforçada no turno 3", ""),
    ("RQE-08", "5.3", "Funções, responsabilidades e autoridades; equipa de gestão de energia a)–e)", "—", "C", "MAN-SGE-01", "RG-SGE-02 tbl_raci", "Conforme", "RACI com um aprovador por processo", ""),
    ("RQE-09", "6.1", "Riscos e oportunidades; ações; integração; avaliação da eficácia", "—", "C", "PR-SGE-04", "RG-SGE-03", "Parcial", "RE-xx ainda por incluir no RG-SGA-02", "CONE-A-26-07"),
    ("RQE-10", "6.2", "Objetivos e metas energéticas a)–h); planos de ação com método de verificação", "Reter", "B", "PR-SGE-04", "RG-SGE-05 tbl_objetivos, tbl_metas_energeticas, tbl_planos_acao", "Conforme", "6 objetivos, 8 metas, 7 planos", ""),
    ("RQE-11", "6.3", "Revisão energética a)–e); métodos e critérios mantidos; resultados retidos; atualização", "Manter + Reter", "A", "PR-SGE-02", "RG-SGE-04", "Conforme", "Consumos por uso estimados até 15/12/2026", ""),
    ("RQE-12", "6.4", "IDE apropriados; metodologia mantida; variáveis relevantes; valores retidos", "Manter + Reter", "A", "PR-SGE-03", "RG-SGE-05 tbl_kpi, tbl_ide_mensal, Metodologia_IDE_LBE", "Conforme", "", ""),
    ("RQE-13", "6.5", "LBE com período apropriado; normalização; revisão a)–c); modificações retidas", "Reter", "B", "PR-SGE-03", "RG-SGE-05 LBE_Modelo, tbl_lbe, tbl_alteracoes_lbe", "Parcial", "p-valor dos graus-dia 0,27 (CONE-A-26-03)", "RD-E-26-D03"),
    ("RQE-14", "6.6", "Plano de recolha de dados a)–e); exatidão e repetibilidade", "Reter", "B", "PR-SGE-05", "RG-SGE-06", "Conforme", "Submedição em serviço desde 15/12/2026", ""),
    ("RQE-15", "7.1", "Recursos", "—", "C", "—", "RG-SGE-13 tbl_decisoes (recursos)", "Conforme", "", ""),
    ("RQE-16", "7.2", "Competência (reter evidência)", "Reter", "B", "PR-SGE-09", "RG-SGE-07 tbl_competencias, tbl_registo_formacao", "Conforme", "", ""),
    ("RQE-17", "7.3", "Consciencialização a)–d)", "—", "C", "PR-SGE-09", "RG-SGE-07 tbl_consciencializacao", "Conforme", "", ""),
    ("RQE-18", "7.4", "Comunicação a)–e); processo de sugestões", "Considerar reter", "C", "PR-SGE-09", "RG-SGE-07 tbl_comunicacao, tbl_sugestoes", "Conforme", "Prazo de 15 dias definido", ""),
    ("RQE-19", "7.5", "Informação documentada (criar, atualizar, controlar)", "—", "C", "PR-SGA-08 (SGI)", "RG-SGE-00 tbl_matriz_mestra; RG-SGA-09 tbl_lista_mestra", "Parcial", "Documentos do SGE a incluir na lista mestra do SGI", "—"),
    ("RQE-20", "8.1", "Critérios de operação e manutenção dos USE; comunicação; controlo; alterações", "Manter (na medida necessária)", "B", "PR-SGE-06; IT-SGE-01 a 03", "RG-SGE-08", "Conforme", "ANI-02 tratado", ""),
    ("RQE-21", "8.2", "Projeto: oportunidades e controlo operacional; resultados na especificação", "Reter", "B", "PR-SGE-07", "RG-SGE-09 tbl_projetos", "Lacuna", "NC menor NCE-26-01 (linhas novas)", "NCE-26-01"),
    ("RQE-22", "8.3", "Aquisições: critérios ao longo da vida; informar fornecedores; especificação da compra de energia", "—", "C", "PR-SGE-07", "RG-SGE-09 tbl_criterios_aquisicao, tbl_aquisicoes, tbl_lcc, tbl_compra_energia", "Parcial", "NCE-26-02 fechada; eficácia por avaliar", "NCE-26-02"),
    ("RQE-23", "9.1.1", "Monitorização (4 características-chave); IDE vs LBE; desvios significativos investigados", "Reter", "B", "PR-SGE-05; PR-SGE-08", "RG-SGE-05; RG-SGE-06; RG-SGE-11", "Conforme", "Melhoria demonstrada em out–dez/2026", ""),
    ("RQE-24", "9.1.2", "Avaliação da conformidade legal em intervalos planeados", "Reter", "B", "PR-SGE-10", "RG-SGE-10", "Conforme", "REP 2026 entregue", ""),
    ("RQE-25", "9.2", "Auditoria interna (programa, critérios, imparcialidade, relato)", "Reter", "B", "PR-SGE-11", "RG-SGE-12", "Conforme", "AUD-E-2026-01", ""),
    ("RQE-26", "9.3", "Revisão pela gestão (9.3.2, 9.3.3, 9.3.4)", "Reter", "B", "PR-SGE-11", "RG-SGE-13", "Conforme", "RD-E-2026-01", ""),
    ("RQE-27", "10.1", "NC e ação corretiva a)–e)", "Reter", "B", "PR-SGE-12", "RG-SGE-14 tbl_nc, tbl_5porques", "Conforme", "2 NC em curso no prazo", ""),
    ("RQE-28", "10.2", "Melhoria contínua do SGE e demonstração da melhoria do desempenho energético", "—", "C", "PR-SGE-12", "RG-SGE-14 tbl_melhoria; RG-SGE-05 LBE_Modelo", "Parcial", "Demonstrada só no 4.º trimestre de 2026", "RD-E-26-D08"),
]

MESTRA = [
    ("DOCE-01", "Âmbito e fronteiras do SGE", "Documento", "4.3", "A (manter)", GE, DG, "Anual", "Até substituição + 5 anos", "MAN-SGE-01; RG-SGE-01 tbl_ambito"),
    ("DOCE-02", "Política energética", "Documento", "5.2", "A (manter)", DG, DG, "Anual", "Até substituição + 5 anos", "POL-SGE-01; RG-SGE-02"),
    ("DOCE-03", "Métodos e critérios da revisão energética", "Documento", "6.3", "A (manter)", GE, DIND, "Anual", "Até substituição + 5 anos", "PR-SGE-02; RG-SGE-04 Metodo_Criterios"),
    ("DOCE-04", "Metodologia para determinar e atualizar os IDE", "Documento", "6.4", "A (manter)", GE, DIND, "Anual", "Até substituição + 5 anos", "PR-SGE-03; RG-SGE-05 Metodologia_IDE_LBE"),
    ("DOCE-05", "Critérios operacionais dos USE (na medida necessária)", "Documento", "8.1", "A (manter)", GE, DIND, "Anual", "Até substituição + 5 anos", "PR-SGE-06; IT-SGE-01 a 03; RG-SGE-08"),
    ("REGE-01", "Objetivos e metas energéticas", "Registo", "6.2", "B (reter)", GE, DG, "Anual", "5 anos", "RG-SGE-05 tbl_objetivos, tbl_metas_energeticas"),
    ("REGE-02", "Planos de ação", "Registo", "6.2.3", "B (reter)", GE, DIND, "Mensal", "5 anos", "RG-SGE-05 tbl_planos_acao"),
    ("REGE-03", "Resultados da revisão energética", "Registo", "6.3", "B (reter)", GE, DIND, "Anual", "10 anos", "RG-SGE-04"),
    ("REGE-04", "Valores dos IDE", "Registo", "6.4", "B (reter)", GE, "—", "Mensal", "10 anos", "RG-SGE-05 tbl_ide_mensal"),
    ("REGE-05", "LBE, dados das variáveis relevantes e modificações", "Registo", "6.5", "B (reter)", GE, DIND, "Por alteração", "10 anos", "RG-SGE-05 tbl_lbe, tbl_base_energia, tbl_alteracoes_lbe"),
    ("REGE-06", "Dados recolhidos; medição, exatidão e repetibilidade", "Registo", "6.6", "B (reter)", GE, "—", "Contínua", "10 anos", "RG-SGE-06"),
    ("REGE-07", "Evidência de competência", "Registo", "7.2", "B (reter)", RH_, "—", "Por ação", "Contrato + 5 anos", "RG-SGE-07 tbl_registo_formacao"),
    ("REGE-08", "Sugestões de melhoria (considerar reter)", "Registo", "7.4", "C", GE, "—", "Contínua", "5 anos", "RG-SGE-07 tbl_sugestoes"),
    ("REGE-09", "Atividades de projeto relativas ao desempenho energético", "Registo", "8.2", "B (reter)", DIND, "—", "Por projeto", "Vida do equipamento", "RG-SGE-09 tbl_projetos"),
    ("REGE-10", "Resultados de monitorização e medição; investigação de desvios", "Registo", "9.1.1", "B (reter)", GE, "—", "Mensal", "10 anos", "RG-SGE-05; RG-SGE-11"),
    ("REGE-11", "Avaliação da conformidade legal", "Registo", "9.1.2", "B (reter)", GE, "—", "Semestral", "10 anos", "RG-SGE-10"),
    ("REGE-12", "Programa e resultados das auditorias", "Registo", "9.2", "B (reter)", AUD, "—", "Anual", "10 anos", "RG-SGE-12"),
    ("REGE-13", "Resultados da revisão pela gestão", "Registo", "9.3", "B (reter)", DG, DG, "Anual", "10 anos", "RG-SGE-13"),
    ("REGE-14", "NC, ações e resultados", "Registo", "10.1", "B (reter)", GE, "—", "Por ocorrência", "10 anos", "RG-SGE-14"),
]

# matriz normativa (pedido: família ISO 50000 + normas de integração, com o requisito da 50001 que apoiam)
NORMATIVA = [
    ("ISO 50001:2018 + Amd 1:2024 (NP EN ISO 50001:2019)", "Sistemas de gestão da energia — Requisitos com orientação para utilização", "Núcleo (requisitos certificáveis)", "Todos (4 a 10)",
     "Todo o SGE da Plasticom", "RG-SGE-00 a 14; MAN-SGE-01", "Prontidão (Resumo_Prontidao)", "Aplicada", ""),
    ("ISO 50006:2023", "Avaliação do desempenho energético com IDE e LBE", "Orientação — muito importante", "6.4, 6.5, 9.1.1, 10.2", "LBE-01 por regressão, normalização, domínio de validade, ajuste não rotineiro, demonstração da melhoria",
     "RG-SGE-05; PR-SGE-03", "IDE-01 (melhoria normalizada)", "Aplicada", "2.ª edição substitui a ISO 50006:2014 (texto da pasta: NBR ISO 50006:2016 — ficheiro PDF corrompido)"),
    ("ISO 50015:2014", "M&V do desempenho energético das organizações — princípios e orientação", "Orientação — muito importante", "6.2.3, 9.1.1", "Planos de M&V (MV-01 a 04); incerteza",
     "RG-SGE-11; PR-SGE-08", "Poupança verificada (MWh)", "Aplicada", ""),
    ("ISO 50047:2016", "Poupanças de energia — determinação nas organizações", "Orientação", "9.1.1, 10.2", "Poupança = LBE ajustada − real; soma por ação vs instalação", "RG-SGE-11 tbl_poupancas", "MWh poupados", "Aplicada", ""),
    ("ISO 50002-1:2025 (+ -2 edifícios, -3 processos)", "Auditorias energéticas — Parte 1: requisitos gerais", "Apoio direto", "6.3", "Método da revisão energética e da auditoria SGCIE", "RG-SGE-04; PR-SGE-02", "Oportunidades (MWh/ano)",
     "Aplicada", "Substitui a ISO 50002:2014 (a lista do ChatGPT cita a edição antiga)"),
    ("EN 16247-1:2022 (série EN 16247)", "Auditorias energéticas — requisitos gerais (CEN)", "Apoio direto (Europa)", "6.3", "Base europeia das auditorias; serviu de ponto de partida à ISO 50002-1:2025", "RG-SGE-04", "—", "Referência", ""),
    ("ISO 50004:2020", "Orientação para implementar, manter e melhorar um SGE", "Apoio direto", "Todos", "Exemplos para cada requisito (critérios de USE, equipa, controlo operacional)", "PR-SGE-01 a 12", "—", "Aplicada", ""),
    ("ISO 50005:2021", "Orientação para uma implementação faseada (níveis de maturidade)", "Apoio (útil a PME)", "Todos", "Autoavaliação de maturidade do SGE", "RG-SGE-00 Autoavaliacao_Maturidade", "Nível médio", "Aplicada", ""),
    ("ISO/TS 50011:2023", "Avaliar a GESTÃO da energia com base na ISO 50001 (pontuação EMPS)", "Apoio (autoavaliação)", "Todos", "Inspira a autoavaliação (estrutura, operação, metas)", "RG-SGE-00", "EMPS", "Referência",
     "Correção: não é 'avaliar a melhoria do desempenho energético' — avalia o estado da gestão da energia"),
    ("ISO 50003:2021", "Requisitos para organismos que auditam e certificam SGE", "Certificação", "9.1.1, 10.2 (melhoria demonstrada)", "Critério de prontidão para a certificação em 2027", "RG-SGE-05 LBE_Modelo; RG-SGE-12", "Melhoria demonstrada (Sim/Não)", "Aplicada", ""),
    ("ISO 19011 (edição em vigor)", "Diretrizes para auditoria de sistemas de gestão", "Auditoria", "9.2", "Programa, imparcialidade, relato", "RG-SGE-12; PR-SGE-11", "Cumprimento do programa", "Aplicada", ""),
    ("ISO 50009:2021", "Implementação de um SGE comum a várias organizações", "Apoio (não aplicável)", "4.3", "Só uma unidade — não aplicável", "—", "—", "Não aplicável", ""),
    ("ISO/PAS 50010:2023", "Orientação para 'net zero energy' nas operações com um SGE ISO 50001", "Apoio técnico (clima)", "6.2, 8.3 b)", "Visão de longo prazo da UPAC e da eletricidade renovável (OBJ-E-05)", "RG-SGE-05; PL-SGA-01", "IDE-10", "Referência",
     "Correção: não trata de 'desempenho energético e financeiro'"),
    ("ISO 50007:2017", "Serviços de energia — avaliação e melhoria do serviço prestado aos UTILIZADORES", "Apoio (contratos de serviços)", "8.3", "Contratos de manutenção com impacto nos USE", "RG-SGE-09 CAQ-06", "—", "Referência",
     "Correção: é orientado aos utilizadores do serviço, não 'requisitos para provedores'"),
    ("ISO 50046:2019", "Métodos gerais para prever poupanças de energia", "Apoio", "6.3 d)", "Estimativa das poupanças das oportunidades (OPE-xx)", "RG-SGE-04 tbl_oportunidades", "MWh previstos", "Aplicada", ""),
    ("ISO 50014 (inexistente)", "—", "—", "—", "Não existe norma ISO 50014 publicada (verificado em iso.org)", "—", "—", "Não aplicável", "Correção à lista do ChatGPT"),
    ("EN 17267:2019", "Plano de medição e monitorização da energia", "Apoio (medição)", "6.6, 9.1.1", "Árvore de contadores, qualidade dos dados", "RG-SGE-06; PR-SGE-05", "KPI-E-01", "Aplicada", ""),
    ("EN 17463:2021 (VALERI)", "Avaliação de investimentos relacionados com a energia", "Apoio (económico)", "8.2, 8.3", "LCC da IM-002", "RG-SGE-09 tbl_lcc", "LCC anual equivalente", "Aplicada", ""),
    ("ISO 11011:2013", "Ar comprimido — avaliação da eficiência energética", "Técnica (USE-03)", "6.3, 8.1, 9.1.1", "Fugas, pressão, potência específica", "RG-SGE-08 CO-09 a 12; RG-SGE-11", "IDE-08", "Aplicada", ""),
    ("ISO 1217 (anexo E)", "Compressores de deslocamento — ensaios de aceitação", "Técnica (compras)", "8.3", "Potência específica declarada dos compressores", "RG-SGE-09 CAQ-02", "kW/(m³/min)", "Aplicada", ""),
    ("IEC 60034-30-1 / Reg. (UE) 2019/1781", "Classes de eficiência de motores (IE)", "Técnica / legal", "8.3", "Motores IE3/IE4", "RG-SGE-09 CAQ-04; RG-SGE-10 LEG-E-05", "—", "Aplicada", ""),
    ("IEC 62053-22 / IEC 61557-12", "Contadores (classe 0,5S) e analisadores de energia", "Técnica (medição)", "6.6", "Exatidão dos contadores M00–M06 e do PA-01", "RG-SGE-06 tbl_equipamentos", "—", "Aplicada", ""),
    ("EUROMAP 60.1 / 60.2", "Consumo de energia de injetoras e máquinas de sopro", "Técnica (setor)", "8.3", "Declaração do consumo específico nas compras de máquinas", "RG-SGE-09 CAQ-01", "kWh/kg", "Aplicada", ""),
    ("EVO — IPMVP Core Concepts", "Protocolo internacional de M&V (opções A–D)", "Protocolo", "6.2.3, 9.1.1", "Opções A, B e C", "RG-SGE-11", "—", "Aplicada", ""),
    ("ASHRAE Guideline 14", "Medição das poupanças de energia (critérios estatísticos)", "Protocolo", "6.5", "CV(RMSE), NMBE, R²", "RG-SGE-05 LBE_Modelo", "—", "Aplicada", ""),
    ("ISO 14001:2026", "Sistemas de gestão ambiental", "Integração (SGI)", "4.1, 6.1, 9.1.2", "Fonte única de contexto, riscos, requisitos legais e dados de energia", "RG-SGA-00 a 21", "—", "Aplicada", ""),
    ("ISO 9001:2026", "Sistemas de gestão da qualidade", "Integração (SGI)", "7.5, 9.2, 9.3", "Auditorias e revisões integradas; validação do setpoint do chiller", "RG-SGQ-00 a 19", "—", "Aplicada", ""),
    ("ISO 14064-1:2018", "Quantificação e relato de GEE organizacionais", "Integração (clima)", "4.1, 6.2", "Inventário GEE (âmbito 2 depende do SGE)", "RG-SGA-19", "tCO2e", "Aplicada (SGA)", ""),
    ("ISO 14067:2018", "Pegada de carbono de produtos", "Integração (clima)", "4.2", "PCF por SKU pedida pelos clientes (usa kWh por máquina)", "RG-SGA-19; RG-SGE-04 tbl_energia_maquina", "kgCO2e/1.000 un", "Parcial", ""),
    ("ISO 14040 / 14044", "Avaliação do ciclo de vida", "Integração (clima)", "8.2", "Decisões de projeto com ciclo de vida (SGA)", "RG-SGA-20", "—", "Referência", ""),
    ("ISO 14068-1:2023", "Gestão para a neutralidade carbónica", "Integração (clima)", "6.2", "Enquadramento das metas de descarbonização (PL-SGA-01)", "RG-SGA-19", "—", "Referência", ""),
    ("ISO 45001:2018", "Segurança e saúde no trabalho", "Integração (futura)", "8.1", "Segurança nas intervenções em quadros elétricos e ar comprimido", "—", "—", "Não implementada", "Próximo subsistema do SGI"),
    ("ISO 55001:2024", "Gestão de ativos", "Integração", "8.1, 8.2, 8.3", "Plano de substituição das máquinas hidráulicas (OPE-06, OPE-14) com LCC", "RG-SGE-09", "—", "Referência", ""),
    ("ISO 31000:2018", "Gestão do risco — diretrizes", "Integração", "6.1", "Método dos riscos RE-xx", "RG-SGE-03", "—", "Aplicada", ""),
]

NORMAS = [
    ("F-01", "NBR ISO 50001:2018 (texto da norma) — apostila da pasta de interpretação", "ISO 50001 - Sistema de Gestão da Energia - Interpretação/Norma-ISO-500012018-Apostila...pdf", "Texto dos requisitos 4 a 10 usado nas matrizes e nos procedimentos", "Requisito"),
    ("F-02", "ISO 50001:2018/Amd 1:2024 — Climate action changes", "https://www.iso.org/standard/88430.html", "4.1 e 4.2 (alterações climáticas) — RG-SGE-01 tbl_clima", "Requisito"),
    ("F-03", "ISO 50006:2023 (prévia oficial ISO/FDIS: índice, termos 3.1.1–3.1.18, secções 4–5)", "https://www.iso.org/standard/79367.html; https://standards.iteh.ai/catalog/standards/iso/a20f6ab7-5e8e-4f5e-bc3a-2bdda723d80e/iso-50006-2023", "Estrutura do RG-SGE-05 e do PR-SGE-03 (fronteiras, utilizadores, tipos de IDE, período de referência, normalização, fatores estáticos, demonstração)", "Orientação"),
    ("F-04", "Módulo M-9 'Medição e indicadores de desempenho energético — ISO 50006' (curso Bureau Veritas)", "Pasta de interpretação", "Plano de medição (o quê, porquê, como, frequência, valor esperado, desvio, ação) e tipos de IDE", "Orientação"),
    ("F-05", "Módulos M-2 a M-11 do curso de interpretação ISO 50001 (Bureau Veritas) e manuais de SGE da pasta", "Pasta de interpretação", "Interpretação de cada secção nos procedimentos PR-SGE", "Orientação"),
    ("F-06", "ISO 50002-1:2025 — Energy audits, Part 1", "https://www.iso.org/standard/83645.html", "Revisão energética (RG-SGE-04)", "Orientação"),
    ("F-07", "ISO 50003:2021", "https://www.iso.org/standard/77575.html", "Melhoria demonstrada para certificar; competência dos auditores", "Outro requisito"),
    ("F-08", "ISO 50015:2014", "https://www.iso.org/standard/60043.html", "Planos de M&V (RG-SGE-11)", "Orientação"),
    ("F-09", "ISO 50046:2019", "https://www.iso.org/standard/67790.html", "Previsão das poupanças (RG-SGE-04)", "Orientação"),
    ("F-10", "ISO/TS 50011:2023", "https://www.iso.org/standard/81286.html", "Autoavaliação da gestão da energia (EMPS)", "Orientação"),
    ("F-11", "ISO/PAS 50010:2023", "https://www.iso.org/standard/51873.html", "Visão net zero energy (OBJ-E-05)", "Orientação"),
    ("F-12", "LBNL — EnPI Lite 'Valid model requirements' e 50001 Ready Navigator (tarefa 11)", "https://enpilite.lbl.gov/valid-model-requirements; https://navigator.lbl.gov/guidance/task/11", "Testes estatísticos da LBE-01; passos de IDE/LBE", "Orientação"),
    ("F-13", "ASHRAE Guideline 14 (CV(RMSE) ≤ 15%, NMBE ±5% em dados mensais)", "https://www.ashrae.org", "Testes da LBE-01", "Orientação"),
    ("F-14", "EVO — IPMVP Core Concepts", "https://evo-world.org", "Opções de M&V", "Orientação"),
    ("F-15", "EN 17267:2019 — Energy measurement and monitoring plan", "https://standards.iteh.ai/catalog/standards/cen/a4a4eeca-3fee-406b-bd85-95d07aae506d/en-17267-2019", "RG-SGE-06", "Orientação"),
    ("F-16", "Kent, R. — Energy Management in Plastics Processing (BPF / Tangram)", "https://tangram.co.uk/wp-content/uploads/BPF-Energy-Management-in-Plastics-Processing.pdf", "Benchmarks do setor (carga de base 20–40%, fugas 20–40%, +1 °C ≈ −3% no chiller, standby)", "Orientação"),
    ("F-17", "UNIDO — Practical guide / EnMS toolkit", "https://decarbonization.unido.org/resources/practical-guide-for-implementing-an-energy-management-system/", "Estrutura dos registos (oportunidades, IDE)", "Orientação"),
    ("F-18", "SGCIE — perguntas frequentes e conversor; DGEG", "https://sgcie.pt/faq/; https://www.dgeg.gov.pt", "RG-SGE-10 (limiares, metas do PREn, REP)", "Requisito legal"),
    ("F-19", "Decreto-Lei 71/2008 (SGCIE); Despacho 17313/2008; DL 68-A/2015", "https://diariodarepublica.pt", "RG-SGE-10", "Requisito legal"),
    ("F-20", "Diretiva (UE) 2023/1791 e análise DNV", "https://eur-lex.europa.eu/eli/dir/2023/1791/oj; https://www.dnv.com/assurance/Management-Systems/eu-directive-2023-1791-energy-efficiency-iso-50001/", "Art. 11.º (85 TJ / 10 TJ) — LEG-E-03, OE-02", "Requisito legal"),
    ("F-21", "Decreto-Lei 130/2026 (altera o DL 15/2022 — SEN; autoconsumo)", "https://diariodarepublica.pt/dr/detalhe/decreto-lei/130-2026-1139897687", "LEG-E-04 (UPAC)", "Requisito legal"),
    ("F-22", "Registos do SGA (RG-SGA-01, 02, 04, 10, 11, 13, 17, 18, 19) e do SGQ (RG-SGQ-07, 11, 13, 16)", "SGI_Sistema-de-gestao-integrado", "Fontes únicas de contexto, riscos, legal, energia mensal, fatores de emissão, alterações", "Dados"),
    ("F-23", "Dataset do projeto (fact_production, dim_machine_profile)", "datasets/silver; datasets/dim", "Produção, horas e paragens por máquina; ano de instalação", "Dados"),
]

FONTE_UNICA = [
    ("Contexto do SGI (PESTEL, SWOT, TOWS)", "RG-SGA-01", "tbl_pestel; tbl_swot; tbl_tows", "RG-SGE-01 tbl_contexto_energia (ID_Contexto_SGI)", "Questões só de energia com prefixo CTX-E"),
    ("Riscos e oportunidades do SGI", "RG-SGA-02", "tbRiscos; tbOportunidades", "RG-SGE-03 (R22, R25, R34, R35, R39, O6, O15)", "RE-xx / OE-xx propostos para inclusão"),
    ("Requisitos legais do SGI", "RG-SGA-04", "tbl_legal; tbl_obrigacoes", "RG-SGE-10 (LEG-08, OBR-06)", "LEG-E-xx propostos para inclusão"),
    ("Eletricidade mensal por uso mar/2025–dez/2026", "RG-SGA-13", "tbl_dados_ambientais", "RG-SGE-04/05/06/09 (valores lidos por envdata.ambiente())", "Fonte única completa: o RG-SGA-13 inclui o fecho de 2026 (set–dez) com o efeito das ações de energia"),
    ("Parâmetros do rateio (KW_*, AR_FRAC, FRIO_*)", "RG-SGA-13", "tbl_parametros", "RG-SGE-04 tbl_campanha (confirma os kW)", "A campanha sugere rever AR_FRAC após os VSD"),
    ("Equipamentos EQP-01 a EQP-05", "RG-SGA-13", "tbl_equipamentos", "RG-SGE-06 tbl_equipamentos (cópia com Registo_Dono)", "EQE-xx novos do SGE"),
    ("Fatores de emissão e gasóleo (âmbito 1)", "RG-SGA-19", "tbl_fatores_emissao; tbl_inventario_gee", "RG-SGE-04 tbl_tipos_energia; tbl_consumo_uso", "Fatores SGCIE (Despacho 17313/2008) só no SGE"),
    ("Planeamento de alterações", "RG-SGA-18", "tbl_alteracoes", "RG-SGE-09 tbl_projetos; RG-SGE-05 tbl_alteracoes_lbe", "Secção de energia na checklist (RD-E-26-D05)"),
    ("Fornecedor de eletricidade", "RG-SGA-11", "tbl_outros_fornecedores (ENE-01)", "RG-SGE-01 PIE-03; RG-SGE-09 tbl_compra_energia", ""),
    ("IDE, LBE, objetivos e metas energéticas", "RG-SGE-05", "tbl_kpi; tbl_lbe; tbl_objetivos", "RG-SGA-05 OBJ-01/KPI-01 (mesma intenção; SGA mantém para o relato ambiental)", "O SGE é dono da medição normalizada"),
    ("Revisão energética, USE e oportunidades", "RG-SGE-04", "tbl_usos; tbl_use; tbl_oportunidades", "RG-SGA-03 (aspetos de energia AA-007, 010, 021, 030–032)", ""),
    ("Energia por máquina (campanha e rateio)", "RG-SGE-04", "tbl_campanha; tbl_energia_maquina", "RG-SGA-19 (PCF por SKU — futuro)", ""),
    ("Faturas, tarifas e potência", "RG-SGE-09", "tbl_faturas; tbl_precos", "RG-SGE-04/05 (preço médio); RG-SGA-17 (prm_Preco_Eletricidade)", "Atualizar o parâmetro do SGA com o preço de 2026"),
    ("Trajetória SGCIE (PREn)", "RG-SGE-10", "tbl_sgcie", "RG-SGA-04 LEG-08", ""),
    ("Pessoas e operadores", "RG-SGQ-07 / dataset", "tbl_pessoas", "RG-SGE-07 tbl_registo_formacao", ""),
]

PARES = [
    ("HE-01", "Partes interessadas", "SGE-01", "tbl_partes_interessadas", "SGA-01", "tbl_partes_interessadas"),
    ("HE-02", "Política", "SGE-02", "tbl_politica", "SGA-05", "tbl_politica"),
    ("HE-03", "RACI", "SGE-02", "tbl_raci", "SGA-08", "tbl_raci"),
    ("HE-04", "Catálogo de indicadores", "SGE-05", "tbl_kpi", "SGA-05", "tbl_kpi"),
    ("HE-05", "Objetivos", "SGE-05", "tbl_objetivos", "SGA-05", "tbl_objetivos"),
    ("HE-06", "Parâmetros", "SGE-05", "tbl_parametros", "SGA-13", "tbl_parametros"),
    ("HE-07", "Plano de monitorização", "SGE-06", "tbl_plano_monitorizacao", "SGA-13", "tbl_plano_monitorizacao"),
    ("HE-08", "Equipamentos de medição", "SGE-06", "tbl_equipamentos", "SGA-13", "tbl_equipamentos"),
    ("HE-09", "Competências", "SGE-07", "tbl_competencias", "SGA-08", "tbl_competencias"),
    ("HE-10", "Registo de formação", "SGE-07", "tbl_registo_formacao", "SGA-08", "tbl_registo_formacao"),
    ("HE-11", "Comunicação", "SGE-07", "tbl_comunicacao", "SGA-09", "tbl_comunicacao"),
    ("HE-12", "Requisitos legais", "SGE-10", "tbl_legal", "SGA-04", "tbl_legal"),
    ("HE-13", "Obrigações", "SGE-10", "tbl_obrigacoes", "SGA-04", "tbl_obrigacoes"),
    ("HE-14", "Programa de auditorias", "SGE-12", "tbl_programa_auditorias", "SGA-14", "tbl_programa_auditorias"),
    ("HE-15", "Constatações", "SGE-12", "tbl_constatacoes", "SGA-14", "tbl_constatacoes"),
    ("HE-16", "Agenda da revisão", "SGE-13", "tbl_agenda", "SGA-15", "tbl_agenda"),
    ("HE-17", "Decisões da revisão", "SGE-13", "tbl_decisoes", "SGA-15", "tbl_decisoes"),
    ("HE-18", "Não conformidades", "SGE-14", "tbl_nc", "SGA-07", "tbl_nc"),
    ("HE-19", "5 Porquês", "SGE-14", "tbl_5porques", "SGA-07", "tbl_5porques"),
    ("HE-20", "Matriz de requisitos", "SGE-00", "tbl_matriz_iso", "SGA-00", "tbl_matriz_iso"),
    ("HE-21", "Fonte única", "SGE-00", "tbl_fonte_unica", "SGA-00", "tbl_fonte_unica"),
    ("HE-22", "Clima", "SGE-01", "tbl_clima", "SGQ-01", "tbl_clima"),
    ("HE-23", "Âmbito", "SGE-01", "tbl_ambito", "SGQ-01", "tbl_ambito"),
    ("HE-24", "Liderança", "SGE-02", "tbl_lideranca", "SGQ-03", "tbl_lideranca"),
    ("HE-25", "Consciencialização", "SGE-07", "tbl_consciencializacao", "SGQ-07", "tbl_consciencializacao"),
    ("HE-26", "Sugestões", "SGE-07", "tbl_sugestoes", "SGQ-19", "tbl_sugestoes"),
    ("HE-27", "Auditores", "SGE-12", "tbl_auditores", "SGQ-16", "tbl_auditores"),
    ("HE-28", "Checklist de auditoria", "SGE-12", "tbl_checklist_aud", "SGQ-16", "tbl_checklist_aud"),
    ("HE-29", "Ações anteriores da revisão", "SGE-13", "tbl_acoes_anteriores", "SGQ-17", "tbl_acoes_anteriores"),
]
EQUIV = {("tbl_partes_interessadas", "Tratado_pelo_SGE"): ("Torna_se_Obrigacao", "Tratado_pelo_SGQ"), ("tbl_kpi", "Tipo_IDE_ISO50006"): ("Tipo_ISO14031",),
         ("tbl_raci", "Processo_SGE"): ("Processo_SGA", "Processo_SGQ"), ("tbl_raci", "Gestor_Energia"): ("Gestor_SGA",), ("tbl_competencias", "ID_USE"): ("ID_Aspeto_AAS",),
         ("tbl_competencias", "USE_Descricao"): ("Aspeto_Significativo",), ("tbl_legal", "IDs_USE"): ("IDs_Aspetos",), ("tbl_legal", "ID_Acao"): ("ID_PAM",),
         ("tbl_decisoes", "Tipo_Saida_9_3_4"): ("Tipo_Saida_9_3_3",), ("tbl_matriz_iso", "Alterado_Amd1_2024"): ("Novo_ou_Alterado_2026",)}

MATURIDADE = [
    ("MAT-01", "Contexto, âmbito e fronteiras (4)", 3, "Documentado e ligado ao contexto do SGI; clima considerado.", "Rever com a transposição da Diretiva 2023/1791."),
    ("MAT-02", "Liderança e equipa (5)", 3, "Equipa nomeada, política aprovada, revisão pela gestão realizada.", "Integrar energia nos investimentos (RD-E-26-D05)."),
    ("MAT-03", "Revisão energética e USE (6.3)", 3, "USE e oportunidades quantificados; consumos por uso ainda estimados.", "Refazer com 12 meses de submedição (2027)."),
    ("MAT-04", "IDE e LBE (6.4, 6.5)", 3, "Modelo estatístico com testes e domínio; melhoria demonstrada no 4.º trimestre.", "LBE por USE medidas; LBE-01b."),
    ("MAT-05", "Recolha de dados e medição (6.6, 9.1)", 3, "Árvore de contadores com 95% medido desde 15/12/2026.", "Dados de 15 min no BI e alarmes."),
    ("MAT-06", "Controlo operacional (8.1)", 3, "16 critérios e rondas semanais com tendência positiva.", "Automatizar alarmes de desvio por USE."),
    ("MAT-07", "Projeto e aquisições (8.2, 8.3)", 2, "Critérios e LCC definidos, mas NC em projeto e serviços em 2026.", "Aplicar a checklist de energia a 100% das alterações."),
    ("MAT-08", "Competência e comunicação (7)", 3, "Formação com eficácia e sugestões com prazo.", "Substituto do Gestor de Energia (RE-05)."),
    ("MAT-09", "Avaliação, auditoria e melhoria (9, 10)", 3, "Auditoria, revisão e NC com ciclo completo; M&V coerente.", "Demonstrar melhoria anual em 2027."),
]


def load_tables(folder, extra=()):
    out, cat = {}, []
    files = sorted(glob.glob(os.path.join(folder, "SGE-[0-9][0-9]_*.xlsx"))) + list(extra)
    for f in files:
        base = os.path.basename(f)
        if base.startswith("SGE-00"):
            continue
        wb = openpyxl.load_workbook(f, data_only=True)
        for ws in wb.worksheets:
            for name in ws.tables:
                t = ws.tables[name]
                rows = list(ws[t.ref])
                hdr = [c.value for c in rows[0]]
                df = pd.DataFrame([[c.value for c in r] for r in rows[1:]], columns=hdr)
                out[name if base.startswith("SGE-") else f"SGA:{name}"] = df
                if base.startswith("SGE-"):
                    hr = rows[0][0].row
                    cat.append(dict(Ficheiro=base, Folha=ws.title, Tabela=name, Linha_Cabecalho=hr, N_Linhas=int(df.iloc[:, 0].notna().sum()), N_Colunas=len(hdr),
                                    Chave_Primaria=hdr[0], Leitura_pandas=f'pd.read_excel("{base}", sheet_name="{ws.title}", header={hr - 1})'))
    return out, cat


def split_ids(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return []
    return [x.strip() for x in re.split(r"[;,]", str(v)) if x.strip() and x.strip() not in ("—", "-")]


def integridade(T):
    res = []
    for rel in REL:
        src, colname, dst, dcol = rel[:4]
        pat = rel[4] if len(rel) > 4 else PAT
        if src not in T or colname not in T[src].columns:
            res.append(dict(Tabela_Origem=src, Coluna=colname, Tabela_Destino=dst.replace("|", " / ").replace("SGA:", "SGA · "), N_Referencias=0, N_Invalidas=0, IDs_Invalidos="tabela/coluna não encontrada", Resultado="Rever"))
            continue
        valid = set()
        for tt, cc in zip(dst.split("|"), dcol.split("|")):
            if tt in T and cc in T[tt].columns:
                valid |= set(T[tt][cc].dropna().astype(str))
        refs = [i for v in T[src][colname] for i in split_ids(v)]
        refs = [r for r in refs if re.match(pat, r)]
        missing = sorted({r for r in refs if r not in valid})
        res.append(dict(Tabela_Origem=src, Coluna=colname, Tabela_Destino=dst.replace("|", " / ").replace("SGA:", "SGA · "), N_Referencias=len(refs), N_Invalidas=len(missing),
                        IDs_Invalidos="; ".join(missing[:25]) or None, Resultado="OK" if not missing else "Rever"))
    return res


def table_headers(path):
    wb = openpyxl.load_workbook(path)
    return {n: [c.name for c in ws.tables[n].tableColumns] for ws in wb.worksheets for n in ws.tables}


def harmonizacao(folder, own):
    files = {os.path.basename(f)[:6]: f for f in glob.glob(os.path.join(folder, "SGE-[0-9][0-9]_*.xlsx")) + glob.glob(os.path.join(SGA_REG, "SGA-[0-9][0-9]_*.xlsx"))
             + glob.glob(os.path.join(SGQ_REG, "SGQ-[0-9][0-9]_*.xlsx"))}
    cache = {"SGE-00": own}
    det, res = [], []
    for pid, assunto, re_, te, ro, to in PARES:
        for r in (re_, ro):
            if r not in cache:
                cache[r] = table_headers(files[r])
        ce, co = cache[re_].get(te, []), cache[ro].get(to, [])
        usados = set()
        for i, c in enumerate(ce, 1):
            alvo, tipo = None, "Só SGE"
            if c in co:
                alvo, tipo = c, "Igual"
            else:
                for eq in EQUIV.get((te, c), ()):
                    if eq in co and eq not in ce:
                        alvo, tipo = eq, "Equivalente (renomear na consolidação)"
                        break
            if alvo:
                usados.add(alvo)
            det.append(dict(ID_Par=pid, Tabela_SGE=te, Coluna_SGE=c, Posicao_SGE=i, Sistema_Padrao=ro[:3], Tabela_Padrao=to, Coluna_Padrao=alvo,
                            Posicao_Padrao=(co.index(alvo) + 1) if alvo else None, Correspondencia=tipo))
        for j, c in enumerate(co, 1):
            if c not in usados:
                det.append(dict(ID_Par=pid, Tabela_SGE=te, Coluna_SGE=None, Posicao_SGE=None, Sistema_Padrao=ro[:3], Tabela_Padrao=to, Coluna_Padrao=c, Posicao_Padrao=j,
                                Correspondencia=f"Só {ro[:3]}"))
        res.append(dict(ID_Par=pid, Assunto=assunto, Registo_SGE="RG-" + re_, Tabela_SGE=te, Registo_Padrao="RG-" + ro, Tabela_Padrao=to, N_Colunas_SGE=len(ce), N_Colunas_Padrao=len(co)))
    return res, det


def build(out, folder):
    extra = [os.path.join(SGA_REG, f) for f in ("SGA-01_Contexto_SWOT_PESTEL.xlsx", "SGA-02_Gestao_Riscos_Oportunidades.xlsx", "SGA-03_Aspetos_Impactes_Ambientais.xlsx",
                                                "SGA-04_Requisitos_Legais_Conformidade.xlsx", "SGA-06_Plano_Acoes_Melhoria_PAM.xlsx", "SGA-13_Monitorizacao_Medicao_Desempenho.xlsx",
                                                "SGA-16_Melhoria_Kaizen_EMAS.xlsx", "SGA-18_Planeamento_Alteracoes.xlsx")]
    T, cat = load_tables(folder, extra)
    b = Book("RG-SGE-00", "Índice do SGE, Matriz ISO 50001, Matriz Normativa e Modelo de Dados",
             activities="Ponto de entrada do SGE: onde está cada requisito, que normas e fontes foram usadas, que tabelas existem, como se ligam entre si e ao SGA/SGQ, e o estado de prontidão para a certificação.",
             clauses="7.5 (informação documentada 'manter' / 'reter'); 4.4; 9.1 (dados para análise); preparação da certificação (ISO 50003:2021).",
             purpose="Catálogo dos 14 registos e tabelas; matriz requisito → documento → evidência → estado com prontidão por capítulo; matriz mestra de informação documentada; matriz normativa "
                     "(família ISO 50000 e normas de integração, com correções à lista recebida); todas as fontes citadas; fonte única SGE ↔ SGA ↔ SGQ; harmonização coluna a coluna; "
                     "integridade referencial; autoavaliação de maturidade (ISO 50005 / ISO/TS 50011).",
             guidance=[("ISO 50001:2018 7.5.1 e anexo A", "A norma exige 'manter' (documentos) e 'reter' (registos) só em pontos definidos; o resto é determinado pela organização."),
                       ("ISO 50005:2021 e ISO/TS 50011:2023", "Níveis de maturidade e avaliação da gestão da energia (folha Autoavaliacao_Maturidade)."),
                       ("Pedido do utilizador (lista de normas relacionadas, 30/09/2026)", "Matriz normativa com número, título, relação, requisito que apoia, aplicação, evidência e KPI.")],
             legal=[("DL 71/2008; Despacho 17313/2008; DL 68-A/2015; Diretiva (UE) 2023/1791; DL 15/2022 + DL 130/2026", "Lista completa com avaliação no RG-SGE-10.")])
    b.add_list("Estado", ESTADO)
    b.add_list("Classe", ["A", "B", "C"])
    b.add_list("SimNao", SIMNAO)

    fcols = [col("Ficheiro", 46, desc="Ficheiro."), col("Codigo", 10, key="PK", desc="Código do registo."), col("Conteudo", 90, desc="Conteúdo."), col("Clausulas", 18, desc="Cláusulas ISO 50001.")]
    b.table("Ficheiros", "tbl_ficheiros", fcols, [dict(Ficheiro=f, Codigo=c, Conteudo=t, Clausulas=cl) for f, c, t, cl in FILES], "Lista dos registos do SGE.",
            title="ÍNDICE DO SISTEMA DE REGISTOS DO SGE — PLASTICOM (ISO 50001:2018 + Amd 1:2024; ISO 50006:2023)",
            subtitle=f"Pasta Registos_SGE_Plasticom · 14 registos + este índice · Procedimentos em Documentos_SGE_Plasticom · Data de referência {DATA_REF:%d/%m/%Y}", row_height=30, tab_color="1F4E5F")

    rcols = [col("ID_Requisito", 7, key="PK", desc="Requisito."), col("Clausula", 10, desc="Cláusula."),
             col("Grupo", 6, f='=IF(@Clausula@="","",IF(LEFT(@Clausula@,2)="10","10",LEFT(@Clausula@,1)))', desc="Capítulo 4–10."),
             col("Requisito", 50, desc="Requisito (resumo)."),
             col("Alterado_Amd1_2024", 8, f='=IF(@ID_Requisito@="","",IF(OR(@Clausula@="4.1",@Clausula@="4.2"),"Sim","Não"))', desc="Alterado pela Amd 1:2024 (clima) — posição de Novo_ou_Alterado_2026."),
             col("Informacao_Documentada", 20, desc="Manter / reter / não exigida."), col("Documento", 18, desc="Documento controlado."), col("Evidencia_Registo_Tabelas", 44, desc="Evidência."),
             col("Estado", 9, dv="Estado", desc="Conforme · Parcial · Lacuna (31/12/2026)."), col("Observacao", 40, desc="Observação.", req=False), col("Acao", 14, desc="Ação.", req=False),
             col("Pontos", 6, "num", f='=IF(@Estado@="","",IF(@Estado@="Conforme",1,IF(@Estado@="Parcial",0.5,0)))', desc="1 · 0,5 · 0."),
             col("Classe", 6, dv="Classe", desc="[Só SGE/SGQ] A manter · B reter · C organização.")]
    b.table("Matriz_ISO50001", "tbl_matriz_iso", rcols, rows_from(["ID_Requisito", "Clausula", "Requisito", "Informacao_Documentada", "Classe", "Documento", "Evidencia_Registo_Tabelas", "Estado", "Observacao", "Acao"], REQ),
            "Requisito → documento → evidência → estado (mesma estrutura do RG-SGA-00 / RG-SGQ-00).", title="MATRIZ DE REQUISITOS ISO 50001:2018 + Amd 1:2024 — DOCUMENTO, EVIDÊNCIA E ESTADO",
            subtitle="Verificação de 31/12/2026 · Lacuna = NC em aberto · Parcial = existe com falha de evidência ou eficácia",
            cf=[("Estado", {"Conforme": "green", "Parcial": "orange", "Lacuna": "red"}), ("Classe", {"A": "blue", "B": "purple"})], row_height=32, freeze_col=2)
    ws = b.sheet("Resumo_Prontidao", "Índice de prontidão por capítulo (calculado a partir da matriz).", tab_color="C00000")
    title(ws, "PRONTIDÃO PARA A CERTIFICAÇÃO ISO 50001 — RESUMO POR CAPÍTULO — calculado", "Índice = (Conforme + 0,5 × Parcial) ÷ requisitos")
    header_row(ws, 3, ["Capítulo", "Tema", "Requisitos", "Conforme", "Parcial", "Lacuna", "Índice"], widths=[10, 26, 11, 10, 10, 10, 12])
    R = lambda c_: f"tbl_matriz_iso[{c_}]"
    temas = [("4", "Contexto"), ("5", "Liderança"), ("6", "Planeamento"), ("7", "Suporte"), ("8", "Operação"), ("9", "Avaliação do desempenho"), ("10", "Melhoria")]
    for k, (g, t_) in enumerate(temas):
        r = 4 + k
        cell(ws, r, 1, g)
        cell(ws, r, 2, t_)
        cell(ws, r, 3, f'=COUNTIF({R("Grupo")},A{r})', fmt="0")
        for j, e in enumerate(("Conforme", "Parcial", "Lacuna")):
            cell(ws, r, 4 + j, f'=COUNTIFS({R("Grupo")},A{r},{R("Estado")},"{e}")', fmt="0")
        cell(ws, r, 7, f'=IFERROR(SUMIFS({R("Pontos")},{R("Grupo")},A{r})/C{r},"")', fmt="0%")
    rt = 4 + len(temas)
    cell(ws, rt, 1, "Total", bold=True)
    for c_ in range(3, 7):
        L = get_column_letter(c_)
        cell(ws, rt, c_, f"=SUM({L}4:{L}{rt - 1})", fmt="0", bold=True)
    cell(ws, rt, 7, f'=IFERROR(SUM({R("Pontos")})/C{rt},"")', fmt="0%", bold=True)
    ws.conditional_formatting.add(f"G4:G{rt}", FormulaRule(formula=["G4<0.75"], fill=PatternFill("solid", fgColor=CF_COLORS["orange"][0])))

    mcols = [col("Codigo", 8, key="PK", desc="Código."), col("Documento_Registo", 44, desc="Informação documentada exigida pela ISO 50001."), col("Tipo", 10, desc="Documento / registo."),
             col("Requisito", 8, desc="Cláusula."), col("Obrigatoriedade", 10, desc="A manter / B reter / C organização."), col("Responsavel", 22, desc="Responsável."), col("Aprovacao", 18, desc="Aprovação."),
             col("Frequencia", 12, desc="Frequência."), col("Retencao", 18, desc="Retenção."), col("Onde_no_SGE", 40, desc="Onde está.")]
    b.table("Matriz_Mestra_Info_Doc", "tbl_matriz_mestra", mcols, rows_from(input_names(mcols), MESTRA),
            "Informação documentada exigida pela ISO 50001 ('manter' e 'reter') e onde está.", title="MATRIZ MESTRA DE INFORMAÇÃO DOCUMENTADA — ISO 50001",
            cf=[("Obrigatoriedade", {"A": "blue", "B": "purple"})], row_height=24)

    ncols = [col("Norma", 34, key="PK", desc="Norma / protocolo."), col("Titulo", 50, desc="Título."), col("Relacao_ISO50001", 22, desc="Relação com a ISO 50001."),
             col("Requisitos_50001_Apoiados", 18, desc="Requisitos da ISO 50001 que apoia."), col("Aplicacao_Plasticom", 50, desc="Como foi aplicada na Plasticom."),
             col("Evidencia", 30, desc="Registo / documento."), col("KPI", 18, desc="Indicador associado."), col("Estado_Uso", 12, desc="Aplicada / referência / não aplicável."),
             col("Nota_Correcao", 44, desc="Nota (inclui correções verificadas em iso.org).", req=False)]
    b.table("Matriz_Normativa", "tbl_matriz_normativa", ncols, [dict(zip(input_names(ncols), n)) for n in NORMATIVA],
            "Matriz normativa do SGE: família ISO 50000, normas técnicas e de integração no SGI.", title="MATRIZ NORMATIVA DO SGE — FAMÍLIA ISO 50000 E INTEGRAÇÃO",
            subtitle="A ISO 50001 é o núcleo certificável; as restantes são orientação, protocolos ou sistemas integrados · Títulos e edições verificados em iso.org (30/09/2026)",
            cf=[("Estado_Uso", {"Aplicada": "green", "Referência": "blue", "Não": "gray", "Parcial": "orange"}), ("Nota_Correcao", {"Correção": "yellow"})], row_height=40, freeze_col=1)
    b.table("Normas_Fontes", "tbl_fontes", [col("ID", 6, key="PK"), col("Referencia", 56, desc="Norma / fonte."), col("Onde", 60, desc="Localização / URL."),
                                             col("Aplicada_em", 56, desc="Onde foi aplicada."), col("Tipo", 14, desc="Requisito / orientação / dados.")],
            [dict(zip(["ID", "Referencia", "Onde", "Aplicada_em", "Tipo"], n)) for n in NORMAS], "Todas as normas, requisitos e fontes consultadas.",
            title="NORMAS, REQUISITOS E FONTES CITADAS", subtitle="As orientações (ISO 5000x, EN, IPMVP, ASHRAE, BPF) não são requisitos auditáveis da ISO 50001 — foram usadas como boas práticas", row_height=30)
    b.table("Matriz_Fonte_Unica", "tbl_fonte_unica", [col("Assunto", 40, key="PK"), col("Registo_Dono", 12), col("Tabelas_Donas", 32), col("So_Referenciado_Em", 50), col("Regra_Nota", 44, req=False)],
            [dict(zip(["Assunto", "Registo_Dono", "Tabelas_Donas", "So_Referenciado_Em", "Regra_Nota"], f)) for f in FONTE_UNICA],
            "Fonte única do SGI (SGE ↔ SGA ↔ SGQ): cada assunto é escrito num só registo.", title="MATRIZ DE FONTE ÚNICA — SGE ↔ SGA ↔ SGQ", row_height=30)

    ccols = [col("Ficheiro", 44), col("Folha", 26), col("Tabela", 26, key="PK"), col("Linha_Cabecalho", 8, "int"), col("N_Linhas", 8, "int"), col("N_Colunas", 8, "int"),
             col("Chave_Primaria", 18), col("Leitura_pandas", 80)]
    b.table("Catalogo_Tabelas", "tbl_catalogo", ccols, cat, "Catálogo de todas as tabelas de dados do SGE (gerado automaticamente).")
    icols = [col("Tabela_Origem", 24), col("Coluna", 18), col("Tabela_Destino", 44), col("N_Referencias", 10, "int"), col("N_Invalidas", 9, "int"), col("IDs_Invalidos", 50, req=False), col("Resultado", 9)]
    b.table("Integridade_Referencial", "tbl_integridade", icols, integridade(T), "Verificação das ligações entre registos do SGE e com o SGA (instantâneo da construção).",
            cf=[("Resultado", {"Rever": "red", "OK": "green"})])

    b.table("Dim_Uso", "tbl_dim_uso", [col("Codigo", 9, key="PK"), col("Uso", 60), col("Tipo_Energia", 18), col("Processo_SGA", 8), col("Equipamentos", 26), col("Variavel_Candidata", 36)],
            [dict(zip(["Codigo", "Uso", "Tipo_Energia", "Processo_SGA", "Equipamentos", "Variavel_Candidata"], u)) for u in USOS], "Dimensão uso de energia (códigos do RG-SGA-13).")
    b.table("Dim_Funcao", "tbl_dim_funcao", [col("ID_Funcao", 8, key="PK"), col("Funcao", 50), col("Area", 8)], [dict(ID_Funcao=a, Funcao=f, Area=p) for a, f, p in FUNCOES], "Dimensão função.")
    b.table("Dim_Clausula_ISO50001", "tbl_dim_clausula", [col("Clausula", 8, key="PK"), col("Titulo", 60), col("Informacao_Documentada", 44), col("Alterado_Amd1_2024", 10)],
            [dict(zip(["Clausula", "Titulo", "Informacao_Documentada", "Alterado_Amd1_2024"], c)) for c in CLAUSULAS], "Cláusulas da ISO 50001:2018 + Amd 1:2024.", row_height=20)
    b.table("Dim_Maquina", "tbl_dim_maquina", [col("ID_Maquina", 10, key="PK"), col("Processo", 8), col("Ano_Instalacao", 9, "int"), col("Tecnologia", 36)],
            [dict(zip(["ID_Maquina", "Processo", "Ano_Instalacao", "Tecnologia"], m)) for m in MAQUINAS], "Dimensão máquina (dataset + tecnologia simulada).")

    b.add_list("Nivel", [1, 2, 3, 4])
    acols = [col("ID", 7, key="PK"), col("Elemento", 36), col("Nivel_1a4", 8, "int", dv="Nivel", desc="1 inicial · 2 em desenvolvimento · 3 implementado · 4 otimizado (ISO 50005)."),
             col("Evidencia", 60), col("Proximo_Passo", 50)]
    b.table("Autoavaliacao_Maturidade", "tbl_maturidade", acols, [dict(zip(["ID", "Elemento", "Nivel_1a4", "Evidencia", "Proximo_Passo"], m)) for m in MATURIDADE],
            "Autoavaliação de maturidade do SGE (inspirada na ISO 50005 e na ISO/TS 50011).", title="AUTOAVALIAÇÃO DE MATURIDADE DO SGE (ISO 50005 / ISO/TS 50011)",
            subtitle="Nível 3 = implementado e eficaz na maioria dos requisitos · Autoavaliação da equipa de gestão de energia em 31/12/2026", row_height=30)
    ws = b.wb["Autoavaliacao_Maturidade"]
    tm = b.tables["tbl_maturidade"]
    cell(ws, tm["last"] + 2, 2, "Nível médio", bold=True)
    cell(ws, tm["last"] + 2, 3, "=AVERAGE(tbl_maturidade[Nivel_1a4])", fmt="0.0", bold=True)

    own = {n: list(b.tables[n]["colmap"]) for n in ("tbl_ficheiros", "tbl_matriz_iso", "tbl_fonte_unica")}
    hres, hdet = harmonizacao(folder, own)
    b.add_list("Correspondencia", ["Igual", "Equivalente (renomear na consolidação)", "Só SGE", "Só SGA", "Só SGQ"])
    H = lambda c_: f"tbl_harmonizacao[{c_}]"
    hcols = [col("ID_Par", 7, key="PK"), col("Assunto", 28), col("Registo_SGE", 10), col("Tabela_SGE", 22), col("Registo_Padrao", 10, desc="Registo do SGA (padrão do SGI) ou do SGQ."),
             col("Tabela_Padrao", 22), col("Mesmo_Nome_Tabela", 9, f='=IF(@ID_Par@="","",IF(@Tabela_SGE@=@Tabela_Padrao@,"Sim","Não"))'),
             col("N_Colunas_SGE", 8, "int"), col("N_Colunas_Padrao", 8, "int"),
             col("N_Iguais", 8, "int", f=f'=COUNTIFS({H("ID_Par")},@ID_Par@,{H("Correspondencia")},"Igual")'),
             col("N_Equivalentes", 9, "int", f=f'=COUNTIFS({H("ID_Par")},@ID_Par@,{H("Correspondencia")},"Equivalente*")'),
             col("Pct_Comum_Padrao", 9, "pct", f='=IF(@ID_Par@="","",IFERROR((@N_Iguais@+@N_Equivalentes@)/@N_Colunas_Padrao@,""))', desc="% das colunas do padrão presentes no SGE.")]
    b.table("Harmonizacao_Resumo", "tbl_harmonizacao_resumo", hcols, hres, "Resumo da normalização SGE ↔ SGA/SGQ por par de tabelas.",
            title="HARMONIZAÇÃO SGE ↔ SGA / SGQ — RESUMO POR TABELA", subtitle="Padrão do SGI = SGA (e SGQ nos temas que o SGA não tem) · Consolidar juntando tabelas homónimas e renomeando as 'Equivalentes'",
            cf=[("Mesmo_Nome_Tabela", {"Sim": "green", "Não": "orange"})], row_height=18, freeze_col=2)
    dcols = [col("ID_Par", 7), col("Tabela_SGE", 22), col("Coluna_SGE", 26, req=False), col("Posicao_SGE", 7, "int", req=False), col("Sistema_Padrao", 7), col("Tabela_Padrao", 22),
             col("Coluna_Padrao", 26, req=False), col("Posicao_Padrao", 7, "int", req=False), col("Correspondencia", 22, dv="Correspondencia"),
             col("Nome_Consolidado", 26, f='=IF(@ID_Par@="","",IF(@Coluna_Padrao@<>"",@Coluna_Padrao@,@Coluna_SGE@))')]
    b.table("Matriz_Harmonizacao_SGI", "tbl_harmonizacao", dcols, hdet, "Correspondência coluna a coluna SGE ↔ SGA/SGQ (gerada a partir dos ficheiros).",
            title="MATRIZ DE HARMONIZAÇÃO — COLUNA A COLUNA", cf=[("Correspondencia", {"Igual": "green", "Equivalente": "yellow", "Só SGE": "blue", "Só SGA": "gray", "Só SGQ": "gray"})], row_height=15, freeze_col=3)

    ws = b.sheet("Modelo_Dados", "Descrição do modelo de dados e convenções de IDs.")
    ws.column_dimensions["A"].width = 26
    title(ws, "MODELO DE DADOS DO SGE")
    notes(ws, 3, [
        ("Estrutura", "Modelo em estrela: factos mensais (tbl_base_energia, tbl_ide_mensal, tbl_consumo_uso, tbl_energia_maquina, tbl_faturas, tbl_rondas_energia) ligados às dimensões mês, uso, máquina, "
                      "função e cláusula, e aos registos de gestão (IDE, LBE, objetivos, planos, riscos, NC) por IDs estáveis."),
        ("Convenção de IDs", "CTX-E / PIE (contexto) · POL-E / LIDE / EGE / PE (liderança) · RE / OE (riscos) · TE / USE / OPE (revisão energética) · IDE / KPI-E / LBE / ALE / FEST / VAR / OBJ-E / MET-E / PA-E "
                             "(desempenho) · DAD / MEDE / M00–M10 / EQE (dados) · COMPE / FOR-E / REGE / CONE / COME / SUGE (pessoas) · CO / RE-26 (operação) · PRJ-E / CAQ / AQE / LCC / FT (compras) · "
                             "LEG-E / OBRE (legal) · DSV / MV / POU (desvios e M&V) · AUDE / AUD-E / CHKE / CONE-A (auditoria) · RD-E (revisão) · NCE / MEL (melhoria)."),
        ("Ligação ao SGA e SGQ", "Contexto, riscos, legal, energia mensal, equipamentos EQP, alterações e kaizen são lidos dos registos do SGA (prefixo SGA: na integridade). As tabelas com assunto igual "
                                 "têm o mesmo nome e colunas (Matriz_Harmonizacao_SGI)."),
        ("Casos de análise / ML", "1) Previsão do consumo mensal (modelo da LBE como baseline e ML para resíduos). 2) Deteção de anomalias em dados de 15 min dos analisadores. "
                                  "3) Consumo por SKU a partir de tbl_energia_maquina (PCF). 4) Otimização do setpoint do chiller com a temperatura."),
        ("Atualização", "python _build/build_all.py reconstrói tudo e recalcula no Excel; python _build/build_docs.py gera os procedimentos Word."),
        ("Localização (Excel PT)", "As fórmulas usam números de série de data e funções clássicas (LINEST, TDIST, FDIST, TINV) que funcionam em Excel português."),
    ])
    return b.save(out)


if __name__ == "__main__":
    import sys
    print(build(sys.argv[1], sys.argv[2]))
