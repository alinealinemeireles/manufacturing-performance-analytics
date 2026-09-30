"""RG-SGQ-04 — Riscos e oportunidades da qualidade (ISO 9001:2026 6.1.1, 6.1.2, 6.1.3; 9.1.3 e/f; 9.3.2 g/h).

Fonte única: o registo corporativo de riscos e oportunidades do SGI é o RG-SGA-02 (mantido à mão). Este registo
é uma VISTA GERADA dos riscos/oportunidades que afetam a conformidade do produto e a satisfação do cliente,
com os campos que a ISO 9001:2026 acrescenta (efeito na conformidade, disrupção, opção de tratamento da Nota 2,
integração nos processos e avaliação da eficácia separada para riscos e para oportunidades).
Inclui a PFMEA de processo com a ocorrência calculada a partir das inspeções por atributos do dataset.
"""
import os
import openpyxl
import pandas as pd
from sgqlib import *
from dimsq import *
import qdata as Q

SGA02 = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "SGA_Sistema-de-gestao-ambiental", "Registos_SGA_Plasticom",
                                     "SGA-02_Gestao_Riscos_Oportunidades.xlsx"))

OPCOES = ["Evitar o risco", "Assumir o risco para perseguir uma oportunidade", "Eliminar a fonte de risco", "Alterar a probabilidade",
          "Alterar as consequências", "Partilhar o risco", "Manter o risco por decisão informada"]

# riscos da qualidade: ID → (efeito na conformidade/satisfação, disrupção (Nota 1), opção de tratamento, processo SGQ, integração no processo (6.1.2 b1),
#                            método de avaliação da eficácia (6.1.2 b2), KPI, resultado da eficácia, ligações SGQ)
RQ = {
    "R1": ("Frascos com espessura fora de especificação e lotes rejeitados", "Não", "Alterar a probabilidade", "SOP", "PC-SOP-01 rev. 04: SPC de espessura por cavidade; manutenção do perfil de aquecimento",
           "Taxa de rejeição ISBM-003 ≤ 2,5% durante 3 meses", "KPI-Q-01", "Por avaliar", "CAPA-Q-26-04"),
    "R2": ("Flash e fugas em frascos do molde M-SOP-007", "Não", "Eliminar a fonte de risco", "MAN", "Reforma do molde e manutenção por contador de ciclos (PFMEA RPN 280 → 120)",
           "Leakage em M-SOP-007 < 300 ppm após reforma", "KPI-Q-01", "Eficaz", "MOC-Q-26-07"),
    "R3": ("Paragens da ISBM-005 atrasam entregas", "Sim", "Alterar a probabilidade", "MAN", "Plano preventivo hidráulico por horas; peças críticas em stock",
           "MTBF ISBM-005 ≥ 12 h", "KPI-Q-07", "Não eficaz", "RG-SGQ-02 tbl_maquinas"),
    "R4": ("Short shot e variação de peso em tampas da IM-002", "Não", "Eliminar a fonte de risco", "INJ", "Nova janela de parâmetros validada por DOE (MOC-Q-26-05)",
           "Taxa de defeito IM-002 ≤ 0,38% (meta DMAIC)", "KPI-Q-01", "Eficaz", "MOC-Q-26-05; PRJ-Q-01"),
    "R5": ("Flash em tampas em campanhas longas", "Não", "Alterar a probabilidade", "INJ", "Limpeza de respiros a cada 50 000 ciclos no MES", "Flash IM-004 < 4 000 ppm", "KPI-Q-01", "Por avaliar", ""),
    "R6": ("Falha de aderência da decoração (reclamação de cliente)", "Não", "Alterar a probabilidade", "SER", "Ensaio de aderência no arranque e a meio do lote (PC-DEC-01 rev. 03)",
           "Zero lotes com adesão NC em 3 meses", "KPI-Q-02", "Por avaliar", "CAPA-Q-26-03"),
    "R7": ("Defeitos e atrasos na serigrafia SS-001", "Sim", "Alterar a probabilidade", "MAN", "Plano de rolete/rasqueta por contador; backlog de manutenção a zero", "MTTR SS-001 ≤ 1,8 h", "KPI-Q-07", "Por avaliar", ""),
    "R8": ("Transferência incompleta de foil", "Não", "Alterar a probabilidade", "HFS", "Registo da velocidade real e limite de 60 golpes/min", "Foil Transfer < 250 ppm", "KPI-Q-01", "Por avaliar", ""),
    "R9": ("Mais defeitos no turno 2 em todos os processos", "Não", "Alterar a probabilidade", "PCP", "Checklist de passagem de turno; reforço de inspetor no turno 2",
           "Diferença de rejeição turno 2 vs. média ≤ 0,2 pp", "KPI-Q-01", "Por avaliar", "RG-SGQ-03 GW-26-02"),
    "R10": ("Variação de peso atribuída ao operador OP-INJ-003", "Não", "Alterar a probabilidade", "RH", "Retreino e validação no posto (RG-SGQ-07)", "Desvio-padrão do operador ≤ 1,3× pares", "KPI-Q-18", "Por avaliar", ""),
    "R11": ("Pontos negros (contaminação) em frascos — reclamação crítica", "Não", "Partilhar o risco", "CMP", "Certificado de análise por lote de masterbatch e cláusula de qualidade no contrato SUP-008",
            "Black Specks < 8 000 ppm", "KPI-Q-01", "Por avaliar", "RG-SGQ-12"),
    "R12": ("Defeitos elevados nas primeiras semanas dos produtos novos", "Não", "Alterar a probabilidade", "RD", "Lançamento com plano de controlo reforçado (safe launch) 6 semanas",
            "FPY dos SKU novos ≥ 90% à 6.ª semana", "KPI-Q-02", "Parcial", "DD-26-01 a DD-26-06"),
    "R13": ("Lotes de resina fora de especificação chegam à produção", "Sim", "Evitar o risco", "CMP", "SUP-005 em aprovação condicional; 100% dos lotes com ensaio de MFI e humidade",
            "Lotes aceites SUP-005 ≥ 90%", "KPI-Q-08", "Não eficaz", "RG-SGQ-12; CAPA-Q-26-05"),
    "R14": ("Variação de qualidade e falta de PCR/rPET", "Sim", "Partilhar o risco", "CMP", "Segundo fornecedor de HDPE-PCR em qualificação", "Dois fornecedores aprovados por material PCR", "KPI-Q-08", "Por avaliar", ""),
    "R15": ("Falta de MP interrompe ordens e entregas", "Sim", "Alterar as consequências", "PCP", "Stock de segurança de 10 dias para resinas críticas; aviso ao cliente (8.2.1 e)",
            "Zero paragens por falta de MP em 3 meses", "KPI-Q-20", "Por avaliar", "RG-SGQ-08 COMC-06"),
    "R16": ("Lote marginal passa na amostragem AQL e gera reclamação", "Não", "Alterar a probabilidade", "LAB", "Inspeção reforçada (ISO 2859-1:2026) após 2 lotes rejeitados em 5; ensaio de fuga 100% nas linhas com leak tester",
            "PPM de cliente ≤ 150", "KPI-Q-03", "Parcial", "RG-SGQ-13"),
    "R17": ("Capacidade marginal: variação dimensional leva a NC e reclamação", "Não", "Alterar a probabilidade", "LAB", "Plano Cpk: características críticas com Cpk ≥ 1,33 até 2027",
            "% grupos críticos com Cpk ≥ 1,33", "KPI-Q-13", "Por avaliar", "OBJ-Q-04"),
    "R18": ("Rastreabilidade incompleta impede contenção rápida", "Sim", "Alterar as consequências", "TI", "Rastreabilidade lote → ordem → cliente no ERP; exercício de recolha semestral",
            "Tempo de rastreio ≤ 4 h nos exercícios", "KPI-Q-02", "Parcial", "RG-SGQ-13 tbl_rastreio"),
    "R19": ("Parâmetros não registados: deteção tardia de desvios", "Não", "Alterar a probabilidade", "TI", "Recolha automática de parâmetros no MES (O3)", "% máquinas com parâmetros registados", "KPI-Q-07", "Por avaliar", ""),
    "R20": ("CAPA vencidas e ineficazes: recorrência das NC", "Não", "Eliminar a fonte de risco", "QUA", "Revisão semanal de CAPA; 8D para NC maiores; verificação de eficácia a 90 dias",
            "CAPA no prazo ≥ 85% e eficácia ≥ 80%", "KPI-Q-10; KPI-Q-11", "Não eficaz", "RG-SGQ-18"),
    "R21": ("Reclamações de serviço (atraso, quantidade, produto trocado)", "Sim", "Alterar a probabilidade", "EXP", "Leitura de código de barras na carga; contagem por balança; confirmação de datas",
            "Reclamações de serviço ≤ 5/mês", "KPI-Q-05", "Por avaliar", "CAPA-Q-26-02"),
    "R23": ("Falta de operadores reduz capacidade e competência no posto", "Sim", "Alterar as consequências", "RH", "Polivalência mínima de 2 pessoas por posto crítico", "Cobertura da matriz de polivalência ≥ 80%", "KPI-Q-18", "Por avaliar", ""),
    "R24": ("Custo da não qualidade acima da meta", "Não", "Alterar a probabilidade", "GES", "COQ mensal na revisão pela gestão", "COQ falhas ≤ 3% das vendas", "KPI-Q-16", "Por avaliar", "RG-SGQ-19"),
    "R26": ("Dados corrompidos levam a decisões erradas sobre produto", "Sim", "Alterar as consequências", "TI", "Cópias de segurança diárias e testes de pipeline", "Incidentes de dados = 0", "—", "Eficaz", ""),
    "R33": ("Portfólio não cumpre requisitos PPWR e de clientes europeus", "Não", "Assumir o risco para perseguir uma oportunidade", "RD", "Dossier técnico PPWR como saída obrigatória do D&D", "Famílias com dossier PPWR", "—", "Por avaliar", "DD-26-01"),
    "R34": ("Calor e seca afetam o arrefecimento e a continuidade", "Sim", "Alterar as consequências", "MAN", "Plano de verão: chiller de reserva, janela de processo de calor validada", "Zero lotes NC por temperatura no verão", "KPI-Q-01", "Parcial", "RG-SGQ-02 AMB-01"),
    "R36": ("Resina spot fora de especificação e declarações PCR não verificadas", "Sim", "Evitar o risco", "CMP", "Sem novas compras spot sem aprovação da Qualidade", "Compras spot aprovadas = 100%", "KPI-Q-08", "Por avaliar", ""),
}
OQ = {
    "O1": ("Menos defeitos na frota de injeção", "INJ", "Replicar a janela validada do DOE às IM-001/003–008", "Taxa de defeito média da injeção −30%", "KPI-Q-01", "Por avaliar", "MOC-Q-26-05"),
    "O2": ("Menos flash/fugas por manutenção de moldes baseada em ciclos", "MAN", "Contadores de ciclos no MES", "Defeitos por molde após reforma", "KPI-Q-01", "Por avaliar", ""),
    "O3": ("Deteção precoce de desvios com painéis em tempo quase real", "TI", "Painel de SPC e KPI no Power BI com alertas", "Tempo entre desvio e reação < 1 turno", "KPI-Q-01", "Por avaliar", ""),
    "O4": ("Mais disponibilidade = entregas no prazo", "PCP", "Ataque às 3 maiores causas de paragem", "Disponibilidade 87,4% → 90%", "KPI-Q-07", "Por avaliar", ""),
    "O5": ("Diferenciação com clientes europeus (design for recycling)", "RD", "Critérios de reciclabilidade nas entradas de D&D", "N.º de projetos ganhos com argumento de circularidade", "—", "Por avaliar", "DD-26-01"),
    "O8": ("Fornecedores mais fiáveis através de scorecard", "CMP", "Scorecard mensal e reavaliação semestral (RG-SGQ-12)", "Lotes aceites ≥ 95%", "KPI-Q-08", "Parcial", "RG-SGQ-12"),
    "O9": ("Mais flexibilidade para lotes pequenos alimentar/farma", "PCP", "SMED nas trocas de molde", "Tempo médio de troca −25%", "KPI-Q-07", "Por avaliar", "PRJ-Q-03"),
    "O11": ("Eliminar refações escondidas (fábrica escondida)", "QUA", "Registar retrabalho como NC e atacar as causas", "Custo de refação −50%", "KPI-Q-16", "Por avaliar", ""),
    "O12": ("Contenção rápida e menor custo de reclamação", "TI", "Rastreabilidade ponta a ponta", "Tempo de rastreio ≤ 4 h", "KPI-Q-03", "Por avaliar", "RG-SGQ-13"),
    "O13": ("Deteção a 100% de defeitos de decoração", "LAB", "Visão artificial na HF-001 (investimento aprovado)", "Reclamações de decoração −50%", "KPI-Q-04", "Por avaliar", "PRJ-Q-04"),
    "O14": ("Cumprir o pico set–nov sem perder qualidade", "PCP", "Plano de capacidade e reforço de inspetores", "Rejeição set–nov ≤ média anual", "KPI-Q-01", "Por avaliar", ""),
}

S_CLASSE = {"Critical": 9, "Major": 7, "Minor": 4}
S_ESPECIAL = {"Leakage": 10, "Sealing": 10, "Tamper Band Separation": 10, "Migration Test (Food Contact)": 10, "Drop Test": 8}
D_METODO = {"Visual": 7, "Leak Test": 3, "Go/No-Go": 4, "Pull Test": 5, "Drop Test": 5, "Cross Hatch": 6, "MEK Rub Test": 6, "Colorimeter": 5, "Visual (vision)": 3}
MODO = {"Leakage": "Frasco com fuga", "Flash": "Rebarba no plano de junta", "Bubbles": "Bolhas na parede", "Black Specks": "Pontos negros (contaminação)",
        "Stain": "Mancha", "Transparency": "Falta de transparência", "Color": "Cor fora do padrão", "Burr": "Rebarba de corte", "Sealing": "Tampa não veda",
        "Short Shot": "Peça incompleta", "Thread": "Rosca deformada", "Tamper Band Separation": "Anel de inviolabilidade não separa/separa antes da abertura",
        "Drop Test": "Rotura em queda", "Registration": "Desalinhamento da impressão", "Centering": "Decoração descentrada", "Coverage": "Cobertura insuficiente",
        "Smudge": "Borrão", "Pinholes": "Microfuros na tinta", "Adhesion": "Tinta destaca", "Cure": "Tinta mal curada", "Foil Transfer": "Foil não transfere",
        "Foil Adhesion": "Foil destaca", "Edge Definition": "Contorno mal definido", "Rub Resistance": "Decoração sai por fricção"}
EFEITO = {"Critical": "Perda do conteúdo / segurança do consumidor; reclamação crítica e possível recolha", "Major": "Falha funcional ou estética grave; rejeição pelo cliente na linha de enchimento",
          "Minor": "Defeito estético; reclamação menor"}
CAUSA = {"Leakage": "Desgaste do molde / parâmetros de sopro", "Flash": "Força de fecho insuficiente, desgaste do plano de junta", "Bubbles": "Humidade da resina",
         "Black Specks": "Contaminação do masterbatch / degradação no cilindro", "Stain": "Óleo/limpeza do molde", "Transparency": "Resina húmida ou mal seca",
         "Color": "Dosagem de masterbatch", "Burr": "Faca de corte gasta", "Sealing": "Variação dimensional do vedante", "Short Shot": "Temperatura / velocidade de injeção (IM-002)",
         "Thread": "Arrefecimento insuficiente", "Tamper Band Separation": "Pontes do anel mal dimensionadas na ferramenta nova (TE-012)", "Drop Test": "Espessura baixa na base",
         "Registration": "Folga do suporte do ecrã", "Centering": "Posicionamento do frasco", "Coverage": "Pressão da rasqueta / viscosidade da tinta", "Smudge": "Velocidade excessiva",
         "Pinholes": "Ecrã sujo", "Adhesion": "Tratamento de superfície insuficiente / resina marginal", "Cure": "Potência UV degradada", "Foil Transfer": "Velocidade alta / temperatura baixa",
         "Foil Adhesion": "Temperatura do cunho", "Edge Definition": "Pressão do cunho", "Rub Resistance": "Foil inadequado ao substrato"}
METODO_CP = {"Leakage": "Leak Test", "Sealing": "Leak Test", "Thread": "Go/No-Go", "Tamper Band Separation": "Pull Test", "Drop Test": "Drop Test",
             "Adhesion": "Cross Hatch", "Cure": "MEK Rub Test", "Color": "Visual", "Registration": "Visual", "Centering": "Visual"}


def sga02():
    wb = openpyxl.load_workbook(SGA02, data_only=True)
    out = {}
    for n in ("tbRiscos", "tbOportunidades"):
        for ws in wb.worksheets:
            if n in ws.tables:
                t = ws.tables[n]
                rows = list(ws[t.ref])
                df = pd.DataFrame([[c.value for c in r] for r in rows[1:]], columns=[c.value for c in rows[0]])
                out[n] = df[df["ID"].notna()].set_index("ID")
    return out["tbRiscos"], out["tbOportunidades"]


def pfmea_rows():
    auto = set(Q.rd("dim_machine_profile.csv", Q.DIM).query("HasAutomatedDefectDetection")["MachineId"])
    rows = []
    for f in ("fact_bottle_attribute_inspection_cq_processed.csv", "fact_cap_attribute_inspection_cq_processed.csv", "fact_ink_attribute_inspection_cq_processed.csv"):
        d = Q.periodo(Q.rd(f), "ProductionDate")
        d["Proc"] = d["MachineId"].str.extract(r"^([A-Z]+)")[0].map({"IM": "INJ", "ISBM": "SOP", "SS": "SER", "HF": "HFS"})
        d["Auto"] = d["MachineId"].isin(auto)
        g = d.groupby(["Proc", "Characteristic", "Class"]).agg(Inspecionadas=("SampleSize", "sum"), Defeitos=("DefectsFound", "sum"),
                                                               Lotes=("LotId", "nunique"), Pct_Auto=("Auto", "mean")).reset_index()
        rows.append(g)
    g = pd.concat(rows).groupby(["Proc", "Characteristic", "Class"], as_index=False).agg(Inspecionadas=("Inspecionadas", "sum"), Defeitos=("Defeitos", "sum"),
                                                                                         Lotes=("Lotes", "sum"), Pct_Auto=("Pct_Auto", "mean"))
    out = []
    for k, r in enumerate(g.sort_values(["Proc", "Class", "Characteristic"]).itertuples(), 1):
        met = METODO_CP.get(r.Characteristic, "Visual")
        dval = D_METODO[met]
        if met == "Visual" and r.Pct_Auto >= 0.5:
            met, dval = "Visual (vision)", 3
        out.append(dict(ID_FMEA=f"PFMEA-{k:02d}", Processo=r.Proc, Caracteristica=Q.T_CARACT.get(r.Characteristic, {"Edge Definition": "Definição de contornos", "Rub Resistance": "Resistência à fricção"}.get(r.Characteristic, r.Characteristic)),
                        Classe=Q.T_CLASSE[r.Class], Modo_Falha=MODO.get(r.Characteristic, r.Characteristic), Efeito=EFEITO[r.Class],
                        S=S_ESPECIAL.get(r.Characteristic, S_CLASSE[r.Class]), Causa=CAUSA.get(r.Characteristic, "A investigar"),
                        Unidades_Inspecionadas=int(r.Inspecionadas), Defeitos=int(r.Defeitos), Controlo_Detecao=met, D=dval))
    return out


def build(out):
    R, O = sga02()
    b = Book("RG-SGQ-04", "Riscos e Oportunidades da Qualidade (vista do registo SGI) e PFMEA",
             activities="Determinar, analisar e avaliar riscos (6.1.2) e oportunidades (6.1.3) que afetam a conformidade dos produtos e a satisfação do cliente; planear ações proporcionais, integrá-las nos processos e avaliar a sua eficácia.",
             clauses="6.1.1 Determinação de riscos e oportunidades; 6.1.2 Ações para tratar riscos (Nota 1 disrupção; Nota 2 opções); 6.1.3 Ações para tratar oportunidades (novo em 2026); 9.1.3 e) f); 9.3.2 g) h); 4.4.1 g)",
             purpose="Vista da qualidade do registo corporativo de riscos do SGI (RG-SGA-02, fonte única): os campos de identificação, pontuação, dono e estado são COPIADOS na geração (não editar aqui); os campos da ISO 9001:2026 (efeito na conformidade, disrupção, opção de tratamento, integração no processo, método e resultado da eficácia) são registados aqui. Inclui a PFMEA de processo com a ocorrência calculada pelas inspeções reais.",
             links=[("RG-SGA-02 (fonte)", "tbRiscos / tbOportunidades — alterar pontuação, dono e estado só no RG-SGA-02 e regenerar."),
                    ("RG-SGQ-05 KPI", "KPI usado para avaliar a eficácia."), ("RG-SGQ-18 / RG-SGQ-06", "CAPA e alterações que implementam as ações.")],
             guidance=[("ISO/TC 176 APG — Risk-based thinking", "Não é exigida gestão de risco formal (ISO 31000), mas as ações devem ser proporcionais e a eficácia avaliada: colunas Proporcionalidade e Resultado_Eficacia."),
                       ("ISO 31000:2018 / IEC 31010:2019 (pasta gestão de riscos)", "Escala P×I do registo SGI; PFMEA como técnica de identificação de riscos de processo."),
                       ("AIAG-VDA FMEA / Academy cap. 81–82", "S, O, D e RPN; ocorrência derivada de ppm observado; prioridade de ação (AP simplificada)."),
                       ("ICH Q9(R1) (pasta de interpretação)", "Formalidade proporcional ao risco — usado para os produtos farmacêuticos (anel de inviolabilidade, S = 10).")])
    b.add_list("SimNao", ["Sim", "Não"])
    b.add_list("Opcao", OPCOES)
    b.add_list("Processo", PROC_CODES)
    b.add_list("Eficacia", ["Eficaz", "Parcial", "Não eficaz", "Por avaliar"])
    b.add_list("Classe", ["Crítica", "Maior", "Menor"])

    cp = lambda s, k: s.get(k) if k in s.index else None
    rrows = []
    for rid, m in RQ.items():
        s = R.loc[rid]
        rrows.append(dict(ID_Risco=rid, Area_SGI=cp(s, "Área / processo responsável"), Descricao=cp(s, "Descrição do risco"), Causa=cp(s, "Causa"),
                          Score_Inerente=cp(s, "Score inerente"), Score_Atual=cp(s, "Score atual"), Nivel_Atual=cp(s, "Nível atual"), Dono=cp(s, "Dono do risco"),
                          Acao_Adotada_SGI=cp(s, "Ação adotada"), Estado_SGI=cp(s, "Status atual"),
                          Efeito_Conformidade_Satisfacao=m[0], Disrupcao=m[1], Opcao_Tratamento=m[2], Processo_SGQ=m[3], Integracao_Processo=m[4],
                          Metodo_Eficacia=m[5], KPI=m[6], Resultado_Eficacia=m[7], Data_Avaliacao=DATA_REF if m[7] != "Por avaliar" else None, Ligacoes_SGQ=m[8] or None))
    gen = "Copiado do RG-SGA-02 na geração (não editar aqui)."
    rcols = [
        col("ID_Risco", 7, key="PK / FK → RG-SGA-02", desc="ID no registo corporativo."),
        col("Area_SGI", 16, desc=gen), col("Descricao", 42, desc=gen), col("Causa", 28, desc=gen, req=False),
        col("Score_Inerente", 8, "num0", desc=gen), col("Score_Atual", 8, "num0", desc=gen), col("Nivel_Atual", 8, desc=gen), col("Dono", 20, desc=gen),
        col("Acao_Adotada_SGI", 12, desc=gen, req=False), col("Estado_SGI", 11, desc=gen, req=False),
        col("Efeito_Conformidade_Satisfacao", 34, desc="Efeito indesejado na capacidade de fornecer produto conforme e de aumentar a satisfação (6.1.2)."),
        col("Disrupcao", 9, dv="SimNao", desc="Risco relacionado com disrupção — capacidade de fornecer durante/após uma disrupção (6.1.2 Nota 1)."),
        col("Opcao_Tratamento", 24, dv="Opcao", desc="Opção de tratamento (6.1.2 Nota 2)."),
        col("Processo_SGQ", 7, dv="Processo", desc="Processo onde a ação é integrada.", key="FK → dim processo"),
        col("Integracao_Processo", 40, desc="Como a ação é integrada e implementada no processo (6.1.2 b1)."),
        col("Metodo_Eficacia", 30, desc="Como se avalia a eficácia (6.1.2 b2)."), col("KPI", 12, desc="Indicador.", key="FK → RG-SGQ-05"),
        col("Resultado_Eficacia", 11, dv="Eficacia", desc="Resultado da avaliação da eficácia (9.1.3 e)."), col("Data_Avaliacao", 11, "date", desc="Data da avaliação.", req=False),
        col("Ligacoes_SGQ", 18, desc="CAPA, MOC, projeto ou registo SGQ.", req=False),
        col("Proporcionalidade", 12, f='=IF(@ID_Risco@="","",IF(AND(@Score_Atual@>=12,@Opcao_Tratamento@="Manter o risco por decisão informada"),"Rever",IF(AND(@Score_Atual@<=4,@Opcao_Tratamento@="Evitar o risco"),"Excessiva?","Adequada")))',
            desc="Ações proporcionais ao impacto potencial (6.1.2): risco alto mantido → rever; risco baixo evitado → avaliar custo."),
        col("Alerta_Eficacia", 14, f='=IF(@Resultado_Eficacia@="Não eficaz","Replanear ação",IF(AND(@Resultado_Eficacia@="Por avaliar",@Nivel_Atual@="Alto"),"Avaliar já","—"))',
            desc="Sinal para a revisão pela gestão (9.3.2 g)."),
    ]
    b.table("Riscos_Qualidade", "tbl_riscos_q", rcols, rrows,
            "Riscos que afetam a conformidade e a satisfação (6.1.2) — vista do RG-SGA-02 com os campos da ISO 9001:2026.",
            title="RISCOS DA QUALIDADE — AÇÕES E EFICÁCIA (6.1.2)",
            subtitle="Colunas 2–10 copiadas do RG-SGA-02 (fonte única) · Colunas seguintes: requisitos ISO 9001:2026 · Proporcionalidade e alerta calculados",
            cf=[("Nivel_Atual", {"Alto": "red", "Médio": "yellow", "Baixo": "green"}), ("Resultado_Eficacia", {"Não eficaz": "red", "Parcial": "orange", "Eficaz": "green"}),
                ("Disrupcao", {"Sim": "purple"}), ("Alerta_Eficacia", {"Replanear": "red", "Avaliar": "orange"}), ("Proporcionalidade", {"Rever": "red", "Excessiva": "yellow"})],
            row_height=62, freeze_col=1)

    orows = []
    for oid, m in OQ.items():
        s = O.loc[oid]
        orows.append(dict(ID_Oport=oid, Area_SGI=cp(s, "Área / processo responsável"), Descricao=cp(s, "Descrição da oportunidade"), Beneficio_SGI=cp(s, "Benefício esperado"),
                          Score=cp(s, "Score"), Nivel=cp(s, "Nível (relevância)"), Decisao_SGI=cp(s, "Decisão adotada"), Dono=cp(s, "Dono da oportunidade"),
                          Efeito_Desejado=m[0], Processo_SGQ=m[1], Integracao_Processo=m[2], Metodo_Eficacia=m[3], KPI=m[4], Resultado_Eficacia=m[5],
                          Data_Avaliacao=DATA_REF if m[5] != "Por avaliar" else None, Ligacoes_SGQ=m[6] or None))
    ocols = [col("ID_Oport", 7, key="PK / FK → RG-SGA-02", desc="ID no registo corporativo."), col("Area_SGI", 16, desc=gen), col("Descricao", 42, desc=gen),
             col("Beneficio_SGI", 30, desc=gen, req=False), col("Score", 7, "num0", desc=gen), col("Nivel", 9, desc=gen), col("Decisao_SGI", 12, desc=gen, req=False), col("Dono", 20, desc=gen),
             col("Efeito_Desejado", 34, desc="Efeito desejado na conformidade e na satisfação do cliente (6.1.3)."),
             col("Processo_SGQ", 7, dv="Processo", desc="Processo onde a ação é integrada."), col("Integracao_Processo", 40, desc="6.1.3 b1)."),
             col("Metodo_Eficacia", 30, desc="6.1.3 b2)."), col("KPI", 12, desc="Indicador.", key="FK → RG-SGQ-05"),
             col("Resultado_Eficacia", 11, dv="Eficacia", desc="9.1.3 f)."), col("Data_Avaliacao", 11, "date", desc="Data.", req=False), col("Ligacoes_SGQ", 16, req=False, desc="Ligações.")]
    b.table("Oportunidades_Qualidade", "tbl_oport_q", ocols, orows,
            "Oportunidades que podem ter efeito desejado na conformidade e na satisfação (6.1.3 — subcapítulo novo na ISO 9001:2026).",
            title="OPORTUNIDADES DA QUALIDADE (6.1.3 — novo em 2026)", subtitle="Avaliação da eficácia separada da dos riscos (9.1.3 f; 9.3.2 h)",
            cf=[("Resultado_Eficacia", {"Não eficaz": "red", "Parcial": "orange", "Eficaz": "green"})], row_height=48)

    fcols = [col("ID_FMEA", 9, key="PK", desc="Linha da PFMEA."), col("Processo", 7, dv="Processo", desc="Processo."), col("Caracteristica", 22, desc="Característica do plano de controlo."),
             col("Classe", 8, dv="Classe", desc="Classe da característica."), col("Modo_Falha", 30, desc="Modo de falha potencial."), col("Efeito", 34, desc="Efeito no cliente."),
             col("S", 4, "int", desc="Severidade 1–10 (10 = segurança do consumidor / requisito legal)."), col("Causa", 32, desc="Causa principal conhecida."),
             col("Unidades_Inspecionadas", 11, "num0", desc="Unidades inspecionadas por atributos no período (dataset)."), col("Defeitos", 9, "num0", desc="Defeitos encontrados (dataset)."),
             col("PPM_Observado", 10, "num0", f='=IFERROR(@Defeitos@/@Unidades_Inspecionadas@*1000000,0)', desc="Defeitos por milhão observados na inspeção."),
             col("O", 4, "int", f='=IF(@PPM_Observado@>=100000,10,IF(@PPM_Observado@>=50000,9,IF(@PPM_Observado@>=20000,8,IF(@PPM_Observado@>=10000,7,IF(@PPM_Observado@>=2000,6,IF(@PPM_Observado@>=500,5,IF(@PPM_Observado@>=100,4,IF(@PPM_Observado@>=10,3,IF(@PPM_Observado@>0,2,1)))))))))',
                 desc="Ocorrência calculada pelo ppm observado (escala: ≥100 000→10 … >0→2; 0→1)."),
             col("Controlo_Detecao", 14, desc="Método de deteção do plano de controlo."), col("D", 4, "int", desc="Deteção 1–10 (visão 3, ensaio de fuga 3, passa/não passa 4, visual por amostragem 7)."),
             col("RPN", 6, "int", f='=@S@*@O@*@D@', desc="S × O × D."),
             col("Prioridade_Acao", 10, f='=IF(OR(AND(@S@>=9,@O@>=4),AND(@S@>=7,@O@>=6,@D@>=5)),"Alta",IF(OR(AND(@S@>=9,@O@>=2,@D@>=5),AND(@S@>=5,@O@>=5)),"Média","Baixa"))',
                 desc="Prioridade de ação simplificada (inspirada na tabela AP AIAG-VDA): severidade primeiro, depois ocorrência e deteção."),
             col("Acao_Recomendada", 36, desc="Ação para reduzir O ou D.", req=False), col("ID_Risco", 7, desc="Risco SGI relacionado.", key="FK → RG-SGA-02", req=False)]
    frows = pfmea_rows()
    acao = {"Anel de inviolabilidade": ("Rever pontes da ferramenta TE-012 e ensaio 100% em linha até Cpk ≥ 1,33", "R12"),
            "Estanquidade (fuga)": ("Leak tester 100% em linha nas ISBM sem deteção automática", "R16"), "Vedação": ("Controlo dimensional do vedante por SPC", "R16"),
            "Pontos negros": ("Certificado por lote de masterbatch; purga do cilindro", "R11"), "Injeção incompleta": ("Janela de parâmetros validada por DOE", "R4"),
            "Aderência": ("Ensaio no arranque e a meio do lote", "R6"), "Rebarba": ("Manutenção de moldes por contador de ciclos", "R2"), "Transferência do foil": ("Limitar velocidade; visão artificial HF-001", "R8")}
    for r in frows:
        a = acao.get(r["Caracteristica"])
        r["Acao_Recomendada"], r["ID_Risco"] = (a if a else (None, None))
    b.table("PFMEA_Processo", "tbl_pfmea", fcols, frows,
            "PFMEA resumida por processo × característica com ocorrência calculada a partir das inspeções por atributos do período.",
            title="PFMEA DE PROCESSO — OCORRÊNCIA A PARTIR DOS DADOS REAIS DE INSPEÇÃO",
            subtitle=f"Inspeção por atributos ISO 2859-1, {PER_INI} a {PER_FIM} · O, RPN e prioridade calculados · S e D definidos pela equipa (ver comentários)",
            cf=[("Prioridade_Acao", {"Alta": "red", "Média": "orange", "Baixa": "green"}), ("Classe", {"Crítica": "red", "Maior": "orange"})], row_height=30, freeze_col=3)

    ws = b.sheet("Resumo_Riscos", "Resumo calculado: riscos por opção de tratamento e resultado de eficácia; PFMEA por prioridade.")
    title(ws, "RESUMO DE RISCOS E OPORTUNIDADES DA QUALIDADE — calculado", "Entradas 9.3.2 g) e h) da revisão pela gestão")
    header_row(ws, 4, ["Resultado da eficácia", "Riscos (6.1.2)", "Oportunidades (6.1.3)"], widths=[34, 16, 20])
    for k, e in enumerate(["Eficaz", "Parcial", "Não eficaz", "Por avaliar"]):
        cell(ws, 5 + k, 1, e)
        cell(ws, 5 + k, 2, f'=COUNTIF(tbl_riscos_q[Resultado_Eficacia],"{e}")', fmt="0")
        cell(ws, 5 + k, 3, f'=COUNTIF(tbl_oport_q[Resultado_Eficacia],"{e}")', fmt="0")
    cell(ws, 9, 1, "Total", bold=True)
    cell(ws, 9, 2, "=SUM(B5:B8)", fmt="0", bold=True)
    cell(ws, 9, 3, "=SUM(C5:C8)", fmt="0", bold=True)
    header_row(ws, 11, ["Opção de tratamento (6.1.2 Nota 2)", "N.º de riscos", "dos quais disrupção"])
    for k, o in enumerate(OPCOES):
        cell(ws, 12 + k, 1, o)
        cell(ws, 12 + k, 2, f'=COUNTIF(tbl_riscos_q[Opcao_Tratamento],"{o}")', fmt="0")
        cell(ws, 12 + k, 3, f'=COUNTIFS(tbl_riscos_q[Opcao_Tratamento],"{o}",tbl_riscos_q[Disrupcao],"Sim")', fmt="0")
    r0 = 13 + len(OPCOES)
    header_row(ws, r0, ["PFMEA — prioridade de ação", "N.º de linhas", "RPN máximo"])
    for k, p in enumerate(["Alta", "Média", "Baixa"]):
        cell(ws, r0 + 1 + k, 1, p)
        cell(ws, r0 + 1 + k, 2, f'=COUNTIF(tbl_pfmea[Prioridade_Acao],"{p}")', fmt="0")
        cell(ws, r0 + 1 + k, 3, f'=_xlfn.MAXIFS(tbl_pfmea[RPN],tbl_pfmea[Prioridade_Acao],"{p}")', fmt="0")
    return b.save(out)


if __name__ == "__main__":
    import sys
    print(build(sys.argv[1]))
