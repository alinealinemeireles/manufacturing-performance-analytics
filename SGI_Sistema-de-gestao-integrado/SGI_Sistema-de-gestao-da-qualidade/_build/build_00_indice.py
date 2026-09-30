"""RG-SGQ-00 — Índice do SGQ e modelo de dados: catálogo de ficheiros e tabelas, integridade referencial (inclui ligações ao SGA),
matriz de requisitos ISO 9001:2026 com estado de prontidão, matriz mestra de informação documentada (A/B/C), fontes e normas usadas."""
import glob
import os
import re
import openpyxl
import pandas as pd
from sgqlib import *
from dimsq import *

SGA_REG = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "SGA_Sistema-de-gestao-ambiental", "Registos_SGA_Plasticom"))

FILES = [
    ("SGQ-01_Contexto_SWOT_PESTEL.xlsx", "RG-SGQ-01", "Contexto com a MESMA estrutura do RG-SGA-01: PESTEL, SWOT, TOWS, 5 forças e partes interessadas (lente da qualidade); alterações climáticas, âmbito e aplicabilidade", "4.1; 4.2; 4.3"),
    ("SGQ-02_Processos_Infraestrutura_Ambiente.xlsx", "RG-SGQ-02", "Mapa de processos (tartaruga), matriz de interações, 22 máquinas (MTBF, OEE, rejeição), ambiente dos processos", "4.4; 7.1.3; 7.1.4"),
    ("SGQ-03_Lideranca_Politica_Cultura.xlsx", "RG-SGQ-03", "Política da qualidade, liderança (5.1.1 a–l), RACI e autoridades, cultura ISO 10010, gemba walks, relatos de integridade", "5.1; 5.2; 5.3; 7.3 e)"),
    ("SGQ-04_Gestao_Riscos_Oportunidades.xlsx", "RG-SGQ-04", "Riscos (6.1.2) e oportunidades (6.1.3) da qualidade — vista do RG-SGA-02 — e PFMEA com ocorrência dos dados reais", "6.1; 9.1.3 e) f)"),
    ("SGQ-05_Objetivos_Metas_KPI.xlsx", "RG-SGQ-05", "Base mensal do dataset, 14 KPI mensais por fórmula, painel com gráficos, objetivos SMART (6.2.2 a–e), plano de monitorização", "6.2; 9.1.1; 9.1.3"),
    ("SGQ-06_Planeamento_Alteracoes.xlsx", "RG-SGQ-06", "Pedidos de alteração (MOC) com análise 6.3 a–g, autorização, aprovação do cliente e checklist (ISO 10007)", "6.3; 8.5.6; 8.2.4; 8.3.6"),
    ("SGQ-07_Competencias_Consciencializacao.xlsx", "RG-SGQ-07", "Requisitos por função, matriz pessoa × competência, formação com eficácia em 4 níveis (ISO 10015), consciencialização, conhecimento", "7.1.2; 7.1.6; 7.2; 7.3"),
    ("SGQ-08_Comunicacao_Controlo_Documental.xlsx", "RG-SGQ-08", "Lista mestra, documentos externos, retenção de registos, matriz de comunicação e comunicação com o cliente (contingência)", "7.4; 7.5; 8.2.1"),
    ("SGQ-09_Metrologia_Calibracao_MSA.xlsx", "RG-SGQ-09", "Inventário de 39 equipamentos (ISO 10012:2026), calibrações com banda de guarda, fora de tolerância, Gage R&R real por fórmulas", "7.1.5"),
    ("SGQ-10_Requisitos_Cliente_Encomendas.xlsx", "RG-SGQ-10", "18 clientes e requisitos, requisitos de produto, revisão das 539 encomendas de ago/2026, propriedade do cliente, pós-entrega", "5.1.2; 8.2; 8.5.3; 8.5.5"),
    ("SGQ-11_Design_Desenvolvimento.xlsx", "RG-SGQ-11", "6 projetos D&D (17 SKU alimentar/farma/potes): planeamento, entradas, stage-gate, saídas, alterações e FPY de lançamento", "8.3"),
    ("SGQ-12_Fornecedores_Avaliacao.xlsx", "RG-SGQ-12", "Fornecedores e processos externos, desempenho mensal do dataset, reavaliação por scorecard (A–D), SCAR, requisitos 8.4.3, auditorias", "8.4; 9.1.3 g)"),
    ("SGQ-13_Producao_Rastreabilidade_Libertacao.xlsx", "RG-SGQ-13", "Plano de controlo consolidado, libertação dos 6 361 lotes, exercícios de rastreio reais, erro humano, validação de processos, preservação", "8.1; 8.5.1; 8.5.2; 8.5.4; 8.6"),
    ("SGQ-14_Saidas_Nao_Conformes.xlsx", "RG-SGQ-14", "1 026 NC do dataset com tratamento 8.7.1 a–d, autoridade, clientes do lote e CAPA; 193 concessões com aprovação do cliente", "8.7"),
    ("SGQ-15_Reclamacoes_Satisfacao_Cliente.xlsx", "RG-SGQ-15", "158 reclamações reais com ciclo ISO 10002; inquérito ISO 10004 (CSI, NPS, CES, importância × satisfação), Pareto", "8.2.1 c); 9.1.2"),
    ("SGQ-16_Auditoria_Interna.xlsx", "RG-SGQ-16", "Auditores (competência/imparcialidade), programa baseado no risco, riscos do programa (ISO 19011:2026), checklist e constatações", "9.2"),
    ("SGQ-17_Revisao_pela_Gestao.xlsx", "RG-SGQ-17", "Revisão pela gestão RPG-2026-01: entradas 9.3.2 a–h com dados, decisões 9.3.3, ações anteriores e ata", "9.3; 5.1.1 l)"),
    ("SGQ-18_Nao_Conformidades_CAPA.xlsx", "RG-SGQ-18", "622 CAPA do dataset (estado à data de referência), CAPA de sistema 10.2.1 a–f, 8D das reclamações críticas, 5 Porquês", "10.2"),
    ("SGQ-19_Melhoria_Continua_Custo_Qualidade.xlsx", "RG-SGQ-19", "Carteira de projetos (DMAIC IM-002, PFMEA, SMED, TOC, visão), custo da qualidade PAF mensal, sugestões", "10.1; 9.1.3"),
]

# relações FK: (tabela origem, coluna, tabelas destino, colunas destino)
REL = [
    # contexto (RG-SGQ-01, mesma estrutura do RG-SGA-01); "SGA:" = tabela lida do registo do SGA
    ("tbl_swot", "ID_PESTEL", "tbl_pestel", "ID_PESTEL"), ("tbl_pestel", "ID_SWOT", "tbl_swot", "ID_SWOT"),
    ("tbl_pestel", "ID_RO", "tbl_riscos_q|tbl_oport_q", "ID_Risco|ID_Oport"), ("tbl_swot", "ID_RO", "tbl_riscos_q|tbl_oport_q", "ID_Risco|ID_Oport"),
    ("tbl_partes_interessadas", "ID_RO", "tbl_riscos_q|tbl_oport_q", "ID_Risco|ID_Oport"),
    ("tbl_pestel", "IDs_Risco_SGI", "SGA:tbRiscos|SGA:tbOportunidades", "ID|ID"), ("tbl_swot", "IDs_Risco_SGI", "SGA:tbRiscos|SGA:tbOportunidades", "ID|ID"),
    ("tbl_tows", "IDs_Risco_SGI", "SGA:tbRiscos|SGA:tbOportunidades", "ID|ID"), ("tbl_cinco_forcas", "IDs_Risco_SGI", "SGA:tbRiscos|SGA:tbOportunidades", "ID|ID"),
    ("tbl_partes_interessadas", "IDs_Risco_SGI", "SGA:tbRiscos|SGA:tbOportunidades", "ID|ID"),
    ("tbl_pestel", "ID_Contexto_SGI", "SGA:tbl_pestel|SGA:tbl_swot", "ID_PESTEL|ID_SWOT"), ("tbl_swot", "ID_Contexto_SGI", "SGA:tbl_pestel|SGA:tbl_swot", "ID_PESTEL|ID_SWOT"),
    ("tbl_tows", "IDs_Origem", "tbl_pestel|tbl_swot", "ID_PESTEL|ID_SWOT"), ("tbl_tows", "ID_Objetivo", "tbl_objetivos", "ID_OBJ"),
    ("tbl_partes_interessadas", "ID_PI_SGI", "SGA:tbl_partes_interessadas", "ID_PI"), ("tbl_partes_interessadas", "IDs_Legal", "SGA:tbl_legal", "ID_Legal"),
    ("tbl_ambiente", "ID_Contexto", "tbl_pestel|tbl_swot", "ID_PESTEL|ID_SWOT"),
    ("tbl_processos", "KPIs", "tbl_kpi", "ID_KPI"), ("tbl_processos", "Riscos_Oportunidades", "SGA:tbRiscos|SGA:tbOportunidades", "ID|ID"),
    ("tbl_interacoes", "De", "tbl_processos", "Codigo"), ("tbl_interacoes", "Para", "tbl_processos", "Codigo"),
    ("tbl_riscos_q", "ID_Risco", "SGA:tbRiscos", "ID"), ("tbl_riscos_q", "KPI", "tbl_kpi", "ID_KPI"), ("tbl_oport_q", "ID_Oport", "SGA:tbOportunidades", "ID"), ("tbl_oport_q", "KPI", "tbl_kpi", "ID_KPI"),
    ("tbl_pfmea", "ID_Risco", "SGA:tbRiscos", "ID"),
    ("tbl_objetivos", "ID_KPI", "tbl_kpi", "ID_KPI"), ("tbl_plano_monitorizacao", "ID_KPI", "tbl_kpi", "ID_KPI"),
    ("tbl_alteracoes", "IDs_Risco_SGI", "SGA:tbRiscos|SGA:tbOportunidades", "ID|ID"), ("tbl_checklist_alt", "ID_Alteracao", "tbl_alteracoes", "ID_Alteracao"),
    ("tbl_matriz", "ID_Pessoa", "tbl_pessoas", "ID_Pessoa"), ("tbl_matriz", "ID_Competencia", "tbl_competencias", "ID_Competencia"),
    ("tbl_formacoes", "ID_Competencia", "tbl_competencias", "ID_Competencia"), ("tbl_registo_formacao", "ID_Formacao", "tbl_formacoes", "ID_Formacao"),
    ("tbl_registo_formacao", "Colaborador", "tbl_pessoas", "ID_Pessoa"),
    ("tbl_calibracoes", "ID_Equipamento", "tbl_equipamentos", "ID_Equipamento"), ("tbl_fora_tolerancia", "ID_Equipamento", "tbl_equipamentos", "ID_Equipamento"),
    ("tbl_revisao_encomendas", "ID_Cliente", "tbl_clientes", "ID_Cliente"), ("tbl_alteracoes_requisitos", "ID_Cliente", "tbl_clientes", "ID_Cliente"),
    ("tbl_alteracoes_requisitos", "ID_MOC", "tbl_alteracoes", "ID_Alteracao"), ("tbl_propriedade_cliente", "ID_Cliente", "tbl_clientes", "ID_Cliente"),
    ("tbl_entradas_dd", "ID_Projeto", "tbl_projetos_dd", "ID_Projeto"), ("tbl_etapas", "ID_Projeto", "tbl_projetos_dd", "ID_Projeto"), ("tbl_etapas", "ID_Alteracao", "tbl_alteracoes_dd", "ID_Alteracao"),
    ("tbl_saidas_dd", "ID_Projeto", "tbl_projetos_dd", "ID_Projeto"),
    ("tbl_desempenho_forn", "ID_Fornecedor", "tbl_fornecedores", "ID_Fornecedor"), ("tbl_avaliacao", "ID_Fornecedor", "tbl_fornecedores", "ID_Fornecedor"),
    ("tbl_scar", "ID_Fornecedor", "tbl_fornecedores", "ID_Fornecedor"), ("tbl_auditorias_forn", "ID_Fornecedor", "tbl_fornecedores", "ID_Fornecedor"),
    ("tbl_fornecedores", "ID_Fornecedor", "SGA:tbl_fornecedores", "ID_Fornecedor", r"^SUP-"),   # EXT-* (prestadores) só existem no SGQ
    ("tbl_snc", "ID_Concessao", "tbl_concessoes", "ID_Concessao"), ("tbl_snc", "ID_CAPA", "tbl_capa", "ID_CAPA"), ("tbl_concessoes", "ID_NC", "tbl_snc", "ID_NC"),
    ("tbl_reclamacoes", "ID_Cliente", "tbl_clientes", "ID_Cliente"), ("tbl_reclamacoes", "ID_8D", "tbl_8d", "ID_8D"), ("tbl_8d", "ID_Reclamacao", "tbl_reclamacoes", "ID_Reclamacao"),
    ("tbl_inquerito", "ID_Cliente", "tbl_clientes", "ID_Cliente"), ("tbl_inquerito", "ID_Atributo", "tbl_atributos", "ID_Atributo"),
    ("tbl_checklist_aud", "ID_Auditoria", "tbl_programa_auditorias", "ID_Auditoria"), ("tbl_constatacoes", "ID_Auditoria", "tbl_programa_auditorias", "ID_Auditoria"),
    ("tbl_constatacoes", "ID_CAPA", "tbl_capa_sgq", "ID_CAPA"), ("tbl_capa", "ID_NC", "tbl_snc", "ID_NC"),
    ("tbl_5porques", "ID_CAPA", "tbl_capa_sgq", "ID_CAPA"), ("tbl_projetos_melhoria", "Riscos_Oport", "SGA:tbRiscos|SGA:tbOportunidades", "ID|ID"),
]
PAT = r"^(R\d+|O\d+|PES-|SWT-|PESQ-|SWTQ-|OBJ-Q-|LEG-|PI-|KPI-Q-|P-|OP-|AUX-|CQ-|FOR-Q-|EQM-|CUST-|MOC-Q-|DC-|SUP-|EXT-|CAPA-|CONC-|NC-|8D-|AUD-Q-|DD-|CC-|AT-|INF-|[A-Z]{2,3}$)"

# matriz de requisitos: (ID, cláusula, requisito, informação documentada exigida, classe A/B/C, documento, registo/tabelas, estado, observação, ação)
REQ = [
    ("RQ-01", "4.1", "Questões externas e internas pertinentes; monitorizar e rever", "Não exigida (C)", "C", "PR-SGQ-01", "RG-SGQ-01 tbl_pestel, tbl_swot, tbl_tows, tbl_cinco_forcas", "Conforme", "11 fatores PESTEL e 16 itens SWOT com efeito, evidência, monitorização e revisão; estrutura igual à do RG-SGA-01", ""),
    ("RQ-02", "4.1", "Determinar se as alterações climáticas são pertinentes (2026)", "Não exigida (C)", "C", "PR-SGQ-01", "RG-SGQ-01 tbl_clima", "Conforme", "Determinação e justificação registadas", ""),
    ("RQ-03", "4.2", "Partes interessadas, requisitos pertinentes e quais são tratados pelo SGQ", "Não exigida (C)", "C", "PR-SGQ-01", "RG-SGQ-01 tbl_partes_interessadas", "Conforme", "4.2 c) decidido por requisito", ""),
    ("RQ-04", "4.3", "Âmbito com tipos de produtos e justificação de não aplicabilidade", "Disponível", "A", "AMB-SGQ-01", "RG-SGQ-01 tbl_ambito, tbl_aplicabilidade", "Conforme", "Nenhum requisito excluído", ""),
    ("RQ-05", "4.4", "Processos, entradas/saídas, sequência, critérios, KPI, recursos, responsabilidades, riscos", "Disponível na medida necessária (4.4.2)", "C", "PR-SGQ-01", "RG-SGQ-02 tbl_processos, tbl_interacoes", "Conforme", "18 processos com tartaruga completa", ""),
    ("RQ-06", "5.1.1 a)–h), j)–k)", "Liderança e compromisso da gestão de topo", "Não exigida (C)", "C", "—", "RG-SGQ-03 tbl_lideranca, tbl_gemba", "Conforme", "Evidências com responsável", ""),
    ("RQ-07", "5.1.1 i)", "Promover a cultura da qualidade e o comportamento ético (2026)", "Não exigida (C)", "C", "—", "RG-SGQ-03 tbl_cultura, tbl_relatos_etica", "Parcial", "Maturidade 'Definida'; 3 relatos confirmados", "RPG-2026-01-D10"),
    ("RQ-08", "5.1.1 l)", "Prestar contas pela eficácia do SGQ (2026)", "Não exigida (C)", "C", "—", "RG-SGQ-17 Ata", "Conforme", "Conclusão assinada pelo DG", ""),
    ("RQ-09", "5.1.2", "Foco no cliente", "Não exigida (C)", "C", "—", "RG-SGQ-03 LID-13 a 15", "Conforme", "", ""),
    ("RQ-10", "5.2", "Política da qualidade estabelecida, comunicada, compreendida", "Disponível (5.2.2 a)", "A", "POL-SGQ-01", "RG-SGQ-03 Politica_Qualidade; tbl_verif_politica", "Conforme", "Reverificado em nov/2026 (CON-26-06, quiz 81%); CON-Q-26-09 fechada", "CON-26-06"),
    ("RQ-11", "5.3", "Funções, responsabilidades e autoridades (incl. f) integridade nas alterações)", "Não exigida (C)", "C", "—", "RG-SGQ-03 tbl_autoridades, tbl_raci", "Conforme", "RACI validada por fórmula", ""),
    ("RQ-12", "6.1.1", "Determinar riscos e oportunidades considerando 4.1 e 4.2", "Não exigida (C)", "C", "PR-SGQ-04", "RG-SGA-02; RG-SGQ-04", "Conforme", "Registo único do SGI", ""),
    ("RQ-13", "6.1.2", "Analisar e avaliar riscos; ações proporcionais; avaliar a eficácia", "Não exigida (C)", "C", "PR-SGQ-04", "RG-SGQ-04 tbl_riscos_q", "Parcial", "17 de 27 riscos com eficácia por avaliar; 3 não eficazes", "RPG-2026-01"),
    ("RQ-14", "6.1.3", "Analisar e avaliar oportunidades; avaliar a eficácia (2026)", "Não exigida (C)", "C", "PR-SGQ-04", "RG-SGQ-04 tbl_oport_q", "Parcial", "10 de 11 por avaliar", ""),
    ("RQ-15", "6.2", "Objetivos mensuráveis e planeamento (a–e)", "Disponível (6.2.1 g)", "A", "PR-SGQ-01", "RG-SGQ-05 tbl_objetivos", "Conforme", "8 objetivos com plano completo", ""),
    ("RQ-16", "6.3", "Alterações planeadas (a–g; f e g novos)", "Não exigida (C)", "C", "PR-SGQ-03", "RG-SGQ-06 tbl_alteracoes", "Conforme", "CON-Q-26-05 fechada; MOC-Q-26-16/17 com planeamento a–g completo", "CAPA-Q-26-10"),
    ("RQ-17", "7.1.1–7.1.3", "Recursos, pessoas e infraestrutura", "Não exigida (C)", "C", "PR-SGQ-15", "RG-SGQ-02 tbl_maquinas", "Conforme", "ISBM-005 com MTBF 8,4 h (R3)", ""),
    ("RQ-18", "7.1.4", "Ambiente para a operação dos processos (social, psicológico, físico)", "Não exigida (C)", "C", "—", "RG-SGQ-02 tbl_ambiente", "Conforme", "", ""),
    ("RQ-19", "7.1.5.1", "Recursos de monitorização e medição adequados", "Disponível como evidência", "B", "PR-SGQ-16", "RG-SGQ-09 tbl_equipamentos, GRR_Calculo, tbl_msa", "Conforme", "GRR 3,5% da TV", ""),
    ("RQ-20", "7.1.5.2", "Calibração rastreável, identificação do estado, avaliação retrospetiva", "Disponível (base, quando não há padrão)", "B", "PR-SGQ-16", "RG-SGQ-09 tbl_calibracoes, tbl_fora_tolerancia", "Conforme", "EQM-022/031 recalibrados (set–out/2026); 0 vencidos em 31/12/2026", "CAPA-Q-26-09"),
    ("RQ-21", "7.1.6", "Conhecimento organizacional", "Não exigida (C)", "C", "—", "RG-SGQ-07 tbl_conhecimento", "Conforme", "3 conhecimentos com risco de perda alto", ""),
    ("RQ-22", "7.2", "Competência", "Disponível como evidência", "B", "PR-SGQ-17", "RG-SGQ-07 tbl_matriz, tbl_registo_formacao", "Conforme", "Eficácia da formação avaliada", ""),
    ("RQ-23", "7.3", "Consciencialização (a–e; e cultura e ética)", "Não exigida (C)", "C", "PR-SGQ-17", "RG-SGQ-07 tbl_consciencializacao", "Conforme", "Sessão CON-26-06 (nov/2026) a todos os turnos", ""),
    ("RQ-24", "7.4", "Comunicação interna e externa", "Não exigida (C)", "C", "—", "RG-SGQ-08 tbl_comunicacao", "Conforme", "", ""),
    ("RQ-25", "7.5", "Informação documentada criada, atualizada e controlada", "— (requisitos de controlo)", "C", "PR-SGQ-02", "RG-SGQ-08 tbl_lista_mestra, tbl_doc_externos, tbl_retencao", "Conforme", "3 documentos com revisão vencida", ""),
    ("RQ-26", "8.1", "Planeamento e controlo operacional; confiança nos processos; alterações não previstas", "Disponível na medida necessária", "B", "PR-SGQ-11; PC-*", "RG-SGQ-13 tbl_plano_controlo", "Conforme", "", ""),
    ("RQ-27", "8.2.1", "Comunicação com o cliente (incl. e) contingência — 2026)", "Não exigida (C)", "C", "PR-SGQ-08", "RG-SGQ-08 tbl_comunicacao_cliente", "Conforme", "", ""),
    ("RQ-28", "8.2.2", "Requisitos dos produtos (incl. legais) e capacidade de cumprir declarações", "Não exigida (C)", "C", "PR-SGQ-08", "RG-SGQ-10 tbl_requisitos_produto", "Parcial", "PPWR parcial; Cpk ≥ 1,33 não demonstrado", "OBJ-Q-04"),
    ("RQ-29", "8.2.3", "Revisão dos requisitos antes do compromisso", "Disponível como evidência (8.2.3.2)", "B", "PR-SGQ-08", "RG-SGQ-10 tbl_revisao_encomendas", "Conforme", "0 exceções em 40 encomendas de out/2026 (CON-Q-26-14)", "CAPA-Q-26-13"),
    ("RQ-30", "8.2.4", "Alterações aos requisitos", "Documentação atualizada", "B", "PR-SGQ-08", "RG-SGQ-10 tbl_alteracoes_requisitos", "Conforme", "", ""),
    ("RQ-31", "8.3.2–8.3.5", "Planeamento, entradas, controlos e saídas do design", "Disponível como evidência (8.3.3, 8.3.4 f, 8.3.5)", "B", "PR-SGQ-10", "RG-SGQ-11", "Conforme", "", ""),
    ("RQ-32", "8.3.6", "Alterações do design (revisão, autorização, ações)", "Disponível como evidência", "B", "PR-SGQ-10", "RG-SGQ-11 tbl_alteracoes_dd", "Conforme", "DC-26-01 implementada; auditoria da configuração aos 17 SKU sem desvios", "CAPA-Q-26-11"),
    ("RQ-33", "8.4", "Critérios de avaliação, seleção, monitorização e reavaliação de fornecedores", "Disponível como evidência (8.4.1)", "B", "PR-SGQ-12", "RG-SGQ-12", "Parcial", "SUP-004/006 aprovados com classe C (rever estado)", "RPG-2026-01-D06"),
    ("RQ-34", "8.5.1", "Condições controladas (incl. f validação e g erro humano)", "Disponível (a — características, atividades, resultados)", "B", "PR-SGQ-13; IT-*", "RG-SGQ-13 tbl_plano_controlo, tbl_validacao, tbl_erro_humano", "Conforme", "2 medidas de erro humano em implementação", ""),
    ("RQ-35", "8.5.2", "Identificação e rastreabilidade", "Disponível como evidência (d)", "B", "PR-SGQ-13", "RG-SGQ-13 tbl_rastreio", "Parcial", "Rastreio de resina por FIFO (CON-Q-26-08)", "RPG-2026-01-D07"),
    ("RQ-36", "8.5.3", "Propriedade dos clientes ou fornecedores externos", "Disponível como evidência do que ocorreu", "B", "PR-SGQ-08", "RG-SGQ-10 tbl_propriedade_cliente", "Conforme", "", ""),
    ("RQ-37", "8.5.4–8.5.5", "Preservação e atividades pós-entrega", "Não exigida (C)", "C", "PR-SGQ-13", "RG-SGQ-13 tbl_preservacao; RG-SGQ-10 tbl_pos_entrega", "Conforme", "", ""),
    ("RQ-38", "8.5.6", "Controlo de alterações na produção (revisão, quem autoriza, ações)", "Disponível como evidência", "B", "PR-SGQ-03", "RG-SGQ-06", "Conforme", "CAPA-Q-26-10 fechada; alterações do 4.º trimestre autorizadas antes da implementação", "CAPA-Q-26-10"),
    ("RQ-39", "8.6", "Libertação com evidência de conformidade e de quem autorizou", "Disponível como evidência (a, b)", "B", "PR-SGQ-13", "RG-SGQ-13 tbl_libertacao", "Conforme", "6 361 lotes com inspetor e decisão", ""),
    ("RQ-40", "8.7", "Controlo de saídas não conformes (a–d) e concessões", "Disponível como evidência (8.7.2 a–d)", "B", "PR-SGQ-14", "RG-SGQ-14 tbl_snc, tbl_concessoes", "Parcial", "CON-Q-26-03 fechada: bloqueio no ERP desde 30/10/2026; eficácia por verificar (histórico do dataset mantém 10 concessões sem aprovação)", "CAPA-Q-26-08"),
    ("RQ-41", "9.1.1", "Monitorização, medição, análise e avaliação; evidência dos resultados", "Disponível como evidência", "B", "PR-SGQ-07", "RG-SGQ-05 tbl_kpi_mensal, tbl_plano_monitorizacao", "Conforme", "", ""),
    ("RQ-42", "9.1.2", "Satisfação do cliente (métodos para obter, monitorizar e rever)", "Não exigida (C)", "C", "PR-SGQ-09", "RG-SGQ-15", "Conforme", "CSI 78,5% (meta 80%)", ""),
    ("RQ-43", "9.1.3", "Análise e avaliação (a–h; e/f riscos e oportunidades)", "Não exigida (C)", "C", "PR-SGQ-07", "RG-SGQ-05 Painel_KPI; RG-SGQ-04", "Conforme", "", ""),
    ("RQ-44", "9.2", "Programa de auditoria e resultados", "Disponível como evidência (9.2.2)", "B", "PR-SGQ-05", "RG-SGQ-16", "Conforme", "Programa baseado no risco; 100% cumprido", ""),
    ("RQ-45", "9.3", "Revisão pela gestão (entradas a–h; saídas)", "Disponível como evidência (9.3.3)", "B", "PR-SGQ-01", "RG-SGQ-17", "Conforme", "RPG-2026-01", ""),
    ("RQ-46", "10.1", "Melhoria contínua", "Não exigida (C)", "C", "PR-SGQ-07", "RG-SGQ-19", "Conforme", "8 projetos; COQ mensal", ""),
    ("RQ-47", "10.2", "NC e ação corretiva (a–f); eficácia", "Disponível como evidência (10.2.2)", "B", "PR-SGQ-06", "RG-SGQ-18; RG-SGQ-14", "Parcial", "Comité de CAPA: vencidas 49% → 31%; 17% das NC maiores de out–nov ainda sem CAPA (CON-Q-26-16)", "CAPA-Q-26-12"),
]

# matriz mestra de informação documentada (proposta da pesquisa, preenchida para a Plasticom)
MESTRA = [
    ("DOC-001", "Âmbito do SGQ", "Documento", "GES", "4.3", "A", "Âmbito aprovado com tipos de produtos e aplicabilidade", GQ, DG, "Anual", "Até substituição + 5 anos", "Excel / PDF", "RG-SGQ-01 tbl_ambito", "4.3 (âmbito do SGA — RG-SGA-01)"),
    ("DOC-002", "Política da qualidade", "Documento", "GES", "5.2.2 a)", "A", "POL-SGQ-01 afixada e publicada", DG, DG, "Anual", "Até substituição + 5 anos", "PDF / afixação", "RG-SGQ-03 Politica_Qualidade", "5.2 (política ambiental)"),
    ("DOC-003", "Objetivos da qualidade", "Documento", "GES", "6.2.1 g)", "A", "Objetivos SMART com plano 6.2.2", GQ, DG, "Anual", "5 anos", "Excel / BI", "RG-SGQ-05 tbl_objetivos", "6.2 (RG-SGA-05)"),
    ("DOC-004", "Informação para operar os processos", "Documento", "Todos", "4.4.2 a); 8.5.1 a)", "A", "Procedimentos, planos de controlo, instruções", GQ, DIND, "24 meses", "Até substituição + 5 anos", "Eletrónico", "RG-SGQ-08 tbl_lista_mestra", "7.5 / 8.1"),
    ("REG-001", "Evidência de que os processos são realizados como planeado", "Registo", "Todos", "4.4.2 b); 8.1", "B", "Registos de produção, autocontrolo, SPC", GPROD, "—", "Contínua", "5 anos", "MES / data warehouse", "RG-SGQ-05 tbl_base_mensal; RG-SGQ-13", "8.1"),
    ("REG-002", "Adequação dos recursos de monitorização e medição", "Registo", "MET", "7.1.5.1", "B", "Inventário, MSA", GQ, "—", "Contínua", "Vida do equipamento + 5 anos", "Excel", "RG-SGQ-09 tbl_equipamentos, tbl_msa", "9.1.1 (RG-SGA-13)"),
    ("REG-003", "Base da calibração/verificação e estado", "Registo", "MET", "7.1.5.2 a)", "B", "Certificados, decisão com banda de guarda", GQ, "—", "Por intervalo", "Vida do equipamento + 5 anos", "PDF + Excel", "RG-SGQ-09 tbl_calibracoes", "9.1.1"),
    ("REG-004", "Evidência de competência", "Registo", "RH", "7.2", "B", "Matriz, formação, validação no posto", RH_, "—", "Por ação", "Contrato + 5 anos", "Excel + RH", "RG-SGQ-07 tbl_matriz, tbl_registo_formacao", "7.2 (RG-SGA-08)"),
    ("REG-005", "Resultados da revisão de requisitos e requisitos novos/alterados", "Registo", "COM", "8.2.3.2", "B", "Checklist por encomenda", DCOM, "—", "Por encomenda", "5 anos (farma 10)", "ERP + Excel", "RG-SGQ-10 tbl_revisao_encomendas", "—"),
    ("REG-006", "Entradas, controlos, saídas e alterações do design", "Registo", "RD", "8.3.3–8.3.6", "B", "Stage-gate, V&V, especificações", RD_, DIND, "Por projeto", "Vida do produto + 10 anos", "Repositório de projeto", "RG-SGQ-11", "8.1 (perspetiva de ciclo de vida)"),
    ("REG-007", "Avaliação, seleção, monitorização e reavaliação de fornecedores", "Registo", "CMP", "8.4.1", "B", "Scorecard, decisões, SCAR", CMP_, GQ, "Mensal / semestral", "5 anos", "Excel", "RG-SGQ-12", "8.1 (RG-SGA-11 fornecedores)"),
    ("REG-008", "Rastreabilidade", "Registo", "LAB", "8.5.2 d)", "B", "Lote → ordem → cliente; exercícios", GQ, "—", "Por lote", "Validade + 1 ano (farma 10)", "MES/ERP", "RG-SGQ-13 tbl_libertacao, tbl_rastreio", "—"),
    ("REG-009", "Propriedade do cliente (ocorrências)", "Registo", "COM", "8.5.3", "B", "Perdas/danos comunicados", GQ, "—", "Por ocorrência", "5 anos", "Excel", "RG-SGQ-10 tbl_propriedade_cliente", "—"),
    ("REG-010", "Revisão de alterações na produção", "Registo", "QUA", "8.5.6", "B", "MOC com autorização", GQ, DIND, "Por alteração", "Vida do produto + 5 anos", "Excel", "RG-SGQ-06", "6.3 (RG-SGA-18)"),
    ("REG-011", "Libertação do produto", "Registo", "LAB", "8.6", "B", "Decisão de lote + inspetor", GQ, "—", "Por lote", "Validade + 1 ano (mín. 5)", "MES + Excel", "RG-SGQ-13 tbl_libertacao", "—"),
    ("REG-012", "Saídas não conformes e concessões", "Registo", "LAB", "8.7.2", "B", "Natureza, ação, concessão, autoridade", GQ, "—", "Por NC", "5 anos (farma 10)", "Excel + MES", "RG-SGQ-14", "10.2 (RG-SGA-07)"),
    ("REG-013", "Resultados de monitorização e medição", "Registo", "QUA", "9.1.1", "B", "KPI mensais", GQ, "—", "Mensal", "5 anos", "Excel / BI", "RG-SGQ-05 tbl_kpi_mensal", "9.1.1 (RG-SGA-13)"),
    ("REG-014", "Programa e resultados de auditoria", "Registo", "QUA", "9.2.2", "B", "Programa, checklist, constatações", GQ, DG, "Programa anual", "2 ciclos de certificação", "Excel", "RG-SGQ-16", "9.2 (RG-SGA-14)"),
    ("REG-015", "Resultados da revisão pela gestão", "Registo", "GES", "9.3.3", "B", "Ata e decisões", GQ, DG, "Semestral", "10 anos", "Excel + PDF", "RG-SGQ-17", "9.3 (RG-SGA-15)"),
    ("REG-016", "NC e resultados das ações corretivas", "Registo", "QUA", "10.2.2", "B", "CAPA, 8D, eficácia", GQ, "—", "Por NC", "5 anos", "Excel + CAPA", "RG-SGQ-18", "10.2 (RG-SGA-07)"),
    ("GES-001", "Contexto e partes interessadas", "Registo de gestão", "GES", "4.1; 4.2", "C", "Questões e requisitos monitorizados", GQ, DG, "Semestral", "5 anos", "Excel", "RG-SGQ-01", "4.1; 4.2 (RG-SGA-01)"),
    ("GES-002", "Mapa de processos e interações", "Registo de gestão", "QUA", "4.4.1", "C", "Tartaruga e matriz", GQ, DIND, "Anual", "5 anos", "Excel", "RG-SGQ-02", "4.4"),
    ("GES-003", "Riscos e oportunidades", "Registo de gestão", "QUA", "6.1", "C", "Registo SGI + vista SGQ + PFMEA", GQ, DG, "Trimestral", "5 anos", "Excel", "RG-SGA-02; RG-SGQ-04", "6.1 (RG-SGA-02)"),
    ("GES-004", "Planeamento de alterações", "Registo de gestão", "QUA", "6.3", "C", "MOC a–g", GQ, DIND, "Por alteração", "5 anos", "Excel", "RG-SGQ-06", "6.3 (RG-SGA-18)"),
    ("GES-005", "Cultura da qualidade e liderança", "Registo de gestão", "GES", "5.1.1 i)", "C", "ISO 10010, gemba, relatos", DG, DG, "Semestral", "5 anos", "Excel", "RG-SGQ-03", "5.1"),
    ("GES-006", "Satisfação do cliente e reclamações", "Registo de gestão", "COM", "9.1.2; 8.2.1 c)", "C", "ISO 10002 / 10004", DCOM, DG, "Semestral", "5 anos", "Excel + CRM", "RG-SGQ-15", "7.4 (comunicação externa)"),
    ("GES-007", "Custo da qualidade e melhoria", "Registo de gestão", "QUA", "10.1", "C", "PAF e projetos", GQ, DG, "Mensal", "5 anos", "Excel", "RG-SGQ-19", "10.1 (RG-SGA-16)"),
]

NORMAS = [
    ("REF-01", "ISO 9001:2026 — requisitos (tradução oficial ES)", "Pasta ISO_9001_2026-...-Interpretação", "Todos os registos (cláusulas 4 a 10 lidas no original)", "Norma"),
    ("REF-02", "ISO/TC 176 — ISO 9001 Auditing Practices Group (APG)", "https://committee.iso.org/home/tc176/iso-9001-auditing-practices-group.html", "Contexto, alterações climáticas, processos, riscos, auditoria interna, fornecedores, NC; orientação, não requisitos", "Orientação"),
    ("REF-03", "APG — Guidance on Internal Audits", "committee.iso.org (PDF APG-InternalAudit)", "RG-SGQ-16: priorização por risco, competência e imparcialidade", "Orientação"),
    ("REF-04", "ISO 9001 Navigator Pro (iso9001help.co.uk)", "https://www.iso9001help.co.uk/iso-9001-navigator-pro.html", "Estrutura PDCA; procedimentos e registos por cláusula; preparação para 2026", "Orientação"),
    ("REF-05", "ISO 19011:2026 — auditorias de sistemas de gestão (4.ª ed., mai/2026)", "https://www.iso.org/standard/19011", "RG-SGQ-16: programa, riscos do programa, auditoria remota, competência", "Norma de orientação"),
    ("REF-06", "ISO 10002:2018 — tratamento de reclamações", "Pasta de interpretação (PDF)", "RG-SGQ-15 tbl_reclamacoes", "Norma de orientação"),
    ("REF-07", "ISO 10004:2018 — monitorização e medição da satisfação", "iso.org/standard/71582", "RG-SGQ-15 inquérito, atributos, importância × satisfação", "Norma de orientação"),
    ("REF-08", "ISO 10012:2026 — sistemas de gestão da medição (fev/2026)", "iso.org; Academy cap. 59", "RG-SGQ-09. NOTA: a ISO 10012 trata da medição, não da gestão de mudanças", "Norma"),
    ("REF-09", "ISO 10007:2017 — gestão da configuração", "iso.org/standard/70400", "RG-SGQ-06 (gestão de mudanças) e RG-SGQ-11 (auditoria da configuração)", "Norma de orientação"),
    ("REF-10", "ISO 10010:2022 — cultura da qualidade", "iso.org/standard/38457", "RG-SGQ-03 tbl_cultura", "Norma de orientação"),
    ("REF-11", "ISO 10015:2019 — gestão de competências", "iso.org/standard/69459", "RG-SGQ-07 formação e eficácia", "Norma de orientação"),
    ("REF-12", "ISO 10009:2024 — ferramentas da qualidade", "Pasta de interpretação (PDF)", "RG-SGQ-18 análise de causa", "Norma de orientação"),
    ("REF-13", "ISO 31000:2018 / IEC 31010:2019", "Pasta gestao_de_riscos_ISO_31000", "RG-SGA-02 / RG-SGQ-04", "Norma de orientação"),
    ("REF-14", "ICH Q9(R1) — gestão do risco da qualidade", "Pasta de interpretação (PDF)", "RG-SGQ-04 (produtos farmacêuticos)", "Orientação setorial"),
    ("REF-15", "ISO 2859-1:2026 / ISO 3951", "Academy cap. 63–64, 141–142", "RG-SGQ-13 amostragem e libertação", "Norma"),
    ("REF-16", "Manual de Engenharia da Qualidade (ASQ CQE/CSSBB) — Quality Engineering Academy", "C:/Project Portfolios/quality-engineering-academy", "Cap. 7 COQ, 59 ISO 10012, 88–93 SGQ e auditoria, 99 fornecedores, 157 VoC, 160 8D/CAPA, 175 GRR", "Teoria (ASQ BoK)"),
    ("REF-17", "Curso ISO 9001:2026 — OA 1.1 a 3.7 (M&T)", "Pasta de interpretação/M&T", "Estrutura dos registos e atividades", "Formação"),
    ("REF-18", "Dataset do projeto (bronze/silver/dim) e notebook", "datasets/; manufacturing_performance_analytics.ipynb", "Todos os dados operacionais e de qualidade", "Dados"),
    ("REF-19", "Pesquisa ChatGPT (fornecida pela utilizadora)", "Conversa de 28/09/2026", "Classificação A/B/C da informação documentada e matriz mestra (Matriz_Mestra_Info_Doc)", "Input"),
]

FONTE_UNICA = [
    ("Contexto externo/interno (PESTEL, SWOT, TOWS, 5 forças) — ambiente e SGI", "RG-SGA-01", "tbl_pestel; tbl_swot; tbl_tows; tbl_cinco_forcas", "RG-SGQ-01 (ID_Contexto_SGI)", "Estrutura comum: as tabelas homónimas do RG-SGA-01 e do RG-SGQ-01 juntam-se num único contexto do SGI (coluna Ambito)"),
    ("Contexto externo/interno — lente da qualidade", "RG-SGQ-01", "tbl_pestel; tbl_swot; tbl_tows; tbl_cinco_forcas", "RG-SGQ-04 (ID_RO); RG-SGQ-17", "Cada fator da qualidade liga ao fator de origem no RG-SGA-01 (ID_Contexto_SGI) — não se duplica o texto"),
    ("Partes interessadas (lista geral do SGI)", "RG-SGA-01", "tbl_partes_interessadas", "RG-SGQ-01 (ID_PI_SGI)", "Mesma estrutura no RG-SGQ-01: o SGQ regista os requisitos que trata (4.2 c) numa linha por requisito"),
    ("Riscos e oportunidades (pontuação, dono, estado)", "RG-SGA-02", "tbRiscos; tbOportunidades", "RG-SGQ-04 (vista gerada)", "Alterar só no RG-SGA-02 e regenerar o RG-SGQ-04"),
    ("Riscos: efeito na qualidade, opção de tratamento, eficácia (9001:2026)", "RG-SGQ-04", "tbl_riscos_q; tbl_oport_q", "RG-SGQ-17", ""),
    ("KPI e objetivos da qualidade", "RG-SGQ-05", "tbl_kpi; tbl_objetivos; tbl_kpi_mensal", "RG-SGQ-02/04 (ID_KPI)", ""),
    ("Alterações (MOC)", "RG-SGQ-06", "tbl_alteracoes", "RG-SGQ-10/11 (ID_MOC); RG-SGA-18 (ambiente)", "A mesma alteração tem avaliação ambiental no RG-SGA-18"),
    ("Pessoas, competências e formação da qualidade", "RG-SGQ-07", "tbl_pessoas; tbl_matriz; tbl_formacoes", "RG-SGQ-16 (auditores)", "Formação ambiental no RG-SGA-08"),
    ("Documentos do SGQ", "RG-SGQ-08", "tbl_lista_mestra", "Todos", "Documentos do SGA no RG-SGA-09"),
    ("Equipamentos de medição da qualidade", "RG-SGQ-09", "tbl_equipamentos", "RG-SGQ-13 (plano de controlo)", "Equipamentos ambientais no RG-SGA-13"),
    ("Clientes e requisitos", "RG-SGQ-10", "tbl_clientes", "RG-SGQ-15 (ID_Cliente)", ""),
    ("Fornecedores (dados mestre)", "RG-SGA-11", "tbl_fornecedores", "RG-SGQ-12 (mesmo ID_Fornecedor)", "Desempenho de qualidade só no RG-SGQ-12"),
    ("NC de produto e concessões", "RG-SGQ-14", "tbl_snc; tbl_concessoes", "RG-SGQ-18 (ID_NC)", "NC ambientais no RG-SGA-07"),
    ("Reclamações de clientes", "RG-SGQ-15", "tbl_reclamacoes", "RG-SGQ-18 (tbl_8d)", ""),
    ("Auditorias internas do SGQ", "RG-SGQ-16", "tbl_programa_auditorias; tbl_constatacoes", "RG-SGQ-18", "AUD-Q-26-03 foi integrada com o SGA (RG-SGA-14)"),
    ("Ações corretivas da qualidade", "RG-SGQ-18", "tbl_capa; tbl_capa_sgq", "RG-SGQ-14/16", "Ações ambientais no PAM (RG-SGA-06)"),
]


def load_tables(folder, extra=()):
    out, cat = {}, []
    files = sorted(glob.glob(os.path.join(folder, "SGQ-[0-9][0-9]_*.xlsx"))) + list(extra)
    for f in files:
        base = os.path.basename(f)
        if base.startswith("SGQ-00"):
            continue
        wb = openpyxl.load_workbook(f, data_only=True)
        for ws in wb.worksheets:
            for name, t in ws.tables.items():
                ref = t.ref if hasattr(t, "ref") else t
                rows = list(ws[ref])
                hdr = [c.value for c in rows[0]]
                df = pd.DataFrame([[c.value for c in r] for r in rows[1:]], columns=hdr)
                out[name if base.startswith("SGQ-") else f"SGA:{name}"] = df
                if base.startswith("SGQ-"):
                    hr = rows[0][0].row
                    cat.append(dict(Ficheiro=base, Folha=ws.title, Tabela=name, Linha_Cabecalho=hr, N_Linhas=int(df.iloc[:, 0].notna().sum()), N_Colunas=len(hdr),
                                    Chave_Primaria=hdr[0], Leitura_pandas=f'pd.read_excel("{base}", sheet_name="{ws.title}", header={hr - 1})'))
    return out, cat


def split_ids(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return []
    return [x.strip() for x in re.split(r"[;,]", str(v)) if x.strip() and x.strip() not in ("—", "-")]


def readme_value(path, key):
    """Valor de uma linha do 00_LEIA-ME de um registo (ex.: 'Utilização / atividades')."""
    try:
        ws = openpyxl.load_workbook(path, read_only=True)["00_LEIA-ME"]
        for a, b_ in ws.iter_rows(min_col=1, max_col=2, values_only=True):
            if a == key:
                return b_
    except (FileNotFoundError, KeyError):
        pass
    return None


# ------------------------------------------------------------------ harmonização SGQ ↔ SGA (base dos templates únicos para a gestão de topo)
# (ID do par, assunto, registo SGQ, tabela SGQ, registo SGA, tabela SGA)
PARES = [
    ("H-01", "Contexto — PESTEL", "SGQ-01", "tbl_pestel", "SGA-01", "tbl_pestel"),
    ("H-02", "Contexto — SWOT", "SGQ-01", "tbl_swot", "SGA-01", "tbl_swot"),
    ("H-03", "Contexto — Matriz TOWS", "SGQ-01", "tbl_tows", "SGA-01", "tbl_tows"),
    ("H-04", "Contexto — 5 forças de Porter", "SGQ-01", "tbl_cinco_forcas", "SGA-01", "tbl_cinco_forcas"),
    ("H-05", "Partes interessadas", "SGQ-01", "tbl_partes_interessadas", "SGA-01", "tbl_partes_interessadas"),
    ("H-06", "Catálogo de KPI", "SGQ-05", "tbl_kpi", "SGA-05", "tbl_kpi"),
    ("H-07", "Objetivos", "SGQ-05", "tbl_objetivos", "SGA-05", "tbl_objetivos"),
    ("H-08", "Plano de monitorização e medição", "SGQ-05", "tbl_plano_monitorizacao", "SGA-13", "tbl_plano_monitorizacao"),
    ("H-09", "Planeamento de alterações", "SGQ-06", "tbl_alteracoes", "SGA-18", "tbl_alteracoes"),
    ("H-10", "Formação — catálogo/plano", "SGQ-07", "tbl_formacoes", "SGA-08", "tbl_formacoes"),
    ("H-11", "Formação — registos individuais", "SGQ-07", "tbl_registo_formacao", "SGA-08", "tbl_registo_formacao"),
    ("H-12", "Matriz RACI", "SGQ-03", "tbl_raci", "SGA-08", "tbl_raci"),
    ("H-13", "Lista mestra de documentos", "SGQ-08", "tbl_lista_mestra", "SGA-09", "tbl_lista_mestra"),
    ("H-14", "Matriz de comunicação", "SGQ-08", "tbl_comunicacao", "SGA-09", "tbl_comunicacao"),
    ("H-15", "Equipamentos de medição", "SGQ-09", "tbl_equipamentos", "SGA-13", "tbl_equipamentos"),
    ("H-16", "Fornecedores", "SGQ-12", "tbl_fornecedores", "SGA-11", "tbl_fornecedores"),
    ("H-17", "Não conformidades (produto)", "SGQ-14", "tbl_snc", "SGA-07", "tbl_nc"),
    ("H-18", "Não conformidades de sistema / ação corretiva", "SGQ-18", "tbl_capa_sgq", "SGA-07", "tbl_nc"),
    ("H-19", "5 Porquês", "SGQ-18", "tbl_5porques", "SGA-07", "tbl_5porques"),
    ("H-20", "Programa de auditorias", "SGQ-16", "tbl_programa_auditorias", "SGA-14", "tbl_programa_auditorias"),
    ("H-21", "Constatações de auditoria", "SGQ-16", "tbl_constatacoes", "SGA-14", "tbl_constatacoes"),
    ("H-22", "Revisão pela gestão — entradas / agenda", "SGQ-17", "tbl_entradas_rpg", "SGA-15", "tbl_agenda"),
    ("H-23", "Revisão pela gestão — decisões", "SGQ-17", "tbl_decisoes", "SGA-15", "tbl_decisoes"),
    ("H-24", "Melhoria — sugestões / kaizen", "SGQ-19", "tbl_sugestoes", "SGA-16", "tbl_kaizen"),
    ("H-25", "Índice — ficheiros", "SGQ-00", "tbl_ficheiros", "SGA-00", "tbl_ficheiros"),
    ("H-26", "Índice — matriz de requisitos ISO", "SGQ-00", "tbl_matriz_iso", "SGA-00", "tbl_matriz_iso"),
    ("H-27", "Índice — fonte única", "SGQ-00", "tbl_fonte_unica", "SGA-00", "tbl_fonte_unica"),
    ("H-28", "Dimensão processo", "SGQ-00", "tbl_dim_processo", "SGA-00", "tbl_dim_processo"),
    ("H-29", "Dimensão função", "SGQ-00", "tbl_dim_funcao", "SGA-00", "tbl_dim_funcao"),
]
# colunas com o mesmo significado e nome próprio de cada norma: (tabela SGQ, coluna SGQ) → coluna SGA
EQUIV = {
    ("tbl_pestel", "Efeito_no_SGQ"): "Efeito_no_SGA",
    ("tbl_swot", "Implicacao_SGQ"): "Implicacao_SGA", ("tbl_swot", "Prioritario"): "Selecionado_Atv_3_1",
    ("tbl_partes_interessadas", "Tratado_pelo_SGQ"): "Torna_se_Obrigacao",
    ("tbl_objetivos", "Compromisso_Politica"): "ID_Politica", ("tbl_objetivos", "O_que_Fazer"): "Acoes_PAM",
    ("tbl_alteracoes", "IDs_Risco_SGI"): "ID_RO",
    ("tbl_fornecedores", "Data_Aprovacao"): "Data_Avaliacao",
    ("tbl_snc", "Acao_Tomada"): "Correcao_C1", ("tbl_snc", "ID_CAPA"): "ID_PAM_C2", ("tbl_snc", "Autoridade"): "Responsavel",
    ("tbl_capa_sgq", "ID_CAPA"): "ID_NC",
    ("tbl_5porques", "ID_CAPA"): "ID_NC",
    ("tbl_programa_auditorias", "Data_Planeada"): "Data_Inicio", ("tbl_programa_auditorias", "Data_Real"): "Data_Fim",
    ("tbl_constatacoes", "Constatacao"): "Requisito_Violado_ou_Melhoravel", ("tbl_constatacoes", "ID_CAPA"): "ID_PAM",
    ("tbl_decisoes", "Ligacao"): "IDs_PAM", ("tbl_raci", "Processo_SGQ"): "Processo_SGA",
    ("tbl_sugestoes", "ID_Sugestao"): "ID_KZ", ("tbl_sugestoes", "Sugestao"): "Descricao", ("tbl_sugestoes", "Processo"): "Area",
}


def table_headers(path):
    wb = openpyxl.load_workbook(path)
    return {n: [c.name for c in ws.tables[n].tableColumns] for ws in wb.worksheets for n in ws.tables}


def harmonizacao(folder, own):
    """Compara coluna a coluna as tabelas do SGQ com as do SGA (own = cabeçalhos das tabelas deste índice, ainda não gravado)."""
    files = {os.path.basename(f)[:6]: f for f in glob.glob(os.path.join(folder, "SGQ-[0-9][0-9]_*.xlsx")) + glob.glob(os.path.join(SGA_REG, "SGA-[0-9][0-9]_*.xlsx"))}
    cache = {"SGQ-00": own}
    det, res = [], []
    for pid, assunto, rq, tq, ra, ta in PARES:
        for r in (rq, ra):
            if r not in cache:
                cache[r] = table_headers(files[r])
        cq, ca = cache[rq].get(tq, []), cache[ra].get(ta, [])
        usados = set()
        for i, c in enumerate(cq, 1):
            if c in ca:
                alvo, tipo = c, "Igual"
            elif (tq, c) in EQUIV and EQUIV[(tq, c)] in ca and EQUIV[(tq, c)] not in cq:
                alvo, tipo = EQUIV[(tq, c)], "Equivalente (renomear na consolidação)"
            else:
                alvo, tipo = None, "Só SGQ"
            if alvo:
                usados.add(alvo)
            det.append(dict(ID_Par=pid, Tabela_SGQ=tq, Coluna_SGQ=c, Posicao_SGQ=i, Tabela_SGA=ta, Coluna_SGA=alvo,
                            Posicao_SGA=(ca.index(alvo) + 1) if alvo else None, Correspondencia=tipo))
        for j, c in enumerate(ca, 1):
            if c not in usados:
                det.append(dict(ID_Par=pid, Tabela_SGQ=tq, Coluna_SGQ=None, Posicao_SGQ=None, Tabela_SGA=ta, Coluna_SGA=c, Posicao_SGA=j, Correspondencia="Só SGA"))
        res.append(dict(ID_Par=pid, Assunto=assunto, Registo_SGQ="RG-" + rq, Tabela_SGQ=tq, Registo_SGA="RG-" + ra, Tabela_SGA=ta,
                        N_Colunas_SGQ=len(cq), N_Colunas_SGA=len(ca)))
    return res, det


def integridade(T):
    res = []
    for rel in REL:
        src, colname, dst, dcol = rel[:4]
        pat = rel[4] if len(rel) > 4 else PAT
        if src not in T or colname not in T[src].columns:
            res.append(dict(Tabela_Origem=src, Coluna=colname, Tabela_Destino=dst.replace("|", " / ").replace("SGA:", "SGA · "), N_Referencias=0, N_Invalidas=0, IDs_Invalidos="tabela/coluna não encontrada", Resultado="Rever"))
            continue
        valid = set()
        for tt, cc in zip(dst.split("|"), dcol.split("|")):
            if tt in T and cc in T[tt].columns:
                valid |= set(T[tt][cc].dropna().astype(str))
        refs = [i for v in T[src][colname] for i in split_ids(v)]
        refs = [r for r in refs if re.match(pat, r)]
        missing = sorted({r for r in refs if r not in valid})
        res.append(dict(Tabela_Origem=src, Coluna=colname, Tabela_Destino=dst.replace("|", " / ").replace("SGA:", "SGA · "), N_Referencias=len(refs), N_Invalidas=len(missing),
                        IDs_Invalidos="; ".join(missing[:25]) + (" …" if len(missing) > 25 else "") or None, Resultado="OK" if not missing else "Rever"))
    return res


def build(out, folder):
    extra = [os.path.join(SGA_REG, f) for f in ("SGA-01_Contexto_SWOT_PESTEL.xlsx", "SGA-02_Gestao_Riscos_Oportunidades.xlsx", "SGA-04_Requisitos_Legais_Conformidade.xlsx", "SGA-11_Fornecedores_Ciclo_Vida.xlsx")]
    T, cat = load_tables(folder, extra)
    b = Book("RG-SGQ-00", "Índice do SGQ, Matriz ISO 9001:2026 e Modelo de Dados",
             activities="Ponto de entrada do sistema de registos do SGQ: onde está cada requisito, que tabelas existem, como se ligam entre si e ao SGA, e o estado de prontidão para a ISO 9001:2026.",
             clauses="7.5 Informação documentada (lista de registos e matriz mestra A/B/C); 4.4; 9.1.3 (dados para análise); transição para a ISO 9001:2026",
             purpose="Catálogo dos 19 registos e das suas tabelas, verificação automática da integridade referencial (incluindo ligações ao SGA-01, SGA-02 e SGA-11), matriz requisito → documento → evidência → estado para a ISO 9001:2026 com índice de prontidão por capítulo, matriz mestra de informação documentada (A obrigatória / B evidência / C necessária à eficácia), fontes consultadas e matriz de fonte única SGQ ↔ SGA.",
             guidance=[("ISO 9001:2026 7.5.1 Nota e APG 'Documented information'", "A norma não exige manual nem um procedimento por requisito: a matriz separa o que é obrigatório (A, B) do que a Plasticom determinou necessário (C)."),
                       ("Pesquisa ChatGPT fornecida (28/09/2026)", "Estrutura da matriz mestra (código, tipo, processo, requisito, obrigatoriedade, evidência, responsável, aprovação, frequência, retenção, formato) e relação com a ISO 14001.")])
    b.add_list("Estado", ["Conforme", "Parcial", "Lacuna"])
    b.add_list("Classe", ["A", "B", "C"])
    b.add_list("SimNao", ["Sim", "Não"])

    # mesmas colunas do RG-SGA-00 tbl_ficheiros (Ficheiro, Codigo, Conteudo, Atividades, Clausulas)
    fcols = [col("Ficheiro", 48, desc="Ficheiro."), col("Codigo", 10, key="PK", desc="Código do registo."), col("Conteudo", 90, desc="Conteúdo."),
             col("Atividades", 60, desc="Utilização / atividades (lido do 00_LEIA-ME do registo).", req=False), col("Clausulas", 20, desc="Cláusulas ISO 9001:2026.")]
    frows = []
    for fn, cod, cont, cl in FILES:
        frows.append(dict(Ficheiro=fn, Codigo=cod, Conteudo=cont, Atividades=readme_value(os.path.join(folder, fn), "Utilização / atividades"), Clausulas=cl))
    b.table("Ficheiros", "tbl_ficheiros", fcols, frows, "Lista dos registos do SGQ.",
            title="ÍNDICE DO SISTEMA DE REGISTOS DO SGQ — PLASTICOM (ISO 9001:2026)",
            subtitle=f"Pasta Registos_SGQ_Plasticom · 19 registos Excel + este índice · Gerados por _build/build_all.py a partir do dataset · Data de referência {DATA_REF:%d/%m/%Y}", row_height=30, tab_color="1F4E5F")

    # mesmas colunas e nome de tabela do RG-SGA-00 tbl_matriz_iso; Classe (A/B/C) só no SGQ, no fim
    rcols = [col("ID_Requisito", 6, key="PK", desc="Requisito."), col("Clausula", 12, desc="Cláusula."),
             col("Grupo", 6, f='=IF(@Clausula@="","",IF(LEFT(@Clausula@,2)="10","10",LEFT(@Clausula@,1)))', desc="Capítulo 4–10."),
             col("Requisito", 50, desc="Requisito (resumo)."),
             col("Novo_ou_Alterado_2026", 8, f='=IF(@ID_Requisito@="","",IF(OR(ISNUMBER(SEARCH("2026",@Requisito@)),@Clausula@="6.1.3",@Clausula@="5.1.1 i)",@Clausula@="5.1.1 l)"),"Sim","Não"))', desc="Novo/alterado na edição 2026 (assinalado no texto)."),
             col("Informacao_Documentada", 28, desc="O que a norma exige: 'disponível' / 'disponível como evidência' / não exigida."),
             col("Documento", 18, desc="Documento controlado."), col("Evidencia_Registo_Tabelas", 44, desc="Registo e tabelas com a evidência."),
             col("Estado", 9, dv="Estado", desc="Conforme · Parcial · Lacuna (verificação de 31/12/2026)."), col("Observacao", 44, desc="Observação.", req=False), col("Acao", 16, desc="Ação que trata a lacuna.", req=False),
             col("Pontos", 6, "num", f='=IF(@Estado@="","",IF(@Estado@="Conforme",1,IF(@Estado@="Parcial",0.5,0)))', desc="1 · 0,5 · 0 (índice de prontidão)."),
             col("Classe", 6, dv="Classe", desc="[Só SGQ] A documento obrigatório · B evidência obrigatória · C determinada pela organização.")]
    REQ_N = ["ID_Requisito", "Clausula", "Requisito", "Informacao_Documentada", "Classe", "Documento", "Evidencia_Registo_Tabelas", "Estado", "Observacao", "Acao"]
    b.table("Matriz_ISO9001_2026", "tbl_matriz_iso", rcols, rows_from(REQ_N, REQ),
            "Mapa requisito → informação documentada → evidência → estado para a ISO 9001:2026 (mesma estrutura do RG-SGA-00 Matriz_ISO14001_2026).",
            title="MATRIZ DE REQUISITOS ISO 9001:2026 — DOCUMENTO, EVIDÊNCIA E ESTADO",
            subtitle="Verificação de 31/12/2026 contra os registos do SGQ · Lacuna = NC maior de auditoria · Parcial = processo existe com falha de evidência ou eficácia",
            cf=[("Estado", {"Conforme": "green", "Parcial": "orange", "Lacuna": "red"}), ("Classe", {"A": "blue", "B": "purple"}), ("Novo_ou_Alterado_2026", {"Sim": "blue"})], row_height=32, freeze_col=2)

    ws = b.sheet("Resumo_Prontidao", "Índice de prontidão para a ISO 9001:2026 por capítulo (calculado a partir da matriz).", tab_color="C00000")
    title(ws, "PRONTIDÃO PARA A ISO 9001:2026 — RESUMO POR CAPÍTULO — calculado", "Índice = (Conforme + 0,5 × Parcial) ÷ requisitos")
    header_row(ws, 3, ["Capítulo", "Tema", "Requisitos", "Conforme", "Parcial", "Lacuna", "Novos 2026", "Índice de prontidão"], widths=[10, 26, 11, 10, 10, 10, 11, 14])
    temas = [("4", "Contexto"), ("5", "Liderança"), ("6", "Planeamento"), ("7", "Suporte"), ("8", "Operação"), ("9", "Avaliação do desempenho"), ("10", "Melhoria")]
    R = lambda c_: f"tbl_matriz_iso[{c_}]"
    for k, (g, t) in enumerate(temas):
        r = 4 + k
        cell(ws, r, 1, g)
        cell(ws, r, 2, t)
        cell(ws, r, 3, f'=COUNTIF({R("Grupo")},A{r})', fmt="0")
        for j, e in enumerate(("Conforme", "Parcial", "Lacuna")):
            cell(ws, r, 4 + j, f'=COUNTIFS({R("Grupo")},A{r},{R("Estado")},"{e}")', fmt="0")
        cell(ws, r, 7, f'=COUNTIFS({R("Grupo")},A{r},{R("Novo_ou_Alterado_2026")},"Sim")', fmt="0")
        cell(ws, r, 8, f'=IFERROR(SUMIFS({R("Pontos")},{R("Grupo")},A{r})/C{r},"")', fmt="0%")
    rt = 4 + len(temas)
    cell(ws, rt, 1, "Total", bold=True)
    for c_ in range(3, 8):
        L = get_column_letter(c_)
        cell(ws, rt, c_, f"=SUM({L}4:{L}{rt - 1})", fmt="0", bold=True)
    cell(ws, rt, 8, f'=IFERROR(SUM({R("Pontos")})/C{rt},"")', fmt="0%", bold=True)
    ws.conditional_formatting.add(f"H4:H{rt}", FormulaRule(formula=["H4<0.75"], fill=PatternFill("solid", fgColor=CF_COLORS["orange"][0])))
    for c_ in (3, 4):
        cell(ws, rt + 2 + (c_ - 3), 1, ["Informação documentada obrigatória (A) na matriz", "Evidência obrigatória (B) na matriz"][c_ - 3], border=False)
        cell(ws, rt + 2 + (c_ - 3), 3, f'=COUNTIF({R("Classe")},"{"A" if c_ == 3 else "B"}")', fmt="0", border=False)

    mcols = [col("Codigo", 8, key="PK", desc="Código na matriz mestra."), col("Documento_Registo", 40, desc="Informação documentada."), col("Tipo", 13, desc="Documento / registo / registo de gestão."),
             col("Processo", 7, desc="Processo."), col("Requisito", 14, desc="Requisito ISO 9001:2026."), col("Obrigatoriedade", 7, dv="Classe", desc="A / B / C."),
             col("Evidencia", 34, desc="O que demonstra."), col("Responsavel", 22, desc="Responsável."), col("Aprovacao", 18, desc="Quem aprova."), col("Frequencia", 14, desc="Frequência."),
             col("Retencao", 22, desc="Retenção."), col("Formato", 16, desc="Formato."), col("Onde_no_SGQ", 34, desc="Registo e tabela."), col("Relacao_ISO14001", 28, desc="Correspondência no SGA (SGI).")]
    b.table("Matriz_Mestra_Info_Doc", "tbl_matriz_mestra", mcols, rows_from(input_names(mcols), MESTRA),
            "Matriz mestra de informação documentada da Plasticom: A obrigatória pela ISO, B evidência obrigatória, C necessária para a eficácia.",
            title="MATRIZ MESTRA DE INFORMAÇÃO DOCUMENTADA — ISO 9001:2026",
            subtitle="A = a norma exige que esteja 'disponível' · B = 'disponível como evidência' · C = a Plasticom determinou necessária (7.5.1 b) · Não há manual da qualidade (não exigido)",
            cf=[("Obrigatoriedade", {"A": "blue", "B": "purple", "C": "gray"})], row_height=30, freeze_col=2)

    ccols = [col("Ficheiro", 44, desc="Ficheiro."), col("Folha", 26, desc="Folha."), col("Tabela", 24, key="PK", desc="Tabela Excel."), col("Linha_Cabecalho", 8, "int", desc="Linha do cabeçalho."),
             col("N_Linhas", 8, "int", desc="Registos (1.ª coluna preenchida)."), col("N_Colunas", 8, "int", desc="Colunas."), col("Chave_Primaria", 18, desc="Primeira coluna."),
             col("Leitura_pandas", 80, desc="Código para ler a tabela em Python.")]
    b.table("Catalogo_Tabelas", "tbl_catalogo", ccols, cat, "Catálogo de todas as tabelas de dados do SGQ (gerado automaticamente).")
    icols = [col("Tabela_Origem", 22, desc="Tabela com a chave estrangeira."), col("Coluna", 18, desc="Coluna FK."), col("Tabela_Destino", 40, desc="Tabela referenciada (SGQ ou SGA)."),
             col("N_Referencias", 10, "int", desc="Referências verificadas."), col("N_Invalidas", 9, "int", desc="Sem correspondência."), col("IDs_Invalidos", 50, desc="IDs em falta.", req=False),
             col("Resultado", 9, desc="OK / Rever.")]
    b.table("Integridade_Referencial", "tbl_integridade", icols, integridade(T), "Verificação das ligações entre registos do SGQ e com o SGA (instantâneo da construção).",
            cf=[("Resultado", {"Rever": "red", "OK": "green"})])

    b.table("Normas_Fontes", "tbl_fontes", [col("ID", 7, key="PK"), col("Referencia", 56, desc="Norma / fonte."), col("Onde", 48, desc="Localização / URL."),
                                             col("Aplicada_em", 60, desc="Onde foi aplicada."), col("Tipo", 16, desc="Requisito / orientação / dados.")],
            [dict(zip(["ID", "Referencia", "Onde", "Aplicada_em", "Tipo"], n)) for n in NORMAS], "Normas, orientações e fontes consultadas para construir o SGQ.",
            title="NORMAS, ORIENTAÇÕES E FONTES USADAS", subtitle="As orientações do APG e das normas ISO 100xx não são requisitos auditáveis da ISO 9001 — foram usadas como boas práticas", row_height=30)
    b.table("Matriz_Fonte_Unica", "tbl_fonte_unica", [col("Assunto", 44, key="PK"), col("Registo_Dono", 11, desc="Único registo onde o assunto é escrito."), col("Tabelas_Donas", 32),
                                                        col("So_Referenciado_Em", 40, desc="Onde aparece só por ID ou valor gerado."), col("Regra_Nota", 50, req=False)],
            [dict(zip(["Assunto", "Registo_Dono", "Tabelas_Donas", "So_Referenciado_Em", "Regra_Nota"], f)) for f in FONTE_UNICA],
            "Matriz de fonte única do SGI (SGQ ↔ SGA): evita registar o mesmo assunto em dois ficheiros.", title="MATRIZ DE FONTE ÚNICA — SGQ ↔ SGA", row_height=30)

    b.table("Dim_Processo", "tbl_dim_processo", [col("Codigo", 7, key="PK"), col("Processo", 56), col("Tipo", 11), col("Maquinas_Dataset", 20), col("Dono", 30, desc="[Só SGQ] Dono do processo.")],
            [dict(Codigo=a, Processo=p, Tipo=t, Dono=d, Maquinas_Dataset=m) for a, p, t, d, m in PROCESSOS], "Dimensão processo do SGQ (códigos produtivos iguais aos do SGA).")
    b.table("Dim_Funcao", "tbl_dim_funcao", [col("ID_Funcao", 8, key="PK"), col("Funcao", 44), col("Area", 8, desc="Processo / área da função (código de processo).")],
            [dict(ID_Funcao=a, Funcao=f, Area=p) for a, f, p in FUNCOES], "Dimensão função (mesmas colunas do RG-SGA-00).")
    b.table("Dim_Clausula_ISO9001_2026", "tbl_dim_clausula", [col("Clausula", 9, key="PK"), col("Titulo", 60), col("Novo_ou_Alterado_2026", 10, dv="SimNao"), col("Nota_Alteracao", 80, req=False)],
            [dict(Clausula=a, Titulo=t, Novo_ou_Alterado_2026=n, Nota_Alteracao=x or None) for a, t, n, x in CLAUSULAS], "Cláusulas da ISO 9001:2026 e alterações face à 2015.",
            cf=[("Novo_ou_Alterado_2026", {"Sim": "blue"})], row_height=28)

    # ---------------------------------------------------------------- harmonização SGQ ↔ SGA
    own = {n: list(b.tables[n]["colmap"]) for n in ("tbl_ficheiros", "tbl_matriz_iso", "tbl_fonte_unica", "tbl_dim_processo", "tbl_dim_funcao")}
    hres, hdet = harmonizacao(folder, own)
    b.add_list("Correspondencia", ["Igual", "Equivalente (renomear na consolidação)", "Só SGQ", "Só SGA"])
    H = lambda c_: f"tbl_harmonizacao[{c_}]"
    hcols = [col("ID_Par", 7, key="PK", desc="Par de tabelas SGQ ↔ SGA."), col("Assunto", 34, desc="Assunto comum."),
             col("Registo_SGQ", 10, desc="Registo do SGQ."), col("Tabela_SGQ", 22, desc="Tabela do SGQ."),
             col("Registo_SGA", 10, desc="Registo do SGA (padrão do SGI)."), col("Tabela_SGA", 22, desc="Tabela do SGA."),
             col("Mesmo_Nome_Tabela", 9, f='=IF(@ID_Par@="","",IF(@Tabela_SGQ@=@Tabela_SGA@,"Sim","Não"))', desc="A tabela tem o mesmo nome nos dois sistemas?"),
             col("N_Colunas_SGQ", 8, "int", desc="Colunas da tabela do SGQ."), col("N_Colunas_SGA", 8, "int", desc="Colunas da tabela do SGA."),
             col("N_Iguais", 8, "int", f=f'=COUNTIFS({H("ID_Par")},@ID_Par@,{H("Correspondencia")},"Igual")', desc="Colunas com o mesmo nome."),
             col("N_Equivalentes", 9, "int", f=f'=COUNTIFS({H("ID_Par")},@ID_Par@,{H("Correspondencia")},"Equivalente*")', desc="Mesmo significado, nome próprio da norma."),
             col("N_So_SGQ", 8, "int", f=f'=COUNTIFS({H("ID_Par")},@ID_Par@,{H("Correspondencia")},"Só SGQ")', desc="Colunas só do SGQ (ficam vazias no SGA ao consolidar)."),
             col("N_So_SGA", 8, "int", f=f'=COUNTIFS({H("ID_Par")},@ID_Par@,{H("Correspondencia")},"Só SGA")', desc="Colunas só do SGA (ficam vazias no SGQ ao consolidar)."),
             col("Pct_Comum_SGA", 9, "pct", f='=IF(@ID_Par@="","",IFERROR((@N_Iguais@+@N_Equivalentes@)/@N_Colunas_SGA@,""))',
                 desc="% das colunas do SGA que existem no SGQ (iguais + equivalentes).")]
    b.table("Harmonizacao_Resumo", "tbl_harmonizacao_resumo", hcols, hres,
            "Resumo da normalização SGQ ↔ SGA por par de tabelas (base dos templates únicos para a gestão de topo).",
            title="HARMONIZAÇÃO SGQ ↔ SGA — RESUMO POR TABELA",
            subtitle="Padrão do SGI = registos do SGA · Consolidar: juntar as tabelas homónimas (pd.concat / Power BI Append) e renomear só as colunas 'Equivalentes' · Gerado automaticamente",
            cf=[("Mesmo_Nome_Tabela", {"Sim": "green", "Não": "orange"})], row_height=18, freeze_col=2)
    dcols = [col("ID_Par", 7, desc="Par.", key="FK → tbl_harmonizacao_resumo"), col("Tabela_SGQ", 22, desc="Tabela do SGQ."),
             col("Coluna_SGQ", 26, desc="Coluna no SGQ.", req=False), col("Posicao_SGQ", 7, "int", desc="Posição da coluna no SGQ.", req=False),
             col("Tabela_SGA", 22, desc="Tabela do SGA."), col("Coluna_SGA", 26, desc="Coluna correspondente no SGA.", req=False),
             col("Posicao_SGA", 7, "int", desc="Posição da coluna no SGA.", req=False),
             col("Correspondencia", 22, dv="Correspondencia", desc="Igual · Equivalente (renomear na consolidação) · Só SGQ · Só SGA."),
             col("Nome_Consolidado", 26, f='=IF(@ID_Par@="","",IF(@Coluna_SGA@<>"",@Coluna_SGA@,@Coluna_SGQ@))', desc="Nome a usar no template único do SGI (o do SGA quando existe).")]
    b.table("Matriz_Harmonizacao_SGI", "tbl_harmonizacao", dcols, hdet,
            "Correspondência coluna a coluna entre as tabelas do SGQ e do SGA (gerada automaticamente a partir dos ficheiros).",
            title="MATRIZ DE HARMONIZAÇÃO SGQ ↔ SGA — COLUNA A COLUNA",
            subtitle="Nome_Consolidado = nome da coluna no template único da gestão de topo · 'Só SGQ' / 'Só SGA' = coluna específica da norma (fica vazia no outro sistema)",
            cf=[("Correspondencia", {"Igual": "green", "Equivalente": "yellow", "Só SGQ": "blue", "Só SGA": "gray"})], row_height=15, freeze_col=3)

    ws = b.sheet("Modelo_Dados", "Descrição do modelo de dados e convenções de IDs.")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120
    title(ws, "MODELO DE DADOS DO SGQ")
    lines = [
        ("Estrutura", "Modelo em estrela: factos (base mensal, KPI mensais, libertação de lotes, NC, CAPA, reclamações, receções de MP, calibrações, formação, encomendas) ligados a dimensões partilhadas (mês, processo, função, cláusula, cliente, fornecedor, máquina, produto, pessoa) e a registos de gestão (riscos, objetivos, MOC, auditorias, decisões) por IDs estáveis."),
        ("Ligação ao dataset", "Os IDs de máquina (IM-/ISBM-/SS-/HF-), molde, produto, lote, ordem (WO-), encomenda (SO-), cliente (CUST-), fornecedor (SUP-), NC (NC-), CAPA (CAPA-), reclamação (CC-) e operador (OP-) são os do dataset — permite cruzar os registos do SGQ com o data warehouse, o notebook e o Power BI."),
        ("Convenção de IDs do SGQ", "PESQ / SWTQ / TOWSQ / CFQ (contexto — equivalentes a PES / SWT / TOWS / CF do SGA) · PIQ (partes) · KPI-Q / OBJ-Q / PMM (desempenho) · MOC-Q (alterações) · CQ / FOR-Q / CON (competência) · EQM / CAL / OOT / GRR (metrologia) · RQC / ALR / PRC (cliente) · DD / DC (design) · CONC (concessões) · AUD-Q / CON-Q / CHK (auditoria) · RPG (revisão) · CAPA-Q / 8D (ação corretiva) · PRJ-Q / SUG (melhoria)."),
        ("Ligação ao SGA (SGI)", "Contexto (RG-SGA-01), riscos (RG-SGA-02), requisitos legais (RG-SGA-04) e fornecedores (RG-SGA-11) são lidos na construção (prefixo SGA: na integridade referencial). As tabelas do SGQ com assunto igual ao do SGA têm o MESMO nome e os mesmos nomes de coluna (ver Matriz_Harmonizacao_SGI) para a consolidação em templates únicos para a gestão de topo."),
        ("Casos de análise / ML", "1) Previsão da rejeição mensal por máquina (séries temporais). 2) Classificação da probabilidade de rejeição de lote (tbl_libertacao) por máquina, turno, produto e mês. 3) Sobrevivência do tempo até fecho de CAPA/reclamação. 4) Scorecard preditivo de fornecedores. 5) Texto das NC/8D para agrupar causas."),
        ("Atualização", "python _build/build_all.py reconstrói tudo a partir do dataset e recalcula no Excel; python _build/build_all.py 05 13 reconstrói só os registos indicados. Os registos podem ser editados no Excel; as colunas cinzentas recalculam-se sozinhas."),
        ("Localização (Excel PT)", "As fórmulas evitam TEXT com códigos de data e critérios com casas decimais em texto (dependem do idioma do Excel); usam números de série e concatenação."),
    ]
    for k, (a, t) in enumerate(lines):
        form_block(ws, 3 + k, a, t, vw=1, height=62)
    ordem = ["Ficheiros", "Harmonizacao_Resumo", "Matriz_Harmonizacao_SGI", "Catalogo_Tabelas", "Integridade_Referencial", "Dim_Processo", "Dim_Funcao",
             "Dim_Clausula_ISO9001_2026", "Matriz_Fonte_Unica", "Matriz_ISO9001_2026", "Resumo_Prontidao", "Matriz_Mestra_Info_Doc", "Normas_Fontes", "Modelo_Dados"]
    b.wb._sheets = sorted(b.wb._sheets, key=lambda w: ordem.index(w.title) if w.title in ordem else -1)
    return b.save(out)


if __name__ == "__main__":
    import sys
    print(build(sys.argv[1], sys.argv[2]))
